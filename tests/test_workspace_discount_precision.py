"""Fixed-2 payment-ratio writes and non-destructive transactional policy migration."""
from __future__ import annotations

import csv
from decimal import Decimal, Inexact, localcontext
from io import BytesIO, StringIO
import json
import sqlite3

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from tokensql.workspace.cli import main as cli_main
from tokensql.workspace.data_routes import create_data_router
from tokensql.workspace.discount import normalize_discount_rate
from tokensql.workspace.schemas import PricingBundle, QuoteCreate, QuotePatch
from tokensql.workspace.store import WorkspaceStore


@pytest.mark.parametrize("source,expected", [
    (None, None), ("", None), ("  ", None), ("0", "0.00"), ("-0.000", "0.00"),
    (0, "0.00"), (1, "1.00"), ("0.8", "0.80"), ("0.15", "0.15"),
    ("0.1555", "0.16"), ("0.005", "0.01"), ("0.004999", "0.00"),
    ("1.005", "1.01"), ("1.2", "1.20"), ("8E-1", "0.80"),
    ("15e-2", "0.15"), ("5e-3", "0.01"), (Decimal("0.155"), "0.16"),
    ("9" * 80 + ".995", "1" + "0" * 80 + ".00"),
    ("1E+1000", "1" + "0" * 1000 + ".00"),
    ("1E-1000", "0.00"),
])
def test_exact_helper_fixed_two_half_up_including_large_and_scientific(source, expected):
    with localcontext() as context:
        context.prec = 6
        context.Emax = 9
        context.traps[Inexact] = True
        result = normalize_discount_rate(source)
    if expected is None:
        assert result is None
    else:
        assert format(result, "f") == expected
        assert result.as_tuple().exponent == -2


@pytest.mark.parametrize("source", [
    "-0.000000000000001", "-1E-1000", "-1", "NaN", "sNaN", "Infinity", "-Infinity",
    "garbage", 0.8, float("nan"), float("inf"), True, False, [], {},
])
def test_invalid_and_pre_round_negative_values_are_never_accepted(source):
    with pytest.raises(ValueError):
        normalize_discount_rate(source)
    with pytest.raises(ValueError):
        QuoteCreate.model_validate({"supplier": "Synthetic Supplier", "model_id": "x", "discount_rate": source})
    with pytest.raises(ValueError):
        QuotePatch.model_validate({"expected_revision": 1, "discount_rate": source})


def test_schema_patch_bundle_and_store_enforce_ratio_only_precision():
    amount = "123.45678901234567890123456789"
    quote = QuoteCreate.model_validate({"supplier": "Synthetic Supplier", "model_id": "new", "amount": amount, "discount_rate": "0.1555"})
    assert quote.model_dump(mode="json")["discount_rate"] == "0.16"
    assert quote.model_dump(mode="json")["amount"] == amount
    patch = QuotePatch.model_validate({"expected_revision": 1, "discount_rate": "0.005"})
    assert patch.model_dump(mode="json", exclude_unset=True)["discount_rate"] == "0.01"
    bundle = PricingBundle.model_validate({"update": [{"id": 1, "expected_revision": 1,
                                                       "changes": {"discount_rate": "0.1555"}}]})
    assert bundle.model_dump(mode="json")["update"][0]["changes"]["discount_rate"] == "0.16"
    store = WorkspaceStore(":memory:")
    try:
        inserted = store.apply_bundle({"create": [quote]})["items"][0]
        assert inserted["amount"] == amount and inserted["raw_discount_rate"] == "0.16"
        changed = store.update_quote(inserted["id"], {"discount_rate": "0.005"}, expected_revision=1)
        assert changed.model_dump(mode="json")["discount_rate"] == "0.01"
        assert str(changed.amount) == amount
        # Even trusted/constructed model instances cannot bypass the storage policy.
        direct = QuoteCreate.model_construct(supplier="Synthetic Supplier", model_id="constructed", discount_rate=Decimal("0.1555"))
        assert store.apply_bundle({"create": [direct]})["items"][0]["discount_rate"] == "0.16"
        copied = WorkspaceStore.from_snapshot(store.catalog_snapshot())
        try:
            assert copied.get_quote(1).model_dump(mode="json")["discount_rate"] == "0.01"
        finally:
            copied.close()
    finally:
        store.close()


def seed_legacy_ratios(path):
    store = WorkspaceStore(path)
    store.apply_bundle({"official_prices": [{"model_id": "x", "amount": "2.123456789"}],
                        "create": [{"supplier": "Synthetic Supplier", "model_id": "x", "amount": "123.456789", "currency": "USD"}
                                   for _ in range(5)]}, import_id="historical")
    for quote_id, ratio in enumerate(["0.1555", "0.8", "0.00", None, "0"], start=1):
        row = store._conn.execute("SELECT payload FROM pricing_quote WHERE id=?", (quote_id,)).fetchone()
        payload = json.loads(row["payload"])
        payload["discount_rate"] = ratio
        payload["conditions"] = {"discount_rate": "0.123456789"}
        store._conn.execute("UPDATE pricing_quote SET payload=?,updated_at='2026-01-01T00:00:00+00:00' WHERE id=?",
                            (json.dumps(payload), quote_id))
    store.close()


def dump(path):
    connection = sqlite3.connect(path)
    try:
        return "\n".join(connection.iterdump())
    finally:
        connection.close()


def test_existing_raw_ratios_migrate_audited_without_touching_other_data_and_reopen_is_noop(tmp_path):
    path = tmp_path / "ratios.sqlite"
    seed_legacy_ratios(path)
    connection = sqlite3.connect(path)
    historical = {name: list(connection.execute(f"SELECT * FROM {name}")) for name in
                  ("pricing_model", "pricing_official", "pricing_import", "pricing_audit",
                   "pricing_context", "pricing_context_history")}
    originals = {row[0]: json.loads(row[1]) for row in
                 connection.execute("SELECT id,payload FROM pricing_quote")}
    connection.close()
    store = WorkspaceStore(path)
    try:
        audits = store._conn.execute("SELECT * FROM pricing_audit WHERE actor='migration:discount-2dp'").fetchall()
        assert [row["entity_id"] for row in audits] == ["1", "2", "5"]
        for row in store._conn.execute("SELECT * FROM pricing_quote"):
            payload = json.loads(row["payload"])
            previous = originals[row["id"]]
            assert {k: v for k, v in payload.items() if k != "discount_rate"} == {
                k: v for k, v in previous.items() if k != "discount_rate"}
            changed = row["id"] in {1, 2, 5}
            assert row["revision"] == (2 if changed else 1)
            assert row["updated_at"] != "2026-01-01T00:00:00+00:00" if changed else row["updated_at"] == "2026-01-01T00:00:00+00:00"
        assert [store.get_quote(i).model_dump(mode="json")["discount_rate"] for i in range(1, 6)] == [
            "0.16", "0.80", "0.00", None, "0.00"]
        for audit in audits:
            old, new = json.loads(audit["old_value"]), json.loads(audit["new_value"])
            quote_id = int(audit["entity_id"])
            assert old["discount_rate"] == originals[quote_id]["discount_rate"]
            assert old["revision"] == 1 and new["revision"] == 2
            assert old["amount"] == new["amount"] == "123.456789"
            assert old["currency"] == new["currency"] == "USD"
        for name, expected in historical.items():
            query = f"SELECT * FROM {name}" + (" WHERE actor!='migration:discount-2dp'" if name == "pricing_audit" else "")
            assert [tuple(row) for row in store._conn.execute(query)] == expected
        assert store._conn.execute("PRAGMA foreign_key_check").fetchall() == []
    finally:
        store.close()
    first = dump(path)
    reopened = WorkspaceStore(path)
    reopened.close()
    assert dump(path) == first


def test_discount_migration_injected_failure_rolls_back_all_rows_and_audits(tmp_path, monkeypatch):
    path = tmp_path / "rollback.sqlite"
    seed_legacy_ratios(path)
    before = dump(path)
    original_audit = WorkspaceStore._audit
    def fail_second(self, kind, key, action, actor, old, new):
        if actor == "migration:discount-2dp" and str(key) == "2":
            raise sqlite3.OperationalError("injected second discount migration audit failure")
        return original_audit(self, kind, key, action, actor, old, new)
    monkeypatch.setattr(WorkspaceStore, "_audit", fail_second)
    with pytest.raises(sqlite3.OperationalError, match="injected second"):
        WorkspaceStore(path)
    assert dump(path) == before
    monkeypatch.undo()
    recovered = WorkspaceStore(path)
    assert recovered.get_quote(1).revision == 2
    recovered.close()


def test_legacy_invalid_negative_does_not_partially_migrate_previous_rows(tmp_path):
    path = tmp_path / "invalid.sqlite"
    seed_legacy_ratios(path)
    connection = sqlite3.connect(path)
    row = json.loads(connection.execute("SELECT payload FROM pricing_quote WHERE id=2").fetchone()[0])
    row["discount_rate"] = "-1E-1000"
    connection.execute("UPDATE pricing_quote SET payload=? WHERE id=2", (json.dumps(row),))
    connection.commit()
    connection.close()
    before = dump(path)
    with pytest.raises(ValueError, match="nonnegative"):
        WorkspaceStore(path)
    assert dump(path) == before


def test_api_crud_preview_csv_and_xlsx_all_return_fixed_two_without_rounding_amount():
    store = WorkspaceStore(":memory:")
    app = FastAPI()
    app.include_router(create_data_router(store))
    try:
        with TestClient(app) as client:
            payload = {"items": [{"supplier": "Synthetic Supplier", "model_id": "x", "amount": "1.23456789", "discount_rate": "0.1555"}]}
            preview = client.post("/api/workspace/quotes/preview", json=payload)
            assert preview.status_code == 200 and preview.json()["items"][0]["discount_rate"] == "0.16"
            assert store.list_quotes()[1] == 0
            created = client.post("/api/workspace/quotes", json=payload).json()["items"][0]
            assert created["discount_rate"] == created["raw_discount_rate"] == "0.16"
            assert created["amount"] == "1.23456789"
            changed = client.patch("/api/workspace/quotes/1", json={
                "expected_revision": 1, "discount_rate": "0.005"})
            assert changed.status_code == 200 and changed.json()["discount_rate"] == "0.01"
            assert changed.json()["amount"] == "1.23456789"
            for invalid in ("-0.001", "NaN", "Infinity", 0.1555):
                response = client.patch("/api/workspace/quotes/1", json={
                    "expected_revision": 2, "discount_rate": invalid})
                assert response.status_code == 422
            csv_rows = list(csv.DictReader(StringIO(client.get("/api/workspace/exports").content.decode("utf-8-sig"))))
            assert csv_rows[0]["discount_rate"] == csv_rows[0]["raw_discount_rate"] == "0.01"
            book = load_workbook(BytesIO(client.get("/api/workspace/exports?format=xlsx").content), read_only=True)
            rows = list(book.active.iter_rows(values_only=True))
            data = dict(zip(rows[0], rows[1]))
            assert data["discount_rate"] == data["raw_discount_rate"] == "0.01"
            assert data["amount"] == "1.23456789"
    finally:
        store.close()


def test_cli_preview_apply_and_idempotent_replay_fixed_two(tmp_path, capsys):
    path, source = tmp_path / "cli.sqlite", tmp_path / "bundle.json"
    source.write_text(json.dumps({"create": [{"supplier": "Synthetic Supplier", "model_id": "cli", "amount": "1.23456789", "discount_rate": "0.1555"}]}))
    for command in ("preview", "apply"):
        args = ["--db", str(path), command, str(source)]
        if command == "apply":
            args += ["--import-id", "fixed2"]
        assert cli_main(args) == 0
        result = json.loads(capsys.readouterr().out)
        assert result["items"][0]["discount_rate"] == "0.16"
        assert result["items"][0]["amount"] == "1.23456789"
    assert cli_main(["--db", str(path), "apply", str(source), "--import-id", "fixed2"]) == 0
    assert json.loads(capsys.readouterr().out)["reused"] is True


def test_reused_receipt_output_normalized_without_rewriting_historical_receipt_or_conditions():
    store = WorkspaceStore(":memory:")
    bundle = {"create": [{"supplier": "Synthetic Supplier", "model_id": "x", "discount_rate": "0.8"}]}
    try:
        result = store.apply_bundle(bundle, import_id="same-policy")
        result["items"][0]["discount_rate"] = result["items"][0]["raw_discount_rate"] = "0.8"
        result["items"][0]["conditions"] = {"discount_rate": "0.123456789"}
        result["catalog"]["quotes"][0]["discount_rate"] = "0.8"
        result["catalog"]["quotes"][0]["raw_discount_rate"] = "0.8"
        original = json.dumps(result)
        store._conn.execute("UPDATE pricing_import SET result=? WHERE import_id='same-policy'", (original,))
        replayed = store.apply_bundle(bundle, import_id="same-policy")
        assert replayed["reused"] is True
        assert replayed["items"][0]["discount_rate"] == replayed["items"][0]["raw_discount_rate"] == "0.80"
        assert replayed["catalog"]["quotes"][0]["discount_rate"] == "0.80"
        assert replayed["catalog"]["quotes"][0]["raw_discount_rate"] == "0.80"
        assert replayed["items"][0]["conditions"]["discount_rate"] == "0.123456789"
        assert store._conn.execute("SELECT result FROM pricing_import").fetchone()[0] == original
    finally:
        store.close()


def test_old_policy_import_hash_conflict_remains_safe_and_writes_nothing(tmp_path):
    from hashlib import sha256
    from tokensql.workspace.store import IdempotencyConflictError

    path = tmp_path / "cross-policy.sqlite"
    store = WorkspaceStore(path)
    raw = {"create": [{"supplier": "Synthetic Supplier", "model_id": "x", "discount_rate": "0.1555"}]}
    store.apply_bundle(raw, import_id="legacy-policy")
    old_bundle = PricingBundle.model_validate(raw).model_dump(mode="json")
    old_bundle["create"][0]["discount_rate"] = "0.1555"
    legacy_hash = sha256(json.dumps(old_bundle, ensure_ascii=False, sort_keys=True,
                                   separators=(",", ":")).encode()).hexdigest()
    store._conn.execute("UPDATE pricing_import SET payload_sha256=?", (legacy_hash,))
    before = dump(path)
    try:
        with pytest.raises(IdempotencyConflictError):
            store.apply_bundle(raw, import_id="legacy-policy")
        assert dump(path) == before
    finally:
        store.close()
