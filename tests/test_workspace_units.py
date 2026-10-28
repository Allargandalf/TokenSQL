"""Closed pricing-unit dictionary, strict storage and official-unit discounts."""
from __future__ import annotations

from decimal import Decimal, localcontext
from io import BytesIO
import json

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from pydantic import ValidationError

from tokensql.workspace.data_routes import create_data_router
from tokensql.workspace.schemas import OfficialPriceCreate, QuoteCreate
from tokensql.workspace.store import WorkspaceStore
from tokensql.workspace.units import PRICING_UNITS, pricing_units_payload
from tests.test_workspace_conversion import factor, rates
from tests.test_workspace_data import model, official, quote


@pytest.fixture
def store(tmp_path):
    value = WorkspaceStore(tmp_path / "units.sqlite")
    value.apply_bundle({"models": [model()], "official_prices": [official()]})
    try:
        yield value
    finally:
        value.close()


def test_dictionary_is_exact_closed_product_and_api(tmp_path):
    expected_ids = [f"{currency}_{scale}" for currency in ("usd", "hkd", "cny")
                    for scale in ("ktoken", "mtoken", "second", "image")] + [
            f"{currency}_{scale}" for currency in ("usd", "hkd", "cny")
            for scale in ("mtoken_hour", "request", "account_month")]
    payload = pricing_units_payload()
    assert payload["version"] == "2"
    assert [item["id"] for item in payload["items"]] == expected_ids
    assert len(PRICING_UNITS) == 21
    assert {(item["currency"], item["unit"], item["unit_quantity"])
            for item in payload["items"]} == {
        (currency, unit, quantity) for currency in ("USD", "HKD", "CNY")
        for unit, quantity in (("token", "1000"), ("token", "1000000"),
                               ("second", "1"), ("image", "1"), ("token_hour", "1000000"),
                               ("request", "1"), ("account_month", "1"))}
    local = WorkspaceStore(tmp_path / "api.sqlite")
    app = FastAPI()
    app.include_router(create_data_router(local))
    try:
        with TestClient(app) as client:
            assert client.get("/api/workspace/pricing-units").json() == payload
    finally:
        local.close()


def test_id_is_unique_truth_and_legacy_split_fields_are_strict():
    canonical = QuoteCreate.model_validate({**quote(), "pricing_unit_id": "cny_ktoken"})
    assert canonical.pricing_unit_id == "cny_ktoken"
    assert canonical.currency == "CNY" and canonical.unit == "token"
    assert canonical.unit_quantity == Decimal("1000")
    id_only = OfficialPriceCreate.model_validate({
        "model_id": "maker/model-a", "billing_item": "input", "amount": "1",
        "pricing_unit_id": "usd_mtoken", "explicit_official": True,
        "source": "official", "source_ref": "page#price",
    })
    assert id_only.currency == "USD" and id_only.unit_quantity == Decimal("1000000")
    with pytest.raises(ValidationError, match="conflicts"):
        QuoteCreate.model_validate({**quote(), "pricing_unit_id": "usd_ktoken", "currency": "HKD"})
    with pytest.raises(ValidationError, match="21 pricing units"):
        QuoteCreate.model_validate(quote(unit="unsupported_legacy_request", unit_quantity="1"))
    with pytest.raises(ValidationError):
        QuoteCreate.model_validate({**quote(), "pricing_unit_id": "eur_ktoken"})


def test_store_persists_only_unit_id_but_projects_compatibility_fields(store):
    item = store.apply_bundle({"create": [quote()]})["items"][0]
    raw = json.loads(store._conn.execute("SELECT payload FROM pricing_quote WHERE id=?", (item["id"],)).fetchone()[0])
    assert raw["pricing_unit_id"] == "cny_ktoken"
    assert not {"currency", "unit", "unit_quantity"}.intersection(raw)
    official_raw = json.loads(store._conn.execute(
        "SELECT payload FROM pricing_official WHERE model_id=? AND billing_item=?",
        ("maker/model-a", "input")).fetchone()[0])
    assert official_raw["pricing_unit_id"] == "cny_mtoken"
    assert not {"currency", "unit", "unit_quantity"}.intersection(official_raw)
    assert item["currency"] == "CNY" and item["unit_quantity"] == "1000"
    assert item["official_price"]["pricing_unit_id"] == "cny_mtoken"


def test_same_full_usd_unit_is_preserved_without_backend_discount(store):
    store.apply_bundle({"official_prices": [official(amount="2", pricing_unit_id="usd_mtoken", expected_revision=1)]})
    item = store.apply_bundle({"create": [quote(amount="1.5", pricing_unit_id="usd_mtoken", unit_quantity="1000000")]})["items"][0]
    assert item["amount"] == "1.5" and item["official_price"]["amount"] == "2"
    assert item["discount_rate"] is None and item["converted_supplier_amount"] is None
    assert item["discount_evidence"] == {} and item["normalized"] == {}


def test_token_scale_and_cross_currency_are_raw_frontend_inputs(store):
    store.apply_bundle({"official_prices": [official(amount="1", pricing_unit_id="usd_mtoken", expected_revision=1)]})
    store.set_pricing_context(fx_rates=rates(), token_equivalences=[], expected_revision=1)
    item = store.apply_bundle({"create": [quote(amount="0.0054", pricing_unit_id="hkd_ktoken")]})["items"][0]
    assert item["amount"] == "0.0054" and item["pricing_unit_id"] == "hkd_ktoken"
    assert item["official_price"]["pricing_unit_id"] == "usd_mtoken"
    assert item["converted_supplier_amount"] is None and item["discount_rate"] is None


def test_cross_dimension_context_cannot_generate_stored_or_server_prices(store):
    store.set_pricing_context(fx_rates=[], token_equivalences=[factor()], expected_revision=1)
    item = store.apply_bundle({"create": [quote(amount="0.001", pricing_unit_id="cny_second",
                                               unit="second", unit_quantity="1")]})["items"][0]
    assert item["amount"] == "0.001" and item["pricing_unit_id"] == "cny_second"
    assert item["converted_supplier_amount"] is None and item["discount_rate"] is None
    assert item["conversion_evidence"] == {} and item["discount_evidence"] == {}


def test_cross_dimension_and_currency_keep_source_discount_independent(store):
    store.apply_bundle({"official_prices": [official(amount="2", pricing_unit_id="usd_mtoken", expected_revision=1)]})
    store.set_pricing_context(fx_rates=rates(), token_equivalences=[factor()], expected_revision=1)
    item = store.apply_bundle({"create": [quote(amount="0.001", discount_rate="0.91", pricing_unit_id="hkd_second",
                                               unit="second", unit_quantity="1")]})["items"][0]
    assert item["amount"] == "0.001" and item["discount_rate"] == "0.91"
    assert item["raw_discount_rate"] == "0.91" and item["derived_amount"] is None
    assert item["conversion_evidence"] == {} and item["discount_evidence"] == {}


def test_missing_factor_never_computes_discount_or_converted_amount(store):
    item = store.apply_bundle({"create": [quote(amount="0.001", pricing_unit_id="cny_image",
                                                  unit="image", unit_quantity="1")]})["items"][0]
    assert item["converted_supplier_amount"] is None
    assert item["conversion_status"] == "unavailable"
    assert item["discount_rate"] is None and item["discount_status"] == "unavailable"
    assert item["issues"] == []


def test_unknown_old_unit_is_preserved_with_issue_and_not_guessed(store):
    created = store.apply_bundle({"create": [quote()]})["items"][0]
    raw = json.loads(store._conn.execute("SELECT payload FROM pricing_quote WHERE id=?",
                                        (created["id"],)).fetchone()[0])
    raw.pop("pricing_unit_id")
    raw.update(currency="CNY", unit="unsupported_legacy_request", unit_quantity="1")
    store._conn.execute("UPDATE pricing_quote SET payload=? WHERE id=?",
                        (json.dumps(raw), created["id"]))
    legacy = store.get_quote(created["id"])
    assert legacy.pricing_unit_id is None
    assert legacy.legacy_unit == {"currency": "CNY", "unit": "unsupported_legacy_request", "unit_quantity": "1"}
    assert legacy.converted_supplier_amount is None and legacy.discount_rate is None
    assert "unknown_legacy_pricing_unit" in {entry["code"] for entry in legacy.issues}
    preserved = json.loads(store._conn.execute("SELECT payload FROM pricing_quote WHERE id=?",
                                              (created["id"],)).fetchone()[0])
    assert preserved["unit"] == "unsupported_legacy_request" and "pricing_unit_id" not in preserved
    preview = store.preview_bundle({"create": [quote(supplier="Another supplier")]})
    assert preview["changes"]["added"] == 1
    assert store.list_quotes()[1] == 1


def test_unknown_old_id_without_split_fields_is_readable_and_preserved(store):
    created = store.apply_bundle({"create": [quote()]})["items"][0]
    raw = json.loads(store._conn.execute("SELECT payload FROM pricing_quote WHERE id=?",
                                        (created["id"],)).fetchone()[0])
    raw["pricing_unit_id"] = "future_vendor_unit"
    store._conn.execute("UPDATE pricing_quote SET payload=? WHERE id=?",
                        (json.dumps(raw), created["id"]))
    legacy = store.get_quote(created["id"])
    assert legacy.pricing_unit_id is None
    assert legacy.currency is None and legacy.unit is None and legacy.unit_quantity is None
    assert legacy.legacy_unit == {"pricing_unit_id": "future_vendor_unit"}
    assert legacy.pricing_issue == "unknown_legacy_pricing_unit"
    preview = store.preview_bundle({"create": [quote(supplier="Another supplier")]})
    assert preview["changes"]["added"] == 1
    preserved = json.loads(store._conn.execute("SELECT payload FROM pricing_quote WHERE id=?",
                                              (created["id"],)).fetchone()[0])
    assert preserved["pricing_unit_id"] == "future_vendor_unit"


def test_backend_has_no_display_state_or_derived_formal_discount(store):
    store.apply_bundle({"official_prices": [official(amount="10", pricing_unit_id="usd_second",
        unit="second", unit_quantity="1", expected_revision=1)]})
    item = store.apply_bundle({"create": [quote(amount="8", pricing_unit_id="usd_second",
        unit="second", unit_quantity="1")]})["items"][0]
    assert item["discount_rate"] is None and item["discount_status"] == "unavailable"
    assert item["issues"] == [] and item["display_issues"] == [] and item["normalized"] == {}


def test_unknown_old_official_unit_survives_preview_clone(store):
    row = store._conn.execute(
        "SELECT payload FROM pricing_official WHERE model_id=? AND billing_item=?",
        ("maker/model-a", "input")).fetchone()
    raw = json.loads(row["payload"])
    raw.pop("pricing_unit_id")
    raw.update(currency="CNY", unit="unsupported_legacy_request", unit_quantity="1")
    store._conn.execute(
        "UPDATE pricing_official SET payload=? WHERE model_id=? AND billing_item=?",
        (json.dumps(raw), "maker/model-a", "input"))
    snapshot = store.catalog_snapshot()
    price = snapshot["official_prices"][0]
    assert price["pricing_unit_id"] is None and price["unit"] == "unsupported_legacy_request"
    assert price["legacy_unit"]["unit"] == "unsupported_legacy_request" and "normalized" not in price
    preview = store.preview_bundle({"create": [quote()]})
    assert preview["items"][0]["discount_rate"] is None
    assert preview["items"][0]["official_price"]["legacy_unit"]["unit"] == "unsupported_legacy_request"


def test_view_export_is_memory_only_formula_safe_and_real_xlsx(store):
    app = FastAPI()
    app.include_router(create_data_router(store))
    body = {
        "filename": "converted-view", "columns": [
            {"key": "name", "label": "Name"}, {"key": "amount", "label": "Amount"},
            {"key": "evidence", "label": "Evidence"}],
        "rows": [{"name": "=DANGEROUS()", "amount": "0.12345678901234567890123456789",
                  "evidence": {"factor": "100", "estimated": True}}],
    }
    before = store.catalog_snapshot()
    with TestClient(app) as client:
        csv_response = client.post("/api/workspace/exports/view", json={**body, "format": "csv"})
        assert csv_response.status_code == 200
        assert "converted-view.csv" in csv_response.headers["content-disposition"]
        text = csv_response.content.decode("utf-8-sig")
        assert "'=DANGEROUS()" in text
        assert "0.12345678901234567890123456789" in text
        xlsx_response = client.post("/api/workspace/exports/view", json={**body, "format": "xlsx"})
        assert xlsx_response.status_code == 200 and xlsx_response.content.startswith(b"PK")
        workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True, data_only=False)
        sheet = workbook["View"]
        values = list(sheet.iter_rows(values_only=True))
        assert values[0] == ("Name", "Amount", "Evidence")
        assert values[1][0] == "'=DANGEROUS()"
        assert values[1][1] == "0.12345678901234567890123456789"
        assert json.loads(values[1][2]) == {"estimated": True, "factor": "100"}
        rejected = client.post("/api/workspace/exports/view", json={**body, "format": "csv",
            "rows": [{"name": "x", "amount": 0.1}]})
        assert rejected.status_code == 422
    assert store.catalog_snapshot() == before


def test_patch_pricing_unit_id_replaces_old_projection_and_revalidates(store):
    created = store.apply_bundle({"create": [quote(amount="8", pricing_unit_id="cny_mtoken",
                                                     unit_quantity="1000000")]})["items"][0]
    usd = store.update_quote(created["id"], {"pricing_unit_id": "usd_ktoken"},
                             expected_revision=1)
    assert usd.pricing_unit_id == "usd_ktoken" and usd.currency == "USD"
    assert usd.unit == "token" and usd.unit_quantity == Decimal("1000")
    assert usd.amount == Decimal("8") and usd.discount_rate is None
    second = store.update_quote(created["id"], {"pricing_unit_id": "cny_second"},
                                expected_revision=2)
    assert second.pricing_unit_id == "cny_second" and second.unit == "second"
    assert second.unit_quantity == Decimal("1")
    assert second.discount_rate is None and second.conversion_status == "unavailable"
    with pytest.raises(ValueError, match="conflicts"):
        store.update_quote(created["id"], {"pricing_unit_id": "usd_ktoken", "currency": "HKD"},
                           expected_revision=3)


def test_frontend_temporary_factor_is_not_a_business_payload_field():
    with pytest.raises(ValidationError):
        QuoteCreate.model_validate({**quote(), "display_factor": "123"})
