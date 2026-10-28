"""Nullable source records and transactional in-place compatibility for populated SQLite."""
from __future__ import annotations

from contextlib import contextmanager
from decimal import Decimal
from io import BytesIO
import json
import sqlite3

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from tokensql.workspace.data_routes import create_data_router
from tokensql.workspace.schemas import ModelCreate, OfficialPriceCreate, QuoteCreate
from tokensql.workspace.store import RevisionConflictError, WorkspaceStore

OLD_DDL = "\n                CREATE TABLE IF NOT EXISTS pricing_vendor (\n                    vendor_id TEXT PRIMARY KEY, name TEXT NOT NULL);\n                CREATE TABLE IF NOT EXISTS pricing_model (\n                    model_id TEXT PRIMARY KEY, vendor_id TEXT NOT NULL REFERENCES pricing_vendor(vendor_id),\n                    physical_key TEXT NOT NULL, payload TEXT NOT NULL, revision INTEGER NOT NULL,\n                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL,\n                    UNIQUE(vendor_id, physical_key));\n                CREATE TABLE IF NOT EXISTS pricing_alias (\n                    vendor_id TEXT NOT NULL REFERENCES pricing_vendor(vendor_id),\n                    alias_key TEXT NOT NULL, model_id TEXT NOT NULL REFERENCES pricing_model(model_id),\n                    PRIMARY KEY(vendor_id,alias_key));\n                CREATE TABLE IF NOT EXISTS pricing_official (\n                    model_id TEXT NOT NULL REFERENCES pricing_model(model_id),\n                    billing_item TEXT NOT NULL, payload TEXT NOT NULL, revision INTEGER NOT NULL,\n                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL,\n                    PRIMARY KEY(model_id,billing_item));\n                CREATE TABLE IF NOT EXISTS pricing_supplier (\n                    supplier_id TEXT PRIMARY KEY, name TEXT NOT NULL, name_key TEXT NOT NULL UNIQUE,\n                    provider_type TEXT NOT NULL DEFAULT 'unknown', revision INTEGER NOT NULL DEFAULT 1);\n                CREATE TABLE IF NOT EXISTS pricing_channel (\n                    channel_id INTEGER PRIMARY KEY AUTOINCREMENT,\n                    supplier_id TEXT NOT NULL REFERENCES pricing_supplier(supplier_id),\n                    model_id TEXT NOT NULL REFERENCES pricing_model(model_id),\n                    supplier_model_id TEXT, channel_key TEXT NOT NULL,\n                    UNIQUE(supplier_id,model_id,channel_key));\n                CREATE TABLE IF NOT EXISTS pricing_quote (\n                    id INTEGER PRIMARY KEY AUTOINCREMENT,\n                    model_id TEXT NOT NULL REFERENCES pricing_model(model_id),\n                    billing_item TEXT NOT NULL,\n                    supplier_id TEXT NOT NULL REFERENCES pricing_supplier(supplier_id),\n                    channel_id INTEGER NOT NULL REFERENCES pricing_channel(channel_id),\n                    payload TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'active',\n                    revision INTEGER NOT NULL DEFAULT 1,\n                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL);\n                CREATE INDEX IF NOT EXISTS ix_pricing_quote_model ON pricing_quote(model_id,billing_item,status);\n                CREATE TABLE IF NOT EXISTS pricing_audit (\n                    id INTEGER PRIMARY KEY AUTOINCREMENT, entity_type TEXT NOT NULL, entity_id TEXT NOT NULL,\n                    action TEXT NOT NULL, actor TEXT NOT NULL, old_value TEXT, new_value TEXT,\n                    created_at TEXT NOT NULL);\n                CREATE TABLE IF NOT EXISTS pricing_context (\n                    id INTEGER PRIMARY KEY CHECK(id=1), revision INTEGER NOT NULL,\n                    payload TEXT NOT NULL, updated_at TEXT NOT NULL);\n                CREATE TABLE IF NOT EXISTS pricing_context_history (\n                    revision INTEGER PRIMARY KEY, payload TEXT NOT NULL,\n                    actor TEXT NOT NULL, created_at TEXT NOT NULL);\n                CREATE TABLE IF NOT EXISTS pricing_import (\n                    import_id TEXT PRIMARY KEY, payload_sha256 TEXT NOT NULL,\n                    result TEXT NOT NULL, created_at TEXT NOT NULL);\n            "


def old_database(path):
    conn = sqlite3.connect(path)
    conn.executescript(OLD_DDL)
    model = {"model_id": "old-model", "vendor": "Old Maker", "physical_key": "old-physical",
             "model_name": "Old Model", "aliases": ["old alias"], "source": None, "source_ref": None, "note": None}
    official = {"model_id": "old-model", "billing_item": "input", "amount": "100",
                "pricing_unit_id": "usd_mtoken", "explicit_official": True,
                "source": "old-sheet", "source_ref": "D2", "conversion": None, "effective_date": None, "note": None}
    quote = {"model_id": "old-model", "billing_item": "input", "supplier": "Old Supplier",
             "amount": "8", "discount_rate": "0.83", "pricing_unit_id": "cny_ktoken"}
    conn.execute("INSERT INTO pricing_vendor VALUES ('oldmaker','Old Maker')")
    conn.execute("INSERT INTO pricing_model VALUES ('old-model','oldmaker','oldphysical',?,4,'created','updated')", (json.dumps(model),))
    conn.execute("INSERT INTO pricing_alias VALUES ('oldmaker','oldalias','old-model')")
    conn.execute("INSERT INTO pricing_official VALUES ('old-model','input',?,5,'created','updated')", (json.dumps(official),))
    conn.execute("INSERT INTO pricing_supplier VALUES ('old-supplier','Old Supplier','oldsupplier','unknown',2)")
    conn.execute("INSERT INTO pricing_channel VALUES (6,'old-supplier','old-model','old-channel','old-channel')")
    conn.execute("INSERT INTO pricing_quote VALUES (7,'old-model','input','old-supplier',6,?,'active',3,'2026-01-01T00:00:00Z','2026-01-02T00:00:00Z')", (json.dumps(quote),))
    conn.execute("INSERT INTO pricing_audit VALUES (11,'quote','7','added','old-import',NULL,?,'created')", (json.dumps(quote),))
    conn.execute("INSERT INTO pricing_import VALUES ('old-import','raw-hash','{}','created')")
    initial = json.dumps(WorkspaceStore._default_context())
    conn.execute("INSERT INTO pricing_context VALUES (1,1,?,'created')", (initial,))
    conn.execute("INSERT INTO pricing_context_history VALUES (1,?,'system','created')", (initial,))
    conn.execute("CREATE TABLE workspace_quote(id INTEGER PRIMARY KEY, payload TEXT)")
    conn.execute("INSERT INTO workspace_quote VALUES (19,'untouched legacy')")
    conn.execute("CREATE INDEX custom_quote_index ON pricing_quote(revision)")
    conn.execute("CREATE TRIGGER custom_quote_trigger AFTER UPDATE ON pricing_quote BEGIN SELECT 1; END")
    conn.execute("UPDATE sqlite_sequence SET seq=999 WHERE name='pricing_quote'")
    conn.commit()
    conn.close()


def logical_dump(path):
    conn = sqlite3.connect(path)
    try:
        return "\n".join(conn.iterdump())
    finally:
        conn.close()


def test_populated_not_null_schema_migrates_without_rewriting_business_history(tmp_path):
    path = tmp_path / "old-populated.sqlite"
    old_database(path)
    conn = sqlite3.connect(path)
    names = [row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")]
    before = {table: list(conn.execute(f'SELECT * FROM "{table}" ORDER BY rowid')) for table in names}
    auxiliary = list(conn.execute("SELECT type,name,sql FROM sqlite_master WHERE name LIKE 'custom_%' ORDER BY name"))
    conn.close()
    store = WorkspaceStore(path)
    try:
        after = {table: [tuple(row) for row in store._conn.execute(f'SELECT * FROM "{table}" ORDER BY rowid')] for table in names}
        before["sqlite_sequence"].sort()
        after["sqlite_sequence"].sort()
        previous_model = before["pricing_model"][0]
        payload = json.loads(previous_model[3])
        payload.pop("physical_key")
        before["pricing_model"] = [(previous_model[0], previous_model[1],
            json.dumps(payload, ensure_ascii=False, separators=(",", ":")), *previous_model[4:])]
        for table in set(names) - {"pricing_official", "pricing_quote", "pricing_audit", "sqlite_sequence"}:
            assert after[table] == before[table]
        for previous, current in zip(before["pricing_official"], after["pricing_official"], strict=True):
            assert current[1:3] == previous[:2]
            assert current[4] == previous[3] + 1 and current[5] == previous[4]
            old_payload, new_payload = json.loads(previous[2]), json.loads(current[3])
            assert all(new_payload[key] == value for key, value in old_payload.items())
            assert new_payload["official_price_id"] == current[0]
        for previous, current in zip(before["pricing_quote"], after["pricing_quote"], strict=True):
            assert current[:5] == previous[:5] and current[6] == previous[6]
            assert current[7] == previous[7] + 1 and current[8] == previous[8]
            old_payload, new_payload = json.loads(previous[5]), json.loads(current[5])
            assert all(new_payload[key] == value for key, value in old_payload.items())
        assert after["pricing_audit"][:len(before["pricing_audit"])] == before["pricing_audit"]
        assert dict(after["sqlite_sequence"])["pricing_quote"] == dict(before["sqlite_sequence"])["pricing_quote"]
        assert store._conn.execute("SELECT count(*) FROM pricing_audit WHERE actor='migration:pricing-conditions'").fetchone()[0] == 2
        assert "physical_key" not in {row["name"] for row in store._conn.execute("PRAGMA table_info(pricing_model)")}
        assert [tuple(row) for row in store._conn.execute(
            "SELECT type,name,sql FROM sqlite_master WHERE name LIKE 'custom_%' ORDER BY name")] == auxiliary
        assert store._conn.execute("PRAGMA foreign_key_check").fetchall() == []
        for table, fields in {"pricing_model": {"vendor_id"},
                              "pricing_official": {"billing_item"},
                              "pricing_quote": {"billing_item", "supplier_id", "channel_id"}}.items():
            assert all(not row["notnull"] for row in store._conn.execute(f"PRAGMA table_info({table})") if row["name"] in fields)
        prior = store.get_quote(7)
        assert prior.revision == 4 and prior.amount == Decimal("8") and prior.discount_rate == Decimal("0.83")
        assert prior.official_price["amount"] == "100" and prior.derived_amount is None
        new = store.apply_bundle({"create": [{"model_id": "new-only-id"}]})["items"][0]
        assert new["id"] == 1000 and new["supplier_id"] is None and new["vendor"] is None
        assert store._conn.execute("SELECT payload FROM workspace_quote WHERE id=19").fetchone()[0] == "untouched legacy"
    finally:
        store.close()
    reopened = WorkspaceStore(path)
    assert reopened.get_quote(1000).model_id == "new-only-id"
    reopened.close()


def test_schema_migration_failure_rolls_back_all_rebuilt_tables_and_history(tmp_path, monkeypatch):
    path = tmp_path / "migration-rollback.sqlite"
    old_database(path)
    before = logical_dump(path)
    real_connect = sqlite3.connect
    class InterruptedConnection(sqlite3.Connection):
        def execute(self, sql, *args, **kwargs):
            if sql.startswith("ALTER TABLE __nullable_pricing_official"):
                raise sqlite3.OperationalError("synthetic migration interruption")
            return super().execute(sql, *args, **kwargs)
    def connect(*args, **kwargs):
        return real_connect(*args, **kwargs, factory=InterruptedConnection)
    monkeypatch.setattr(sqlite3, "connect", connect)
    with pytest.raises(sqlite3.OperationalError, match="synthetic migration interruption"):
        WorkspaceStore(path)
    monkeypatch.undo()
    assert logical_dump(path) == before
    recovered = WorkspaceStore(path)
    assert recovered.get_quote(7).discount_rate == Decimal("0.83")
    recovered.close()


@pytest.mark.parametrize("record", [
    {"model_id": "mimo-v2.5-pro"},
    {"model_id": "hy4-preview", "discount_rate": "0.8"},
    {"model_id": "Banana Pro", "discount_rate": "0.8", "currency": "CNY", "unit": "Tokens"},
    {"model_id": "partial", "amount": "3", "currency": "HKD"},
    {"model_id": "both", "amount": "3", "discount_rate": "0.93", "pricing_unit_id": "usd_second"},
])
def test_only_model_id_required_and_partial_units_never_guessed(record):
    store = WorkspaceStore(":memory:")
    try:
        item = store.apply_bundle({"create": [record]})["items"][0]
        assert item["vendor"] is None and item["supplier"] is None and item["billing_item"] is None
        expected_rate = record.get("discount_rate")
        if expected_rate is not None:
            expected_rate = format(Decimal(expected_rate).quantize(Decimal("0.01")), "f")
        assert item["amount"] == record.get("amount") and item["discount_rate"] == expected_rate
        assert item["derived_amount"] is None and item["converted_supplier_amount"] is None
        assert item["normalized"] == {}
        raw = json.loads(store._conn.execute("SELECT payload FROM pricing_quote").fetchone()[0])
        if record.get("pricing_unit_id") is None:
            assert item["pricing_unit_id"] is None and item["unit_quantity"] is None
            assert raw["currency"] == record.get("currency", "CNY")
            assert raw["unit"] == ("token" if record.get("unit") == "Tokens" else record.get("unit"))
        assert store._conn.execute("SELECT count(*) FROM pricing_supplier").fetchone()[0] == 0
    finally:
        store.close()


def test_empty_business_scalars_and_null_containers_normalize_without_defaults():
    item = QuoteCreate.model_validate({"model_id": "only-id", **{key: " " for key in (
        "billing_item", "supplier", "supplier_model_id", "amount", "discount_rate", "pricing_basis",
        "pricing_unit_id", "currency", "unit", "unit_quantity", "provider_type", "component_type",
        "source", "source_ref", "note", "effective_date")}, "tags": None, "conditions": None})
    assert item.tags == [] and item.conditions == {}
    assert all(getattr(item, key) is None for key in ("billing_item", "supplier", "amount", "currency", "unit_quantity", "component_type"))
    for payload in ({}, {"model_id": ""}, {"model_id": None}):
        with pytest.raises(ValueError):
            QuoteCreate.model_validate(payload)
    assert ModelCreate.model_validate({"model_id": "x", "vendor": "", "aliases": None}).vendor is None
    assert OfficialPriceCreate.model_validate({"model_id": "x", "source": ""}).source is None


def test_stub_enrichment_and_unknown_vendor_aliases_do_not_invent_or_split_identity():
    store = WorkspaceStore(":memory:")
    try:
        store.apply_bundle({"models": [{"model_id": "a", "aliases": ["same alias"]}, {"model_id": "b", "aliases": ["same alias"]}]})
        assert store._conn.execute("SELECT count(*) FROM pricing_vendor").fetchone()[0] == 0
        assert store._conn.execute("SELECT count(*) FROM pricing_alias").fetchone()[0] == 0
        store.apply_bundle({"models": [{"model_id": "a", "vendor": "Maker", "expected_revision": 1}]})
        with pytest.raises(ValueError, match="split"):
            store.apply_bundle({"models": [{"model_id": "b", "vendor": "Maker", "expected_revision": 1}]})
        assert store.catalog_snapshot()["models"][1]["vendor"] is None
    finally:
        store.close()


def test_null_billing_official_is_unique_and_never_matches_input_price():
    store = WorkspaceStore(":memory:")
    try:
        store.apply_bundle({"official_prices": [{"model_id": "x", "billing_item": "input", "amount": "10", "pricing_unit_id": "cny_mtoken"}]})
        item = store.apply_bundle({"create": [{"model_id": "x", "discount_rate": "0.8"}]})["items"][0]
        assert item["official_price"] is None
        store.apply_bundle({"official_prices": [{"model_id": "x", "amount": "20"}]})
        unknown = store.get_quote(item["id"])
        assert unknown.official_price is None and unknown.official_match.status == "needs_scope"
        assert len(unknown.official_match.candidate_ids) == 2
        with pytest.raises(RevisionConflictError):
            store.apply_bundle({"official_prices": [{"model_id": "x", "billing_item": "", "amount": "30"}]})
        store.apply_bundle({"official_prices": [{"model_id": "x", "amount": "30",
            "pricing_conditions": {"market": "domestic"}}]})
        assert len(store.catalog_snapshot()["official_prices"]) == 3
        identifier = store.catalog_snapshot()["official_prices"][0]["official_price_id"]
        with pytest.raises(sqlite3.IntegrityError):
            store._conn.execute("INSERT INTO pricing_official VALUES (?,'x',NULL,'{}',1,'x','x')", (identifier,))
    finally:
        store.close()


def test_null_patch_clears_fields_and_independent_raw_values_are_never_refilled():
    store = WorkspaceStore(":memory:")
    try:
        created = store.apply_bundle({"create": [{"model_id": "x", "amount": "5", "discount_rate": "0.8",
            "supplier": "S", "billing_item": "input", "pricing_unit_id": "cny_mtoken"}]})["items"][0]
        item = store.update_quote(created["id"], {"pricing_unit_id": None, "supplier": None, "billing_item": None}, expected_revision=1)
        assert item.pricing_unit_id is None and item.currency is None and item.unit is None and item.unit_quantity is None
        assert item.supplier_id is None and item.channel_id is None
        assert item.amount == 5 and item.discount_rate == Decimal("0.8")
        item = store.update_quote(item.id, {"amount": None, "currency": "USD", "unit": "Tokens"}, expected_revision=2)
        assert item.amount is None and item.discount_rate == Decimal("0.8")
        assert item.currency == "USD" and item.unit == "token" and item.unit_quantity is None
        item = store.update_quote(item.id, {"discount_rate": None, "unit": None}, expected_revision=3)
        assert item.discount_rate is None and item.amount is None and item.unit is None and item.currency == "USD"
        rebuilt = WorkspaceStore.from_snapshot(store.catalog_snapshot())
        assert rebuilt.catalog_snapshot() == store.catalog_snapshot()
        rebuilt.close()
        assert all(value is not None for values in store.filter_options().values() for value in values)
    finally:
        store.close()


def test_partial_api_crud_preview_and_raw_exports_keep_discount_without_unit():
    store = WorkspaceStore(":memory:")
    app = FastAPI()
    app.include_router(create_data_router(store))
    try:
        with TestClient(app) as client:
            before = store.catalog_snapshot()
            preview = client.post("/api/workspace/quotes/preview", json={"items": [{"model_id": "Banana Pro", "discount_rate": "0.8"}]})
            assert preview.status_code == 200 and preview.json()["items"][0]["pricing_unit_id"] is None
            assert store.catalog_snapshot() == before
            created = client.post("/api/workspace/quotes", json={"items": [{"model_id": "Banana Pro", "discount_rate": "0.8"}]})
            assert created.status_code == 201, created.text
            item = created.json()["items"][0]
            assert client.get("/api/workspace/quotes").json()["filter_options"]["vendors"] == []
            assert client.patch(f"/api/workspace/quotes/{item['id']}", json={"expected_revision": 1, "discount_rate": None}).status_code == 200
            assert client.patch(f"/api/workspace/quotes/{item['id']}", json={"expected_revision": 2, "discount_rate": "0.8"}).status_code == 200
            csv = client.get("/api/workspace/exports").content.decode("utf-8-sig")
            assert "Banana Pro" in csv and "0.8" in csv and "raw_discount_rate" in csv
            xlsx = client.get("/api/workspace/exports?format=xlsx")
            workbook = load_workbook(BytesIO(xlsx.content), read_only=True)
            rows = list(workbook.active.iter_rows(values_only=True))
            data = dict(zip(rows[0], rows[1]))
            assert data["discount_rate"] == "0.80" and data["pricing_unit_id"] is None and data["amount"] is None
            assert client.delete(f"/api/workspace/quotes/{item['id']}?expected_revision=3").status_code == 200
    finally:
        store.close()


def test_legacy_unknown_unit_metadata_edit_keeps_unit_evidence():
    store = WorkspaceStore(":memory:")
    try:
        item = store.apply_bundle({"create": [{"model_id": "x"}]})["items"][0]
        raw = {"model_id": "x", "pricing_unit_id": "future_unit", "discount_rate": "0.8"}
        store._conn.execute("UPDATE pricing_quote SET payload=? WHERE id=?", (json.dumps(raw), item["id"]))
        updated = store.update_quote(item["id"], {"note": "human correction"}, expected_revision=1)
        assert updated.legacy_unit == {"pricing_unit_id": "future_unit"}
        assert updated.discount_rate == Decimal("0.8") and updated.note == "human correction"
        preserved = json.loads(store._conn.execute("SELECT payload FROM pricing_quote").fetchone()[0])
        assert preserved["pricing_unit_id"] == "future_unit"
    finally:
        store.close()
