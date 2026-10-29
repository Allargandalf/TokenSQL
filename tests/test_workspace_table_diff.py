"""Content-based source comparison and confirmed-cache acceptance (synthetic only)."""
from copy import deepcopy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pytest

from tokensql.workspace.store import IdempotencyConflictError, RevisionConflictError, WorkspaceStore
from tokensql.workspace.table_diff import (PARSER_VERSION, build_reuse_plan, compare_snapshots,
                                          prompt_fingerprint, read_workbooks)


def workbook(path, rows=None, sheet="Prices"):
    book = Workbook()
    source = book.active
    source.title = sheet
    for row in rows or [["Model", "Item", "Price"], ["A", "input", 1],
                        ["B", "output", 2], ["C", "input", 3]]:
        source.append(row)
    book.save(path)
    book.close()
    return path


def snapshot(path):
    return read_workbooks([path])


def candidate(model="A", row=2, source=None, **extra):
    return {"model_id": model, "amount": str(row), "supplier": "Synthetic",
            "source_sheet": "Prices", "source_ref": "Prices!C" + str(row),
            **({"source": source} if source else {}), **extra}


def baseline(old, entries=None):
    return {"snapshot": old, "bundle": {"create": entries if entries is not None else [
        candidate("A", 2), candidate("B", 3), candidate("C", 4)]},
        "file_refs": [], "profile_revision": 1, "profile_fingerprint": "profile",
        "prompt_fingerprint": prompt_fingerprint("Import"), "parser_version": PARSER_VERSION}


def reuse(base, new, diff=None, **extra):
    return build_reuse_plan(base, new, diff, profile_fingerprint="profile",
                            prompt_fingerprint=prompt_fingerprint("Import"), **extra)


def test_same_bytes_and_display_only_changes_reuse_whole_bundle_without_provenance(tmp_path):
    path = workbook(tmp_path / "original.xlsx")
    old = snapshot(path)
    book = load_workbook(path)
    book.active["C2"].font = Font(name="Arial", bold=True, size=18)
    book.active.row_dimensions[2].height = 45
    book.active.column_dimensions["C"].width = 29
    book.active["C2"].number_format = "0.0000"
    changed = tmp_path / "renamed.xlsx"
    book.save(changed)
    new = snapshot(changed)
    comparison = compare_snapshots(old, new)
    assert not comparison["semantic_changed"]
    assert comparison["groups"]["presentation"]
    assert old["files"][0]["sha256"] != new["files"][0]["sha256"]
    base = baseline(old, [{"model_id": "A", "amount": "1", "supplier": "Synthetic"}])
    plan = reuse(base, new, comparison)
    assert not plan["needs_model"] and plan["reuse_mode"] == "all"
    assert plan["stats"]["reused_entries"] == 1


def test_reordered_rows_small_price_change_and_insertion_only_reinterpret_changed_rows(tmp_path):
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["C", "input", 3], ["A", "input", 1],
        ["B", "output", 2.5], ["D", "input", 4]]))
    comparison = compare_snapshots(old, new)
    assert len(comparison["groups"]["modified"]) == 1
    assert len(comparison["groups"]["added"]) == 1
    assert len(comparison["groups"]["moved"]) == 2
    plan = reuse(baseline(old), new, comparison)
    assert plan["reuse_mode"] == "partial" and plan["needs_model"]
    assert [entry["model_id"] for entry in plan["reused_bundle"]["create"]] == ["A", "C"]
    assert [entry["source_ref"] for entry in plan["reused_bundle"]["create"]] == ["'Prices'!C3", "'Prices'!C2"]
    rows = [row["row"] for region in plan["changed_context"]["regions"]
            for row in region["new_rows"] if row["parse"]]
    assert rows == [4, 5]
    assert plan["stats"]["reused_entries"] == 2


def test_deletion_only_and_reorder_need_no_semantic_model_call(tmp_path):
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["C", "input", 3], ["A", "input", 1]]))
    comparison = compare_snapshots(old, new)
    assert len(comparison["groups"]["deleted"]) == 1
    plan = reuse(baseline(old), new, comparison)
    assert not plan["needs_model"]
    assert {entry["model_id"] for entry in plan["reused_bundle"]["create"]} == {"A", "C"}
    assert comparison["groups"]["moved"]


def test_header_unit_note_percent_currency_and_comment_changes_are_semantic(tmp_path):
    source = workbook(tmp_path / "old.xlsx")
    old = snapshot(source)
    for label, mutate in (
        ("unit", lambda sheet: setattr(sheet["C1"], "value", "USD per million tokens")),
        ("percent", lambda sheet: setattr(sheet["C2"], "number_format", "0.00%")),
        ("currency", lambda sheet: setattr(sheet["C2"], "number_format", '"USD" 0.00')),
        ("comment", lambda sheet: setattr(sheet["C1"], "comment", __import__("openpyxl").comments.Comment("Prices exclude tax", "Synthetic")))):
        book = load_workbook(source)
        mutate(book.active)
        path = tmp_path / (label + ".xlsx")
        book.save(path)
        new = snapshot(path)
        comparison = compare_snapshots(old, new)
        assert comparison["semantic_changed"], label
        if label in {"unit", "comment"}:
            key = new["files"][0]["sheets"][0]["sheet_key"]
            assert comparison["affected_new_rows"][key] == [1, 2, 3, 4]
        else:
            assert comparison["groups"]["modified"][0]["after"]["row"] == 2


def test_merged_multiline_model_and_continuation_rows_expand_as_a_group(tmp_path):
    path = workbook(tmp_path / "old.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], [None, "output", 2],
        ["B", "input", 3]])
    book = load_workbook(path)
    book.active.merge_cells("A2:A3")
    book.save(path)
    old = snapshot(path)
    book.active["C3"] = 4
    current = tmp_path / "new.xlsx"
    book.save(current)
    new = snapshot(current)
    diff = compare_snapshots(old, new)
    key = new["files"][0]["sheets"][0]["sheet_key"]
    assert set(diff["affected_new_rows"][key]) >= {2, 3}


def test_formula_value_dependencies_cross_sheets_and_unknown_names_expand(tmp_path):
    path = workbook(tmp_path / "old.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["B", "output", "=C2*2"]])
    book = load_workbook(path)
    summary = book.create_sheet("Summary")
    summary.append(["Model", "Price"])
    summary.append(["A", "=Prices!C3"])
    summary.append(["B", "=INDIRECT(\"Prices!C2\")"])
    book.save(path)
    old = snapshot(path)
    book["Prices"]["C2"] = 2
    current = tmp_path / "new.xlsx"
    book.save(current)
    new = snapshot(current)
    diff = compare_snapshots(old, new)
    sheets = {sheet["name"]: sheet for sheet in new["files"][0]["sheets"]}
    assert 3 in diff["affected_new_rows"][sheets["Prices"]["sheet_key"]]
    assert set(diff["affected_new_rows"][sheets["Summary"]["sheet_key"]]) >= {2, 3}
    formula = sheets["Summary"]["rows"][1]["cells"][1]
    assert formula["formula"] == "=Prices!C3" and formula["cached_value"] is None
    assert any(region["reason"] == "formula_dependency" for region in diff["regions"])


def test_semantics_or_prompt_changes_force_full_and_missing_provenance_is_safe(tmp_path):
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 4], ["B", "output", 2], ["C", "input", 3]]))
    base = baseline(old, [{"model_id": "A", "amount": "1", "supplier": "Synthetic"}])
    plan = reuse(base, new)
    assert plan["reuse_mode"] == "none" and plan["needs_model"]
    assert plan["reason"] == "source_provenance_requires_full_read"
    assert plan["reused_bundle"]["create"] == []
    base = baseline(old)
    base["profile_fingerprint"] = "different"
    assert reuse(base, old)["reuse_mode"] == "none"
    base["profile_fingerprint"] = "profile"
    base["prompt_fingerprint"] = prompt_fingerprint("Different request")
    assert reuse(base, old)["needs_model"]
    assert reuse(baseline(old), old, force_full=True)["needs_model"]


def test_multiple_files_align_by_content_when_renamed_reordered_and_one_changes(tmp_path):
    first = workbook(tmp_path / "first.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1]], sheet="Prices")
    second = workbook(tmp_path / "second.xlsx", [
        ["Model", "Item", "Price"], ["B", "output", 2]], sheet="Prices")
    old = read_workbooks([first, second])
    renamed = tmp_path / "renamed-first.xlsx"
    renamed.write_bytes(first.read_bytes())
    changed = workbook(tmp_path / "changed-second.xlsx", [
        ["Model", "Item", "Price"], ["B", "output", 4]], sheet="Prices")
    new = read_workbooks([changed, renamed])
    entries = [candidate("A", 2, first.name), candidate("B", 2, second.name)]
    plan = reuse(baseline(old, entries), new)
    assert plan["reuse_mode"] == "partial"
    assert [entry["model_id"] for entry in plan["reused_bundle"]["create"]] == ["A"]
    assert plan["reused_bundle"]["create"][0]["source"] == renamed.name


def test_duplicate_content_ambiguity_is_reinterpreted_and_never_position_mapped(tmp_path):
    old = snapshot(workbook(tmp_path / "old.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["A", "input", 1], ["B", "input", 2]]))
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["B", "input", 4], ["A", "input", 1]]))
    diff = compare_snapshots(old, new)
    plan = reuse(baseline(old, [candidate("A", 2), candidate("A", 3), candidate("B", 4)]), new, diff)
    assert plan["reused_bundle"]["create"] == []
    assert all(item["old"]["row"] == 1 for item in diff["row_matches"] if item["unchanged"])


def test_baseline_scoped_persistent_idempotent_and_never_rolls_back_old_receipts(tmp_path):
    path = tmp_path / "isolated.sqlite3"
    store = WorkspaceStore(path)
    sid = store.confirm_supplier({"name": "Synthetic"})["supplier_id"]
    other = store.confirm_supplier({"name": "Other"})["supplier_id"]
    old = snapshot(workbook(tmp_path / "source.xlsx"))
    payload = baseline(old)
    before = store.catalog_snapshot()
    first = store.record_supplier_baseline(sid, "monthly", "first", payload, expected_revision=0)
    assert first["revision"] == 1 and store.get_supplier_baseline(other, "monthly") is None
    assert store.get_supplier_baseline(sid, "other") is None
    assert store.record_supplier_baseline(sid, "monthly", "first", payload)["revision"] == 1
    changed = deepcopy(payload)
    changed["prompt_fingerprint"] = "new-prompt"
    second = store.record_supplier_baseline(sid, "monthly", "second", changed, expected_revision=1)
    assert second["revision"] == 2
    assert store.record_supplier_baseline(sid, "monthly", "first", payload)["import_id"] == "second"
    with pytest.raises(IdempotencyConflictError):
        store.record_supplier_baseline(sid, "monthly", "first", changed)
    with pytest.raises(RevisionConflictError):
        store.record_supplier_baseline(sid, "monthly", "third", payload, expected_revision=0)
    assert store.catalog_snapshot() == before
    store.close()
    restored = WorkspaceStore(path)
    try:
        assert restored.get_supplier_baseline(sid, "monthly")["import_id"] == "second"
        assert restored.get_supplier_baseline(sid, "monthly")["bundle"]["create"][0]["model_id"] == "A"
    finally:
        restored.close()


def test_baseline_rejects_database_operation_recipe_and_unknown_supplier(tmp_path):
    store = WorkspaceStore(":memory:")
    sid = store.confirm_supplier({"name": "Synthetic"})["supplier_id"]
    payload = baseline(snapshot(workbook(tmp_path / "source.xlsx")))
    payload["bundle"] = {"delete": [{"id": 1, "expected_revision": 1}]}
    with pytest.raises(ValueError, match="recipes"):
        store.record_supplier_baseline(sid, "monthly", "bad", payload)
    assert store.get_supplier_baseline(sid, "monthly") is None
    with pytest.raises(ValueError):
        store.get_supplier_baseline("missing", "monthly")
    store.close()


@pytest.mark.parametrize("values", [
    ("1.00", "2.00", "3.00", "2.50"), ("USD 1.00", "USD 2.00", "USD 3.00", "USD 2.50"),
    ("10%", "20%", "30%", "25%")])
def test_text_numeric_xlsx_change_aligns_and_reuses_unchanged_models(tmp_path, values):
    first, second, third, changed = values
    old = snapshot(workbook(tmp_path / "old.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", first], ["B", "output", second], ["C", "input", third]]))
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["C", "input", third], ["A", "input", first], ["B", "output", changed]]))
    diff = compare_snapshots(old, new)
    plan = reuse(baseline(old), new, diff)
    assert plan["reuse_mode"] == "partial" and plan["stats"]["reused_entries"] == 2
    assert len(diff["groups"]["modified"]) == 1
    assert [row["row"] for region in plan["changed_context"]["regions"]
            for row in region["new_rows"] if row["parse"]] == [4]
    serialized = str(plan["changed_context"])
    assert "font" not in serialized and "semantic_key" not in serialized


def test_csv_text_price_change_is_partial_and_filename_is_not_sheet_identity(tmp_path):
    old_path, new_path = tmp_path / "old.csv", tmp_path / "renamed.csv"
    old_path.write_text("Model,Item,Price\nA,input,1.00\nB,output,2.00\nC,input,3.00\n")
    new_path.write_text("Model,Item,Price\nC,input,3.00\nA,input,1.00\nB,output,2.50\n")
    old, new = snapshot(old_path), snapshot(new_path)
    entries = [candidate("A", 2, source_sheet="CSV", source_ref="CSV!C2"),
               candidate("B", 3, source_sheet="CSV", source_ref="CSV!C3"),
               candidate("C", 4, source_sheet="CSV", source_ref="CSV!C4")]
    plan = reuse(baseline(old, entries), new)
    assert plan["reuse_mode"] == "partial" and plan["stats"]["reused_entries"] == 2


def test_known_model_metadata_is_not_replayed_over_fresh_catalog_on_reuse(tmp_path):
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    base = baseline(old)
    base["bundle"]["models"] = [{"model_id": "A", "model_name": "Prior metadata"}]
    whole = reuse(base, old)
    assert whole["reused_bundle"]["models"] == []
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["B", "output", 4], ["C", "input", 3]]))
    partial = reuse(base, new)
    assert partial["reuse_mode"] == "partial" and partial["reused_bundle"]["models"] == []
    assert partial["stats"]["reused_entries"] == 2


def test_row_reordering_taints_unchanged_formula_references_without_reparsing_plain_rows(tmp_path):
    old = snapshot(workbook(tmp_path / "old.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["B", "input", 3],
        ["Calculated", "output", "=C2*2"]]))
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["B", "input", 3], ["A", "input", 1],
        ["Calculated", "output", "=C2*2"]]))
    diff = compare_snapshots(old, new)
    key = new["files"][0]["sheets"][0]["sheet_key"]
    assert diff["semantic_changed"]
    assert diff["affected_new_rows"][key] == [4]
    plan = reuse(baseline(old, [candidate("A", 2), candidate("B", 3), candidate("Calculated", 4)]), new, diff)
    assert plan["stats"]["reused_entries"] == 2 and plan["needs_model"]
    assert [entry["model_id"] for entry in plan["reused_bundle"]["create"]] == ["A", "B"]


def test_defined_names_and_formula_settings_are_snapshot_semantics(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName
    path = workbook(tmp_path / "old.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["B", "input", 3],
        ["Calculated", "output", "=Rate*2"]])
    book = load_workbook(path)
    book.defined_names.add(DefinedName("Rate", attr_text="'Prices'!$C$2"))
    book.save(path)
    old = snapshot(path)
    book.defined_names["Rate"].attr_text = "'Prices'!$C$3"
    book.calculation.calcMode = "manual"
    current = tmp_path / "new.xlsx"
    book.save(current)
    new = snapshot(current)
    diff = compare_snapshots(old, new)
    assert diff["semantic_changed"]
    key = new["files"][0]["sheets"][0]["sheet_key"]
    assert 4 in diff["affected_new_rows"][key]
    assert any(region["reason"] in {"formula_dependency", "workbook_formula_metadata"}
               for region in diff["regions"])
    metadata = new["files"][0]["workbook_metadata"]
    assert metadata["defined_names"][0]["reference"] == "'Prices'!$C$3"
    assert metadata["calculation"]["calcMode"] == "manual"


def test_structured_table_range_changes_expand_its_context(tmp_path):
    from openpyxl.worksheet.table import Table
    path = workbook(tmp_path / "old.xlsx")
    book = load_workbook(path)
    book.active.add_table(Table(displayName="PricesTable", ref="A1:C4"))
    book.save(path)
    old = snapshot(path)
    book.active.tables["PricesTable"].ref = "A1:C3"
    current = tmp_path / "new.xlsx"
    book.save(current)
    new = snapshot(current)
    diff = compare_snapshots(old, new)
    key = new["files"][0]["sheets"][0]["sheet_key"]
    assert diff["affected_new_rows"][key] == [1, 2, 3, 4]


def test_formula_cached_result_change_is_semantic_even_with_same_formula(tmp_path):
    from zipfile import ZipFile, ZIP_DEFLATED
    path = workbook(tmp_path / "base.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["B", "output", "=C2*2"]])
    def cached(target, result):
        with ZipFile(path) as source, ZipFile(target, "w", ZIP_DEFLATED) as output:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<f>C2*2</f><v></v>", ("<f>C2*2</f><v>" + result + "</v>").encode())
                    data = data.replace(b"<f>C2*2</f><v />", ("<f>C2*2</f><v>" + result + "</v>").encode())
                output.writestr(info, data)
    old_path, new_path = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
    cached(old_path, "2")
    cached(new_path, "4")
    old, new = snapshot(old_path), snapshot(new_path)
    diff = compare_snapshots(old, new)
    assert diff["semantic_changed"]
    changed = diff["groups"]["modified"][0]
    assert changed["before"]["cells"][2]["formula"] == changed["after"]["cells"][2]["formula"]
    assert changed["before"]["cells"][2]["cached_value"] == "2"
    assert changed["after"]["cells"][2]["cached_value"] == "4"


def test_noncontiguous_source_range_relocation_requires_reinterpretation(tmp_path):
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["B", "output", 2], ["C", "input", 3], ["A", "input", 1]]))
    diff = compare_snapshots(old, new)
    assert not diff["semantic_changed"]
    base = baseline(old, [candidate("Composite", 2, source_ref="Prices!A2:C4")])
    plan = reuse(base, new, diff)
    assert plan["needs_model"] and plan["reuse_mode"] == "none"
    assert plan["reason"] == "source_region_requires_full_read"


def test_invalid_workbook_has_safe_full_read_fallback(tmp_path):
    path = tmp_path / "fake.xlsx"
    path.write_bytes(b"synthetic bytes accepted by provider fixture")
    result = snapshot(path)
    assert not result["supported"] and result["issues"]
    diff = compare_snapshots(result, result)
    assert diff["requires_full"]


def test_shared_notes_sheet_change_invalidates_related_workbook_prices(tmp_path):
    path = workbook(tmp_path / "old.xlsx")
    book = load_workbook(path)
    guide = book.create_sheet("Guide")
    guide.append(["Currency is CNY; applies to every price sheet"])
    book.save(path)
    old = snapshot(path)
    book["Guide"]["A1"] = "Currency is USD; applies to every price sheet"
    current = tmp_path / "new.xlsx"
    book.save(current)
    new = snapshot(current)
    diff = compare_snapshots(old, new)
    prices = next(sheet for sheet in new["files"][0]["sheets"] if sheet["name"] == "Prices")
    assert diff["affected_new_rows"][prices["sheet_key"]] == [1, 2, 3, 4]
    plan = reuse(baseline(old), new, diff)
    assert plan["reused_bundle"]["create"] == [] and plan["needs_model"]


def test_sixteen_text_prices_reversed_one_change_reuses_fifteen_entries(tmp_path):
    rows = [["Model", "Item", "Price"]] + [["Model-" + str(index), "input", str(index + 1)]
                                            for index in range(16)]
    old = snapshot(workbook(tmp_path / "old.xlsx", rows))
    reordered = [list(row) for row in reversed(rows[1:])]
    changed = next(row for row in reordered if row[0] == "Model-4")
    changed[2] = "5.25"
    new = snapshot(workbook(tmp_path / "new.xlsx", [rows[0], *reordered]))
    entries = [candidate("Model-" + str(index), index + 2) for index in range(16)]
    base = baseline(old, entries)
    base["bundle"]["models"] = [{"model_id": "Model-" + str(index)} for index in range(16)]
    diff = compare_snapshots(old, new)
    plan = reuse(base, new, diff)
    assert plan["reuse_mode"] == "partial" and plan["stats"]["reused_entries"] == 15
    assert sum(row["parse"] for region in plan["changed_context"]["regions"]
               for row in region["new_rows"]) == 1
    assert plan["reused_bundle"]["models"] == []


def test_local_defined_name_changes_reparse_its_formula_scope(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName
    path = workbook(tmp_path / "old.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["B", "input", 3],
        ["Calculated", "output", "=Rate*2"]])
    book = load_workbook(path)
    book.active.defined_names.add(DefinedName("Rate", attr_text="'Prices'!$C$2", localSheetId=0))
    book.save(path)
    old = snapshot(path)
    book.active.defined_names["Rate"].attr_text = "'Prices'!$C$3"
    current = tmp_path / "new.xlsx"
    book.save(current)
    new = snapshot(current)
    diff = compare_snapshots(old, new)
    key = new["files"][0]["sheets"][0]["sheet_key"]
    assert diff["semantic_changed"] and 4 in diff["affected_new_rows"][key]


@pytest.mark.parametrize("old_format,new_format,semantic", [
    ("0\\%", "0%", True), ('0"%"', "0%", True), ("0,,", "0", True),
    ("0.00", "0.0", False), ("0.00%", "0.0%", False)])
def test_meaningful_number_format_escapes_scaling_and_precision(tmp_path, old_format, new_format, semantic):
    path = workbook(tmp_path / "old.xlsx")
    book = load_workbook(path)
    book.active["C2"].number_format = old_format
    book.save(path)
    old = snapshot(path)
    book.active["C2"].number_format = new_format
    current = tmp_path / "new.xlsx"
    book.save(current)
    diff = compare_snapshots(old, snapshot(current))
    assert diff["semantic_changed"] is semantic


def test_biff_formula_detection_proves_literal_streams_and_preserves_unknown_fallback(tmp_path):
    from struct import pack
    from tokensql.workspace.table_diff import _xls_has_formulas
    from xlrd import biffh
    literal = tmp_path / "literal.xls"
    literal.write_bytes(pack("<HH", 0x0809, 4) + b"1234" + pack("<HH", 0x000A, 0))
    assert not _xls_has_formulas(literal)
    formula = tmp_path / "formula.xls"
    formula.write_bytes(pack("<HH", biffh.XL_FORMULA, 0))
    assert _xls_has_formulas(formula)
    broken = tmp_path / "broken.xls"
    broken.write_bytes(pack("<HH", 0x0809, 100) + b"short")
    assert _xls_has_formulas(broken)


def test_reused_official_prices_keep_their_schema_in_whole_and_partial_paths(tmp_path):
    from tokensql.workspace.schemas import PricingBundle
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    base = baseline(old)
    base["bundle"]["official_prices"] = [{"model_id": "A", "billing_item": "input", "amount": "10",
        "pricing_unit_id": "cny_mtoken", "source": "old.xlsx", "source_ref": "Prices!C2"}]
    whole = reuse(base, old)
    PricingBundle.model_validate(whole["reused_bundle"])
    assert "source_sheet" not in whole["reused_bundle"]["official_prices"][0]
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["C", "input", 3], ["A", "input", 1], ["B", "output", 4]]))
    partial = reuse(base, new)
    PricingBundle.model_validate(partial["reused_bundle"])
    assert partial["reuse_mode"] == "partial"
    assert partial["reused_bundle"]["official_prices"][0]["source_ref"] == "'Prices'!C3"
    assert "source_sheet" not in partial["reused_bundle"]["official_prices"][0]


def test_canonical_reference_with_covered_annotation_reuses_and_preserves_note(tmp_path):
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    base = baseline(old, [candidate("A", 2, source_ref="Prices!A2:C2 (unit cell C2='CNY/1M tokens')"),
                          candidate("B", 3), candidate("C", 4)])
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["C", "input", 3], ["A", "input", 1], ["B", "output", 4]]))
    plan = reuse(base, new)
    assert plan["reuse_mode"] == "partial" and plan["stats"]["reused_entries"] == 2
    entry = next(item for item in plan["reused_bundle"]["create"] if item["model_id"] == "A")
    assert entry["source_ref"] == "'Prices'!A3:C3 (unit cell C3='CNY/1M tokens')"
    assert entry.get("note") is None


@pytest.mark.parametrize("annotation", [
    "unit cell C1='CNY/1M tokens'", "secondary Prices!C2", "other cell Z2", "broken (note"])
def test_annotation_outside_canonical_reference_falls_back_safely(tmp_path, annotation):
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    new = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["A", "input", 1], ["B", "output", 4], ["C", "input", 3]]))
    base = baseline(old, [candidate("A", 2, source_ref="Prices!A2:C2 (" + annotation + ")")])
    plan = reuse(base, new)
    assert plan["needs_model"] and plan["reuse_mode"] == "none"


def test_annotation_relocation_preserves_comparable_business_fields(tmp_path):
    from tokensql.workspace.reconciliation import comparable_quote
    old = snapshot(workbook(tmp_path / "old.xlsx"))
    entry = candidate("A", 2, source_ref="Prices!A2:C2 (unit cell C2='CNY/1M tokens')")
    base = baseline(old, [entry, candidate("B", 3), candidate("C", 4)])
    whole = reuse(base, old)["reused_bundle"]["create"][0]
    assert comparable_quote(whole) == comparable_quote(entry)
    moved = snapshot(workbook(tmp_path / "new.xlsx", [
        ["Model", "Item", "Price"], ["C", "input", 3], ["A", "input", 1], ["B", "output", 4]]))
    relocated = next(item for item in reuse(base, moved)["reused_bundle"]["create"] if item["model_id"] == "A")
    assert comparable_quote(relocated) == comparable_quote(entry)
    assert relocated["source_ref"] == "'Prices'!A3:C3 (unit cell C3='CNY/1M tokens')"
    assert relocated.get("note") is None
