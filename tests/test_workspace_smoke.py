from __future__ import annotations

import asyncio
from copy import deepcopy
from io import BytesIO

import httpx
import json
from openpyxl import load_workbook

from tools.workspace_smoke import (
    EXPECTED, SYNTHETIC_FILENAME, _run, _validation_errors, _workbook_bytes,
)


def proposal():
    models = [
        {"model_id": "synthetic/alpha-1", "vendor": "Synthetic", "model_name": "Alpha-1"},
        {"model_id": "synthetic/beta-2", "vendor": "Synthetic", "model_name": "Beta 2"},
    ]
    creates = []
    for (model_name, billing_item), (amount, cell) in EXPECTED.items():
        model_id = "synthetic/alpha-1" if model_name == "Alpha-1" else "synthetic/beta-2"
        creates.append({
            "model_id": model_id, "billing_item": billing_item,
            "supplier": "Synthetic smoke supplier", "provider_type": "reseller",
            "amount": amount, "pricing_unit_id": "cny_mtoken",
            "source": SYNTHETIC_FILENAME, "source_ref": f"Synthetic Quotes!{cell}",
        })
    return {"models": models, "create": creates}


def rows():
    names = {"synthetic/alpha-1": "Alpha-1", "synthetic/beta-2": "Beta 2"}
    return [{**item, "model_name": names[item["model_id"]], "currency": "CNY",
             "unit": "token", "unit_quantity": "1000000"}
            for item in proposal()["create"]]


def valid_result(status="ready"):
    return {
        "status": status, "database_total": 0 if status == "ready" else 4,
        "changes": {"added": 4, "updated": 0, "deleted": 0, "total": 4},
        "proposal": proposal(), "quotes": rows(),
    }


def ready_job():
    return {
        "id": "job-1", "status": "ready", "revision": 2, "bundle": proposal(),
        "preview": {"items": rows(), "changes": {"added": 4, "updated": 0,
                                                    "deleted": 0, "total": 4}},
        "issues": [], "error": None, "conversation_id": "ses_synthetic",
    }


def test_synthetic_workbook_is_memory_generated_and_exact():
    workbook = load_workbook(BytesIO(_workbook_bytes()), read_only=True, data_only=True)
    values = list(workbook["Synthetic Quotes"].iter_rows(values_only=True))
    assert values == [
        ("model", "input price", "output price"),
        ("Alpha-1", "1.25", "3.5"),
        ("Beta 2", "0.8", "2.4"),
    ]


def test_validation_accepts_review_ready_without_writes_and_confirmed_commit():
    assert _validation_errors(valid_result("ready")) == []
    assert _validation_errors(valid_result("confirmed")) == []


def test_validation_rejects_wrong_unit_billing_model_identity_and_source():
    result = deepcopy(valid_result())
    result["quotes"][0]["pricing_unit_id"] = "cny_ktoken"
    result["quotes"][1]["billing_item"] = "completion"
    result["quotes"][2]["model_id"] = "synthetic/alpha-1"
    result["quotes"][3]["model_id"] = "synthetic/alpha-1"
    result["quotes"][3]["source_ref"] = "wrong"
    errors = _validation_errors(result)
    assert any("pricing_unit_id" in error for error in errors)
    assert any("model/billing_item" in error for error in errors)
    assert any("two independent" in error for error in errors)
    assert any("source reference" in error for error in errors)


def mock_transport(*, isolated=True, commit=False):
    calls = []
    quote_reads = 0

    async def handler(request):
        nonlocal quote_reads
        calls.append((request.method, request.url.path))
        path = request.url.path
        if path == "/health":
            payload = {"status": "ok", "application": "tokensql-workspace"}
            if isolated:
                payload["isolated_test"] = True
            return httpx.Response(200, json=payload)
        if path == "/api/workspace/quotes":
            quote_reads += 1
            if commit and quote_reads > 1:
                return httpx.Response(200, json={"total": 4, "items": rows(), "filter_options": {}})
            return httpx.Response(200, json={"total": 0, "items": [], "filter_options": {}})
        if path == "/api/workspace/imports/upload":
            body = await request.aread()
            assert SYNTHETIC_FILENAME.encode() in body
            return httpx.Response(201, json={"items": [{"id": "upload-1"}]})
        if path == "/api/workspace/imports/start":
            body = json.loads(await request.aread())
            assert body["upload_ids"] == ["upload-1"]
            assert "SYNTHETIC QA DATA ONLY" in body["prompt"]
            assert "pricing_unit_id='cny_mtoken'" in body["prompt"]
            return httpx.Response(202, json={"id": "job-1", "status": "queued"})
        if path == "/api/workspace/imports/job-1" and request.method == "GET":
            return httpx.Response(200, json=ready_job())
        if path == "/api/workspace/imports/job-1/cancel":
            return httpx.Response(200, json={"id": "job-1", "status": "cancelled"})
        if path == "/api/workspace/imports/job-1/confirm":
            assert commit
            body = json.loads(await request.aread())
            assert body["expected_revision"] == 2
            return httpx.Response(200, json={
                "id": "job-1", "status": "confirmed",
                "changes": {"added": 4, "updated": 0, "deleted": 0, "total": 4}})
        return httpx.Response(500, json={"unexpected": path})

    return httpx.MockTransport(handler), calls


def test_http_smoke_defaults_to_review_then_cancel_without_database_writes():
    transport, calls = mock_transport(isolated=False)
    result = asyncio.run(_run("http://test", transport=transport, poll_interval=0,
                              timeout_seconds=5))
    assert result["validated"] is True
    assert result["status"] == "ready" and result["database_total"] == 0
    assert result["final_status"] == "cancelled"
    assert result["post_action_database_total"] == 0
    assert ("POST", "/api/workspace/imports/job-1/cancel") in calls
    assert not any(path.endswith("/confirm") for _, path in calls)


def test_confirm_is_refused_before_upload_without_isolated_health_marker():
    transport, calls = mock_transport(isolated=False, commit=True)
    result = asyncio.run(_run("http://test", confirm=True, transport=transport,
                              poll_interval=0, timeout_seconds=5))
    assert result["status"] == "refused" and result["validated"] is False
    assert result["validation_errors"] == ["--confirm requires health.isolated_test=true"]
    assert not any(path.endswith("/upload") for _, path in calls)


def test_confirm_requires_isolation_and_validates_committed_four_rows():
    transport, calls = mock_transport(isolated=True, commit=True)
    result = asyncio.run(_run("http://test", confirm=True, transport=transport,
                              poll_interval=0, timeout_seconds=5))
    assert result["validated"] is True
    assert result["status"] == "confirmed" and result["database_total"] == 4
    assert result["health"]["isolated_test"] is True
    assert ("POST", "/api/workspace/imports/job-1/confirm") in calls
    assert not any(path.endswith("/cancel") for _, path in calls)
