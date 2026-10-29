"""Durable file and proposal lifecycle; formal writes require explicit confirmation."""
from __future__ import annotations

import asyncio
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
from threading import RLock
from typing import Any, Sequence
from uuid import uuid4

from .memory import (AgentMemory, atomic_json, remember_supplier_confirmation,
                     supplier_profile_context, supplier_confirmation_updates)
from .discount import DISCOUNT_POLICY_VERSION, normalize_discount_rate
from .clarification import NeedsInput
from .schemas import PricingBundle
from .store import RevisionConflictError
from .official_lookup import validate_lookup_bundle
from .source_reading import (freeze_source_baseline, merge_source_bundle, public_source_comparison,
                             record_source_baseline, source_coverage, rebind_source_officials)

_SAFE = re.compile(r"[^\w.()\uFF08\uFF09\[\]{} +\-@]", re.UNICODE)
_ZERO = {"added": 0, "updated": 0, "deleted": 0, "total": 0}
PRICING_POLICY_VERSION = 2

def now() -> str:
    return datetime.now(timezone.utc).isoformat()

def clean_filename(name: str) -> str:
    name = name.replace("\\", "/").rsplit("/", 1)[-1]
    name = _SAFE.sub("_", "".join(c for c in name if ord(c) >= 32)).strip(" .")
    name = name or "upload"
    suffix = Path(name).suffix[:20]
    stem = name[:-len(suffix)] if suffix else name
    while len((stem + suffix).encode("utf-8")) > 180:
        stem = stem[:-1]
    return stem + suffix

def current_model_fields(value: Any) -> Any:
    """Drop only the retired structural model field, never source text/conditions."""
    result = deepcopy(value)
    if not isinstance(result, dict):
        return result
    models = result.get("models")
    if isinstance(models, list):
        for item in models:
            if isinstance(item, dict):
                item.pop("physical_key", None)
    # Preview/receipt catalog snapshots are known application structures.
    for key in ("catalog", "preview", "commit_result"):
        if isinstance(result.get(key), dict):
            result[key] = current_model_fields(result[key])
    return result


def legacy_bundle_fields(value: Any) -> Any:
    """Retain old missing currency semantics rather than apply the new CNY default."""
    result = current_model_fields(value)
    if not isinstance(result, dict):
        return result
    for key in ("create", "official_prices"):
        for item in result.get(key) or []:
            if isinstance(item, dict) and "currency" not in item and not item.get("pricing_unit_id"):
                item["currency"] = None
    return result


def normalize_bundle_discounts(value: Any) -> Any:
    """Apply the fixed-two-place raw ratio policy, not monetary derivation."""
    result = current_model_fields(value)
    if not isinstance(result, dict):
        return result
    for key in ("create", "update"):
        rows = result.get(key)
        if not isinstance(rows, list):
            continue
        for item in rows:
            if not isinstance(item, dict):
                continue
            payload = item if key == "create" else item.get("changes")
            if isinstance(payload, dict) and "discount_rate" in payload:
                rate = normalize_discount_rate(payload["discount_rate"])
                payload["discount_rate"] = None if rate is None else format(rate, "f")
    return result


def json_value(value: Any) -> Any:
    if hasattr(value, "model_dump"):
        return value.model_dump(mode="json")
    return json.loads(json.dumps(value, default=str))

class WorkflowService:
    def __init__(self, config: Any, store: Any, memory: AgentMemory) -> None:
        self.config, self.store, self.memory = config, store, memory
        self.agent_dir = config.agent_dir
        self.input_dir, self.progress_dir, self.data_dir = config.input_dir, config.progress_dir, config.data_dir
        for directory in (self.input_dir, self.progress_dir, self.data_dir,
                          self.agent_dir / "runs", self.agent_dir / "catalogs"):
            directory.mkdir(parents=True, exist_ok=True)
        self.path = self.agent_dir / "workflow.json"
        self._lock = RLock()
        self._job_locks: dict[str, asyncio.Lock] = {}
        self.manager: Any = None
        if self.path.exists():
            self.state = json.loads(self.path.read_text(encoding="utf-8"))
        else:
            self.state = {"uploads": {}, "jobs": {}, "model_id_schema_version": 1}
            self._save()
        self._upgrade_model_fields()
        self._upgrade_discount_policy()
        self._upgrade_pricing_policy()

    def _upgrade_model_fields(self) -> None:
        if self.state.get("model_id_schema_version") == 1:
            return
        for job in self.state.get("jobs", {}).values():
            if job.get("status") != "ready":
                # Historical receipts and in-flight idempotency hashes are immutable.
                continue
            for key in ("bundle", "original_bundle"):
                if isinstance(job.get(key), dict):
                    job[key] = legacy_bundle_fields(job[key])
            for key in ("preview", "commit_result"):
                if isinstance(job.get(key), dict):
                    job[key] = current_model_fields(job[key])
            if job.get("status") == "ready" and isinstance(job.get("bundle"), dict):
                # Refresh in RAM after the data schema changes. No automatic confirmation.
                try:
                    preview = json_value(self.store.preview_bundle(job["bundle"]))
                except (ValueError, RevisionConflictError) as error:
                    job["schema_cleanup_warning"] = "Legacy draft needs manual revision: " + type(error).__name__
                else:
                    job.update(preview=preview, issues=preview.get("issues", []) + job.get("proposal_issues", []))
                job["revision"] = job.get("revision", 1) + 1
                job["updated_at"] = now()
        self.state["model_id_schema_version"] = 1
        self._save()

    def _upgrade_discount_policy(self) -> None:
        """Refresh only unconfirmed reviews once, without changing files or history."""
        changed = False
        for job in self.state.get("jobs", {}).values():
            if job.get("status") != "ready" or job.get("discount_policy_version") == DISCOUNT_POLICY_VERSION:
                continue
            try:
                bundle = normalize_bundle_discounts(job.get("bundle"))
                original = normalize_bundle_discounts(job.get("original_bundle"))
                if not isinstance(bundle, dict):
                    raise ValueError("legacy review has no valid proposal")
                # Normalize both sides before remembering later human edits.
                job.update(bundle=bundle, original_bundle=original if isinstance(original, dict) else deepcopy(bundle))
                preview = json_value(self.store.preview_bundle(bundle))
            except (ValueError, RevisionConflictError) as error:
                # Invalid old proposals remain reviewable, never implicitly committed.
                job.update(preview=None, discount_policy_warning=(
                    "Legacy draft requires manual revision under the two-decimal discount policy: "
                    + type(error).__name__))
            else:
                job.update(preview=preview, issues=preview.get("issues", []) + job.get("proposal_issues", []))
                job.pop("discount_policy_warning", None)
            job.update(discount_policy_version=DISCOUNT_POLICY_VERSION,
                       revision=job.get("revision", 1) + 1, updated_at=now())
            changed = True
        if changed:
            self._save()

    def _upgrade_pricing_policy(self) -> None:
        """Refresh pending reviews without inferring scope or rewriting raw proposals."""
        changed = False
        for job in self.state.get("jobs", {}).values():
            if job.get("status") != "ready" or job.get("pricing_policy_version") == PRICING_POLICY_VERSION:
                continue
            try:
                bundle = job.get("bundle")
                if not isinstance(bundle, dict):
                    raise ValueError("legacy review has no valid proposal")
                validate_lookup_bundle(job, bundle, self.store.catalog_snapshot())
                preview = json_value(self.store.preview_bundle(bundle))
            except (ValueError, RevisionConflictError) as error:
                job.update(preview=None, pricing_policy_warning=(
                    "Legacy draft requires manual revision under the structured pricing policy: "
                    + type(error).__name__))
            else:
                job.update(preview=preview, issues=preview.get("issues", []) + job.get("proposal_issues", []))
                job.pop("pricing_policy_warning", None)
            # Both raw bundle baselines stay byte-for-byte business-equivalent. Defaults
            # projected by the new preview are not confirmed human corrections.
            job.update(pricing_policy_version=PRICING_POLICY_VERSION,
                       revision=job.get("revision", 1) + 1, updated_at=now())
            changed = True
        if changed:
            self._save()

    def _save(self) -> None:
        atomic_json(self.path, self.state)

    def _job(self, identifier: str) -> dict[str, Any]:
        if identifier not in self.state["jobs"]:
            raise KeyError("review was not found")
        return self.state["jobs"][identifier]

    def _location(self, item: dict[str, Any]) -> Path:
        stage = item.get("status", "input")
        if stage == "input":
            return self.input_dir / item["stored_name"]
        base = self.data_dir if stage == "data" else self.progress_dir
        return base / item["job_id"] / item["stored_name"]

    def _file_public(self, item: dict[str, Any]) -> dict[str, Any]:
        return {key: item.get(key) for key in ("id", "filename", "size", "sha256", "status", "job_id")}

    def list_files(self) -> list[dict[str, Any]]:
        with self._lock:
            known = {item["stored_name"] for item in self.state["uploads"].values() if item["status"] == "input"}
            for path in sorted(self.input_dir.iterdir()):
                if path.is_file() and not path.is_symlink() and path.name not in known and not path.name.startswith("."):
                    identifier = uuid4().hex
                    self.state["uploads"][identifier] = {"id": identifier, "filename": path.name,
                        "stored_name": path.name, "size": path.stat().st_size, "sha256": None, "status": "input"}
            self._save()
            return [self._file_public(item) for item in self.state["uploads"].values()]

    def upload(self, files: Sequence[Any]) -> list[dict[str, Any]]:
        if not files or len(files) > self.config.max_files:
            raise ValueError("invalid number of uploads")
        if any(len(item.content) > self.config.max_file_bytes for item in files):
            raise ValueError("an upload exceeds the per-file size limit")
        if sum(len(item.content) for item in files) > self.config.max_total_bytes:
            raise ValueError("uploads exceed the total size limit")
        with self._lock:
            saved: list[dict[str, Any]] = []
            try:
                for file in files:
                    identifier = uuid4().hex
                    name = clean_filename(file.filename)
                    item = {"id": identifier, "filename": name, "stored_name": identifier + "_" + name,
                            "size": len(file.content), "sha256": hashlib.sha256(file.content).hexdigest(), "status": "input"}
                    path = self._location(item)
                    with path.open("xb") as handle:
                        handle.write(file.content)
                    self.state["uploads"][identifier] = item
                    saved.append(item)
                self._save()
            except Exception:
                for item in saved:
                    self.state["uploads"].pop(item["id"], None)
                    self._location(item).unlink(missing_ok=True)
                raise
            return [self._file_public(item) for item in saved]

    def ensure_no_waiting_session(self, conversation_id: str | None) -> None:
        if conversation_id and any(job.get("status") == "waiting_for_input"
                and job.get("conversation_id") == conversation_id for job in self.state["jobs"].values()):
            raise RevisionConflictError("Conversation is waiting for input; respond to its existing import job")

    async def _finish_pending_confirmations(self, supplier_id: str | None,
                                            exclude_identifier: str | None = None) -> None:
        for identifier in list(self.state["jobs"]):
            pending = self._job(identifier)
            if pending["status"] != "committed" or identifier == exclude_identifier:
                continue
            owners = {item["supplier_id"] for item in pending.get("profile_updates", [])}
            if pending.get("supplier_id"):
                owners.add(pending["supplier_id"])
            if supplier_id and owners and supplier_id not in owners and not any(owner.startswith("pending-name:") for owner in owners):
                continue
            async with self._job_locks.setdefault(identifier, asyncio.Lock()):
                with self._lock:
                    pending = self._job(identifier)
                    if pending["status"] != "committed":
                        continue
                    try:
                        self._finish_confirmation(pending)
                    except Exception as error:
                        raise RevisionConflictError("Prior import " + identifier + " is already committed but its confirmation receipt needs recovery; retry its confirmation before this supplier continues (" + type(error).__name__ + ")") from error

    async def start(self, *, upload_ids: Sequence[str], prompt: str,
                    conversation_id: str | None = None, intent: str = "import",
                    lookup_model_id: str | None = None, lookup_billing_item: str | None = None,
                    supplier_id: str | None = None, update_request: dict[str, Any] | None = None,
                    source_scope: str | None = None, persist_profile: bool = True,
                    remember_prompt: bool = False, full_reparse: bool = False) -> dict[str, Any]:
        if intent not in {"import", "official_price_lookup"}:
            raise ValueError("unknown workflow intent")
        await self._finish_pending_confirmations(supplier_id)
        if update_request is not None and (intent != "import" or not upload_ids):
            raise ValueError("supplier updates require uploaded files in an ordinary import session")
        request_supplier = (update_request or {}).get("supplier_id")
        if supplier_id and request_supplier and supplier_id != request_supplier:
            raise ValueError("supplier_id conflicts with update_request supplier")
        supplier_id = supplier_id or request_supplier
        if supplier_id is not None:
            self._confirmed_supplier(supplier_id)
        request_scope = (update_request or {}).get("source_scope")
        if source_scope is not None and request_scope and source_scope.strip() != request_scope:
            raise ValueError("source_scope conflicts with update_request source scope")
        source_scope = (source_scope if source_scope is not None else request_scope) or "default"
        source_scope = str(source_scope).strip()
        if not source_scope or len(source_scope) > 255:
            raise ValueError("source_scope must contain 1 to 255 characters")
        if intent == "import":
            supplier_id, conversation_id = await self.manager.resolve_supplier_session(supplier_id, conversation_id)
        self.ensure_no_waiting_session(conversation_id)
        session_scope = getattr(self.manager, "_sessions", {}).get(conversation_id, {})
        continuation = session_scope.get("intent") == "official_price_lookup"
        if continuation:
            if upload_ids:
                raise ValueError("official lookup conversations accept text clarification only, not attachments")
            # A public ordinary-start request cannot downgrade a network-enabled official session.
            intent = "official_price_lookup"
            lookup_model_id = session_scope["lookup_model_id"]
            lookup_billing_item = session_scope.get("lookup_billing_item")
        if intent == "official_price_lookup" and (
                upload_ids or not lookup_model_id or conversation_id is not None and not continuation):
            raise ValueError("official lookup requires a new independent session or its scoped continuation")
        prompt = prompt.strip()
        if not prompt:
            raise ValueError("prompt must not be blank")
        if len(upload_ids) != len(set(upload_ids)) or len(upload_ids) > self.config.max_files:
            raise ValueError("invalid upload selection")
        if conversation_id is not None and not re.fullmatch(r"ses[A-Za-z0-9_-]+", conversation_id):
            raise ValueError("invalid conversation_id")
        with self._lock:
            items = []
            for identifier in upload_ids:
                item = self.state["uploads"].get(identifier)
                if item is None or item["status"] != "input":
                    raise ValueError("selected upload is not available in input")
                if self._location(item).is_symlink() or not self._location(item).is_file():
                    raise ValueError("selected upload is unavailable")
                if item["size"] > self.config.max_file_bytes:
                    raise ValueError("an upload exceeds the per-file size limit")
                if item.get("sha256") is None:
                    item["sha256"] = hashlib.sha256(self._location(item).read_bytes()).hexdigest()
                items.append(item)
            if sum(item["size"] for item in items) > self.config.max_total_bytes:
                raise ValueError("uploads exceed the total size limit")
            auto_reconciliation = supplier_id is not None and bool(items) and update_request is None and intent == "import"
            if auto_reconciliation:
                update_request = self._default_update_request(supplier_id, source_scope)
            job = self._new_job(prompt, conversation_id, list(upload_ids))
            job.update(auto_reconciliation=auto_reconciliation, intent=intent, lookup_model_id=lookup_model_id, lookup_billing_item=lookup_billing_item,
                       supplier_id=supplier_id, update_request=deepcopy(update_request), source_scope=source_scope,
                       persist_profile=persist_profile, remember_prompt=remember_prompt, full_reparse=full_reparse,
                       profile_context=supplier_profile_context(self.store, supplier_id, source_scope)
                           if intent == "import" else None)
            if update_request is not None:
                job["update_request"]["file_sha256"] = self._upload_fingerprint(items)
                if supplier_id:
                    job["update_request"]["supplier_id"] = supplier_id
            job["status"] = "preparing"
            self._save()
            try:
                destination = self.progress_dir / job["id"]
                destination.mkdir(exist_ok=True)
                for item in items:
                    source = self._location(item)
                    item["job_id"] = job["id"]
                    target = destination / item["stored_name"]
                    source.rename(target)
                    item["status"] = "progress"
                    if item["sha256"] is None:
                        item["sha256"] = hashlib.sha256(target.read_bytes()).hexdigest()
                    self._save()
                job["status"] = "queued"
                self._save()
            except Exception:
                self._return_files(job)
                job.update(status="failed", error="Files could not enter the processing directory")
                self._save()
                raise
        self.manager.schedule(job["id"])
        return self.get(job["id"])

    def _new_job(self, prompt: str, conversation_id: str | None, upload_ids: list[str]) -> dict[str, Any]:
        identifier = uuid4().hex
        job = {"id": identifier, "run_id": identifier, "import_id": identifier,
               "conversation_id": conversation_id, "status": "queued", "revision": 1,
               "discount_policy_version": DISCOUNT_POLICY_VERSION,
               "pricing_policy_version": PRICING_POLICY_VERSION,
               "attempt": 1, "clarifications": [], "questions": [], "question_summary": None,
               "intent": "import", "lookup_model_id": None, "lookup_billing_item": None,
               "source_scope": "default", "persist_profile": True, "remember_prompt": False, "profile_context": None,
               "upload_ids": upload_ids, "prompt": prompt, "bundle": None, "original_bundle": None,
               "preview": None, "issues": [], "proposal_issues": [], "error": None, "assistant_message": None,
               "messages": [], "changes": dict(_ZERO), "created_at": now(), "updated_at": now()}
        self.state["jobs"][identifier] = job
        return job

    def _confirmed_supplier(self, identifier: str) -> dict[str, Any]:
        for supplier in self.store.list_suppliers():
            if supplier.get("supplier_id", supplier.get("id")) == identifier:
                return supplier
        raise ValueError("Choose an existing confirmed supplier or confirm its creation first")

    def _default_update_request(self, supplier_id: str, source_scope: str) -> dict[str, Any]:
        snapshot = self.store.catalog_snapshot()
        existing = any(row.get("supplier_id") == supplier_id and row.get("source_scope") == source_scope
                       for row in snapshot.get("quotes", []))
        receipts = any(row.get("supplier_id") == supplier_id and row.get("source_scope") == source_scope
                       for row in snapshot.get("supplier_update_receipts", []))
        fresh = not existing and not receipts and self.store.get_supplier_baseline(supplier_id, source_scope) is None
        return {"supplier_id": supplier_id, "source_scope": source_scope, "mode": "incremental",
                "latest_confirmed": fresh}

    def accept_source_proposal(self, identifier: str, candidate: dict[str, Any], issues: list[Any] | None = None,
                               reconciliation_request: dict[str, Any] | None = None) -> None:
        job = self._job(identifier)
        bundle = merge_source_bundle(job, candidate)
        observations = source_coverage(job, reconciliation_request)
        self.update_job(identifier, source_read_coverage={
            "complete_read": bool((observations or {}).get("complete_read")),
            "covered_sheets": deepcopy((observations or {}).get("covered_sheets", [])),
            "unresolved_items": (observations or {}).get("unresolved_items", 0)})
        self.accept_proposal(identifier, bundle, issues, observations)
        current = self._job(identifier)
        reused = job.get("_source_reused_bundle") or {}
        reused_count = sum(len(reused.get(key, [])) for key in ("create", "official_prices"))
        interpreted = sum(len(candidate.get(key, [])) for key in ("create", "official_prices"))
        metrics = {**current.get("metrics", {}), "reused_entries": reused_count,
                   "reinterpreted_entries": interpreted, "clarification_count": len(current.get("clarifications", []))}
        plan = deepcopy(current.get("reading_plan"))
        if plan:
            plan.setdefault("stats", {}).update(reused_entries=reused_count, reinterpreted_entries=interpreted)
        self.update_job(identifier, metrics=metrics, reading_plan=plan)

    @staticmethod
    def _upload_fingerprint(items: Sequence[dict[str, Any]]) -> str:
        hashes = sorted(item["sha256"] for item in items if item.get("sha256"))
        if len(hashes) != len(items) or not hashes:
            raise ValueError("supplier updates require verified original file fingerprints")
        return hashes[0] if len(hashes) == 1 else hashlib.sha256(
            json.dumps(hashes, separators=(",", ":")).encode()).hexdigest()

    def _bind_supplier(self, bundle: dict[str, Any], supplier_id: str | None) -> dict[str, Any]:
        result = deepcopy(bundle)
        if not supplier_id:
            return result
        supplier = self._confirmed_supplier(supplier_id)
        def check_scope(row: dict[str, Any]) -> None:
            if row.get("supplier_id") and row["supplier_id"] != supplier_id:
                raise ValueError("quotation belongs to another supplier than this selected supplier session")
            name = " ".join(str(row.get("supplier") or "").split()).casefold()
            allowed = {" ".join(value.split()).casefold()
                       for value in [supplier["name"], *supplier.get("aliases", [])]}
            if name and name not in allowed:
                raise ValueError("quotation supplier name conflicts with this selected supplier session")
        for item in result.get("create", []):
            check_scope(item)
            if not item.get("supplier_id") and not item.get("supplier"):
                item.update(supplier_id=supplier_id, supplier=supplier["name"])
        for item in result.get("update", []):
            changes = item.get("changes", {})
            existing = self.store.get_quote(item["id"]).model_dump(mode="json")
            check_scope(existing)
            check_scope({**existing, **changes})
            if not changes.get("supplier_id") and not changes.get("supplier") and not existing.get("supplier"):
                changes.update(supplier_id=supplier_id, supplier=supplier["name"])
        return result

    def _missing_supplier(self, bundle: dict[str, Any]) -> bool:
        if any(not item.get("supplier_id") and not str(item.get("supplier") or "").strip()
               for item in bundle.get("create", [])):
            return True
        for item in bundle.get("update", []):
            changes = item.get("changes", {})
            existing = self.store.get_quote(item["id"]).model_dump(mode="json")
            if "supplier" in changes or "supplier_id" in changes:
                if not changes.get("supplier_id") and not str(changes.get("supplier") or "").strip():
                    return True
            elif not existing.get("supplier_id") and not existing.get("supplier"):
                return True
        return False

    def _wait_for_supplier(self, job: dict[str, Any], bundle: dict[str, Any],
                           issues: list[Any], reconciliation: dict[str, Any] | None) -> None:
        job.update(status="waiting_for_input", waiting_reason="supplier",
                   questions=["Choose or confirm the supplier for these quotations. The original files and parsed draft are retained."],
                   question_summary="Supplier quotations require a confirmed supplier before review.",
                   pending_bundle=deepcopy(bundle), pending_issues=deepcopy(issues),
                   pending_reconciliation=deepcopy(reconciliation),
                   bundle=None, original_bundle=None, preview=None, issues=[], proposal_issues=[],
                   error=None, changes=dict(_ZERO), revision=job["revision"] + 1, updated_at=now())
        self._save()

    def _sheet_inventory(self, job: dict[str, Any]) -> tuple[set[str], bool]:
        sheets: set[str] = set()
        readable = True
        for key in job["upload_ids"]:
            item = self.state["uploads"][key]
            source = self._location(item)
            suffix = Path(item["filename"]).suffix.casefold()
            try:
                if suffix in {".xlsx", ".xlsm"}:
                    from openpyxl import load_workbook
                    book = load_workbook(source, read_only=True, data_only=False)
                    try:
                        sheets.update(book.sheetnames)
                    finally:
                        book.close()
                elif suffix == ".xls":
                    import xlrd
                    book = xlrd.open_workbook(str(source), on_demand=True)
                    try:
                        sheets.update(book.sheet_names())
                    finally:
                        book.release_resources()
                elif suffix in {".csv", ".tsv"}:
                    sheets.add("CSV")
                else:
                    # The model can read arbitrary business files. An unknown
                    # container cannot establish complete workbook coverage.
                    readable = False
            except Exception:
                readable = False
        return sheets, readable

    def _reconciliation_request(self, job: dict[str, Any], bundle: dict[str, Any],
                                metadata: dict[str, Any] | None = None) -> dict[str, Any]:
        request = deepcopy(job.get("reconciliation_request") or job.get("update_request") or {})
        # Only source interpretation comes from the model. Deletion authorization,
        # scope, supplier selection, rollback permission and mappings remain human-owned.
        allowed = {"entries", "complete_read", "covered_sheets", "unresolved_items",
                   "source_version", "source_effective_date"}
        if metadata:
            observations = {key: deepcopy(value) for key, value in metadata.items() if key in allowed}
            conflicts = []
            for key in ("source_version", "source_effective_date"):
                supplied = (job.get("update_request") or {}).get(key)
                observed = observations.get(key)
                if supplied:
                    if observed and str(observed) != str(supplied):
                        conflicts.append({"field": key, "human": supplied, "source": observed})
                    observations.pop(key, None)
            request.update(observations)
            if conflicts:
                job["source_observation_conflicts"] = conflicts
                request["unresolved_items"] = max(1, request.get("unresolved_items", 0))
        request["entries"] = deepcopy(bundle.get("create", []))
        request["models"] = deepcopy(bundle.get("models", []))
        unchanged_official = set(job.get("_source_unchanged_official_ids", []))
        request["official_prices"] = [deepcopy(row) for row in bundle.get("official_prices", [])
                                      if row.get("official_price_id") not in unchanged_official]
        request["file_sha256"] = self._upload_fingerprint(
            [self.state["uploads"][key] for key in job["upload_ids"]])
        if job.get("supplier_id"):
            request["supplier_id"] = job["supplier_id"]
        if request.get("supplier_id"):
            supplier = self._confirmed_supplier(request["supplier_id"])
            for entry in request["entries"]:
                if entry.get("supplier_id") and entry["supplier_id"] != request["supplier_id"]:
                    raise ValueError("An update cannot change quotations belonging to another supplier")
                if not entry.get("supplier_id") and not entry.get("supplier"):
                    entry.update(supplier_id=request["supplier_id"], supplier=supplier["name"])
        if request.get("mode") == "full":
            inventory, readable = self._sheet_inventory(job)
            expected = set(request.get("expected_sheets") or [])
            covered = set(request.get("covered_sheets") or [])
            job["source_inventory"] = {"sheets": sorted(inventory), "readable": readable}
            if not readable or not expected or not expected.issubset(inventory) or not expected.issubset(covered):
                request["complete_read"] = False
        if bundle.get("update") or bundle.get("delete"):
            raise ValueError("Supplier-update drafts must provide raw entries; use reviewed mappings or explicit_delete_ids")
        return request

    def _review_proposal(self, job: dict[str, Any], bundle: dict[str, Any],
                         metadata: dict[str, Any] | None = None) -> tuple[dict[str, Any], dict[str, Any]]:
        bundle = rebind_source_officials(self, job, bundle)
        if job.get("update_request") is not None:
            request = self._reconciliation_request(job, bundle, metadata)
            preview = json_value(self.store.preview_reconciliation(request))
            job["reconciliation_request"] = preview.get("request", request)
            job["proposal_bundle"] = deepcopy(bundle)
            return preview.get("bundle", {}), preview
        validate_lookup_bundle(job, bundle, self.store.catalog_snapshot())
        if job.get("upload_ids") and job.get("intent") != "official_price_lookup":
            job["proposal_bundle"] = deepcopy(bundle)
            unchanged = set(job.get("_source_unchanged_official_ids", []))
            bundle = {**bundle, "official_prices": [row for row in bundle.get("official_prices", [])
                       if row.get("official_price_id") not in unchanged]}
        return bundle, json_value(self.store.preview_bundle(bundle))

    def create_draft(self, bundle: dict[str, Any], prompt: str = "") -> dict[str, Any]:
        bundle = normalize_bundle_discounts(bundle)
        PricingBundle.model_validate(bundle)
        with self._lock:
            job = self._new_job(prompt or "Human proposal", None, [])
            self.accept_proposal(job["id"], bundle)
            return self.get(job["id"])

    def get(self, identifier: str, *, include_source_comparison: bool = True) -> dict[str, Any]:
        with self._lock:
            raw = self._job(identifier)
            hidden = {"original_bundle", "profile_updates", "baseline_update"}
            job = deepcopy({key: value for key, value in raw.items()
                            if key not in hidden and not key.startswith("_source_")
                            and (include_source_comparison or key != "source_comparison")})
            if not include_source_comparison and raw.get("source_comparison"):
                job["source_comparison_available"] = True
            job.pop("original_bundle", None)
            job.pop("profile_updates", None)
            job.pop("baseline_update", None)
            for key in list(job):
                if key.startswith("_source_"):
                    job.pop(key)
            if job.get("source_comparison"):
                job["source_comparison"] = public_source_comparison(job["source_comparison"])
            for key, default in (("attempt", 1), ("clarifications", []), ("questions", []), ("question_summary", None)):
                job.setdefault(key, default)
            job["files"] = [self._file_public(self.state["uploads"][key]) for key in job["upload_ids"]]
            job["review_required"] = job["status"] == "ready"
            return job

    def list_jobs(self) -> list[dict[str, Any]]:
        jobs = [self.get(key, include_source_comparison=False) for key in reversed(list(self.state["jobs"]))]
        for job in jobs:
            if job.pop("source_comparison", None) is not None:
                job["source_comparison_available"] = True
        return jobs

    def update_job(self, identifier: str, **values: Any) -> None:
        with self._lock:
            job = self._job(identifier)
            if job["status"] == "cancelled":
                return
            job.update(values, updated_at=now())
            self._save()

    def accept_proposal(self, identifier: str, bundle: dict[str, Any], issues: list[Any] | None = None,
                        reconciliation_request: dict[str, Any] | None = None) -> None:
        bundle = normalize_bundle_discounts(bundle)
        with self._lock:
            current = self._job(identifier)
            if current["status"] == "cancelled":
                return
            job = deepcopy(current)
            batch_supplier = job.get("supplier_id") or (
                job.get("reconciliation_request") or job.get("update_request") or {}).get("supplier_id")
            if job.get("upload_ids") and job.get("intent") != "official_price_lookup":
                for row in bundle.get("create", []):
                    if row.get("source_scope") and row["source_scope"] != job.get("source_scope", "default"):
                        raise ValueError("quotation source_scope conflicts with the selected quotation series")
                    row["source_scope"] = job.get("source_scope") or "default"
            bundle = self._bind_supplier(bundle, batch_supplier)
            proposal_issues = json_value(issues or [])[:2000]
            # A missing supplier is a resumable business omission; invalid
            # schema or an illegal official-lookup scope is still a real error.
            PricingBundle.model_validate(bundle)
            validate_lookup_bundle(job, bundle, self.store.catalog_snapshot())
            missing_supplier = self._missing_supplier(bundle)
            if job.get("update_request") is not None:
                missing_supplier = missing_supplier and not (
                    (job.get("reconciliation_request") or job.get("update_request") or {}).get("supplier_id"))
                missing_supplier = missing_supplier or not (
                    job.get("supplier_id") or (job.get("reconciliation_request") or job.get("update_request") or {}).get("supplier_id"))
            if missing_supplier:
                self.state["jobs"][identifier] = job
                self._wait_for_supplier(job, bundle, proposal_issues, reconciliation_request)
                return
            applied_bundle, preview = self._review_proposal(job, bundle, reconciliation_request)
            for key in ("waiting_reason", "pending_bundle", "pending_issues", "pending_reconciliation"):
                job.pop(key, None)
            job.update(bundle=deepcopy(applied_bundle), original_bundle=deepcopy(bundle), proposal_issues=proposal_issues,
                       questions=[], question_summary=None, error=None,
                       preview=preview, issues=preview.get("issues", []) + proposal_issues, revision=job["revision"] + 1,
                       status="ready", updated_at=now())
            self.state["jobs"][identifier] = job
            self._save()

    def accept_questions(self, identifier: str, payload: dict[str, Any]) -> None:
        request = NeedsInput.model_validate(payload)
        with self._lock:
            job = self._job(identifier)
            if job["status"] == "cancelled":
                return
            if job["status"] != "running":
                raise RevisionConflictError("Question response status conflict")
            job.update(status="waiting_for_input", questions=list(request.questions),
                       question_summary=request.summary, bundle=None, original_bundle=None, preview=None,
                       issues=[], proposal_issues=[], error=None, changes=dict(_ZERO),
                       revision=job["revision"] + 1, updated_at=now())
            self._save()

    def _merge_update_context(self, job: dict[str, Any], changes: dict[str, Any] | None = None,
                              supplier_id: str | None = None) -> None:
        if job.get("update_request") is None:
            if changes is not None:
                raise ValueError("this job is not a supplier update")
            if supplier_id is not None:
                self._confirmed_supplier(supplier_id)
                job["supplier_id"] = supplier_id
                job["profile_context"] = supplier_profile_context(self.store, supplier_id, job.get("source_scope", "default"))
                if job.get("upload_ids"):
                    job["update_request"] = self._default_update_request(supplier_id, job.get("source_scope") or "default")
                    job["auto_reconciliation"] = True
            return
        old = deepcopy(job.get("reconciliation_request") or job["update_request"])
        supplied = deepcopy(changes or {})
        if supplier_id is not None:
            supplied["supplier_id"] = supplier_id
        if supplied.get("supplier_id"):
            self._confirmed_supplier(supplied["supplier_id"])
            job["supplier_id"] = supplied["supplier_id"]
        request = {**old, **supplied}
        boundary = ("supplier_id", "source_scope", "mode", "expected_sheets", "model_ids", "supplier_model_ids")
        changed = any(key in supplied and request.get(key) != old.get(key) for key in boundary)
        if changed:
            request["scope_confirmed"] = supplied.get("scope_confirmed") is True
            request["complete_read"] = supplied.get("complete_read") is True
            request["covered_sheets"] = deepcopy(supplied.get("covered_sheets", []))
            request["decisions"] = deepcopy(supplied.get("decisions", []))
            request["explicit_delete_ids"] = deepcopy(supplied.get("explicit_delete_ids", []))
            # Retained parsed observations cannot certify an expanded boundary.
            if isinstance(job.get("pending_reconciliation"), dict):
                job["pending_reconciliation"].update(complete_read=False, covered_sheets=[])
        job["source_scope"] = request.get("source_scope") or job.get("source_scope", "default")
        job["update_request"] = deepcopy(request)
        if job.get("reconciliation_request") is not None:
            job["reconciliation_request"] = deepcopy(request)

    async def respond(self, identifier: str, expected_revision: int, prompt: str,
                      supplier_id: str | None = None, update_request: dict[str, Any] | None = None,
                      persist_profile: bool | None = None, remember_prompt: bool | None = None) -> dict[str, Any]:
        prompt = prompt.strip()
        if not prompt or len(prompt) > 16000:
            raise ValueError("clarification response must contain 1 to 16000 characters")
        await self._finish_pending_confirmations(supplier_id or self._job(identifier).get("supplier_id"), identifier)
        async with self._job_locks.setdefault(identifier, asyncio.Lock()):
            with self._lock:
                current = self._job(identifier)
                if current["status"] != "waiting_for_input" or current["revision"] != expected_revision:
                    raise RevisionConflictError("Clarification status or revision conflict")
                job = deepcopy(current)
                self._merge_update_context(job, update_request, supplier_id)
                if persist_profile is not None:
                    job["persist_profile"] = persist_profile
                if remember_prompt is not None:
                    job["remember_prompt"] = remember_prompt
                if job.get("waiting_reason") == "supplier" and supplier_id and job.get("pending_bundle") is not None:
                    pending = self._bind_supplier(job["pending_bundle"], supplier_id)
                    clarification = {"attempt": job.get("attempt", 1), "prompt": prompt,
                                     "supplier_id": supplier_id, "created_at": now(),
                                     "persist_profile": job.get("persist_profile", True),
                                     "questions": deepcopy(job.get("questions", []))}
                    job["clarifications"] = [*job.get("clarifications", []), clarification]
                    self.state["jobs"][identifier] = job
                    self.accept_proposal(identifier, pending, job.get("pending_issues"),
                                         job.get("pending_reconciliation"))
                    return self.get(identifier)
                next_attempt = job.get("attempt", 1) + 1
                clarification = {"attempt": next_attempt, "prompt": prompt, "created_at": now(),
                                 "question_attempt": job.get("attempt", 1),
                                 "persist_profile": job.get("persist_profile", True),
                                 "questions": deepcopy(job.get("questions", [])),
                                 "question_summary": job.get("question_summary"),
                                 "assistant_message": job.get("assistant_message")}
                job.update(status="queued", attempt=next_attempt,
                           clarifications=[*job.get("clarifications", []), clarification],
                           questions=[], question_summary=None, error=None, assistant_message=None, messages=[],
                           bundle=None, original_bundle=None, preview=None, issues=[], proposal_issues=[],
                           changes=dict(_ZERO), revision=job["revision"] + 1, updated_at=now())
                self.state["jobs"][identifier] = job
                self._save()
            self.manager.schedule(identifier)
            return self.get(identifier)

    async def revise(self, identifier: str, expected_revision: int, bundle: dict[str, Any] | None = None,
                     reconciliation_request: dict[str, Any] | None = None,
                     decisions: list[dict[str, Any]] | None = None,
                     persist_profile: bool | None = None, remember_prompt: bool | None = None) -> dict[str, Any]:
        async with self._job_locks.setdefault(identifier, asyncio.Lock()):
            with self._lock:
                current = self._job(identifier)
                if current["status"] != "ready" or current["revision"] != expected_revision:
                    raise ValueError("review status or revision conflict")
                job = deepcopy(current)
                if persist_profile is not None:
                    job["persist_profile"] = persist_profile
                if remember_prompt is not None:
                    job["remember_prompt"] = remember_prompt
                if reconciliation_request is not None or decisions is not None:
                    changes = deepcopy(reconciliation_request or {})
                    if decisions is not None:
                        changes["decisions"] = deepcopy(decisions)
                    self._merge_update_context(job, changes)
                candidate = normalize_bundle_discounts(bundle if bundle is not None else
                    job.get("proposal_bundle") or job["bundle"])
                batch_supplier = job.get("supplier_id") or (
                    job.get("reconciliation_request") or job.get("update_request") or {}).get("supplier_id")
                candidate = self._bind_supplier(candidate, batch_supplier)
                PricingBundle.model_validate(candidate)
                validate_lookup_bundle(job, candidate, self.store.catalog_snapshot())
                if self._missing_supplier(candidate):
                    self.state["jobs"][identifier] = job
                    self._wait_for_supplier(job, candidate, job.get("proposal_issues", []), None)
                    return self.get(identifier)
                applied_bundle, preview = self._review_proposal(job, candidate)
                job.pop("discount_policy_warning", None)
                job.pop("pricing_policy_warning", None)
                job.update(bundle=deepcopy(applied_bundle), preview=preview,
                           issues=preview.get("issues", []) + job.get("proposal_issues", []),
                           revision=job["revision"] + 1, updated_at=now(), error=None)
                self.state["jobs"][identifier] = job
                self._save()
            return self.get(identifier)

    def _move_files(self, job: dict[str, Any], stage: str) -> None:
        base = self.data_dir if stage == "data" else self.progress_dir
        destination = base / job["id"]
        destination.mkdir(exist_ok=True)
        for identifier in job["upload_ids"]:
            item = self.state["uploads"][identifier]
            source = self._location(item)
            target = destination / item["stored_name"]
            # A crash between rename and manifest save can leave the destination already present.
            actual = source if source.is_file() else target
            if stage == "data" and item.get("sha256") and hashlib.sha256(actual.read_bytes()).hexdigest() != item["sha256"]:
                raise OSError("original upload hash changed; confirmation was refused")
            if source != target and source.exists():
                if target.exists():
                    raise OSError("archive destination already exists")
                source.rename(target)
            elif not target.is_file():
                raise OSError("source file is missing")
            item.update(status=stage, job_id=job["id"])
            self._save()

    def _return_files(self, job: dict[str, Any]) -> None:
        for identifier in job["upload_ids"]:
            item = self.state["uploads"][identifier]
            target = self.input_dir / item["stored_name"]
            candidates = [self._location(item), self.progress_dir / job["id"] / item["stored_name"],
                          self.data_dir / job["id"] / item["stored_name"]]
            source = next((path for path in candidates if path.is_file()), None)
            if source is not None and source != target:
                if target.exists():
                    raise OSError("input destination already exists")
                source.rename(target)
            if not target.is_file():
                raise OSError("original upload could not be recovered")
            item.update(status="input")
            item.pop("job_id", None)
            self._save()

    async def confirm(self, identifier: str, expected_revision: int,
                      idempotency_key: str | None = None, *, recovering: bool = False) -> dict[str, Any]:
        await self._finish_pending_confirmations(self._job(identifier).get("supplier_id"), identifier)
        async with self._job_locks.setdefault(identifier, asyncio.Lock()):
            with self._lock:
                job = self._job(identifier)
                if job["status"] == "confirmed":
                    return self.get(identifier)
                if job["status"] == "committed":
                    self._finish_confirmation(job)
                    return self.get(identifier)
                if job["status"] == "confirming":
                    recovering = True
                if not recovering and (job["status"] != "ready" or job["revision"] != expected_revision):
                    raise ValueError("review status or revision conflict")
                if recovering and job["status"] != "confirming":
                    raise ValueError("review is not awaiting confirmation recovery")
                if job.get("discount_policy_warning") or job.get("pricing_policy_warning"):
                    raise ValueError("legacy draft requires manual revision before confirmation")
                # Reconcile from source entries again, never refresh an obsolete
                # update/delete operation list by replacing only its revision token.
                reconciliation = job.get("reconciliation_request")
                if reconciliation is None:
                    validate_lookup_bundle(job, job["bundle"], self.store.catalog_snapshot())
                if not recovering:
                    if reconciliation is not None and job.get("proposal_bundle") and job.get("upload_ids"):
                        source = rebind_source_officials(self, job, job["proposal_bundle"])
                        job["proposal_bundle"] = source
                        reconciliation = self._reconciliation_request(job, source)
                        job["reconciliation_request"] = deepcopy(reconciliation)
                    latest = json_value(self.store.preview_reconciliation(reconciliation) if reconciliation is not None
                                        else self.store.preview_bundle(job["bundle"]))
                    previous = job.get("preview") or {}
                    stale = latest.get("base_catalog_token") != previous.get("base_catalog_token")
                    if reconciliation is not None:
                        stale = stale or latest.get("review_digest") != previous.get("review_digest")
                    if stale:
                        if reconciliation is not None:
                            job["bundle"] = deepcopy(latest.get("bundle", {}))
                        job.update(preview=latest, issues=latest.get("issues", []) + job.get("proposal_issues", []),
                                   revision=job["revision"] + 1, updated_at=now(),
                                   error="Catalog changed; reconciliation was rebuilt. Review its groups and confirm again")
                        self._save()
                        raise RevisionConflictError(job["error"])
                    job["preview"] = latest
                    if reconciliation is not None and not latest.get("ready"):
                        raise ValueError("Resolve the pending reconciliation choices before confirmation")
                    if any(issue.get("code") == "pricing_basis_required"
                           for issue in latest.get("issues", []) if isinstance(issue, dict)):
                        raise ValueError("Choose amount or discount_rate as the valid pricing basis before confirmation")
                if not recovering:
                    filenames = [self.state["uploads"][key]["filename"] for key in job["upload_ids"]]
                    profile_updates = supplier_confirmation_updates(self.store, job, filenames)
                    candidate = job.get("proposal_bundle") or job.get("bundle") or {}
                    allowed_models = [row["model_id"] for group in ("models", "create")
                                      for row in candidate.get(group, []) if row.get("model_id")]
                    profile_previews = {}
                    for update in profile_updates:
                        if not str(update["supplier_id"]).startswith("pending-name:"):
                            before = self.store.get_supplier_profile(update["supplier_id"], update["source_scope"])
                            after = self.store.preview_confirmed_supplier_profile(**update, allowed_model_ids=allowed_models)
                            profile_previews[update["supplier_id"]] = {"before": before, "after": after}
                    job["profile_updates"] = deepcopy(profile_updates)
                    freeze_source_baseline(self, job, profile_previews)
                job.update(status="confirming", confirmation_key=idempotency_key or identifier, updated_at=now())
                self._save()
                try:
                    self._move_files(job, "data")
                    if reconciliation is not None:
                        result = json_value(self.store.apply_reconciliation(reconciliation, actor="review:" + identifier,
                            import_id="review:" + identifier,
                            expected_catalog_token=job["preview"].get("base_catalog_token"),
                            expected_review_digest=job["preview"].get("review_digest")))
                    else:
                        result = json_value(self.store.apply_bundle(job["bundle"], actor="review:" + identifier,
                            import_id="review:" + identifier,
                            expected_catalog_token=job["preview"].get("base_catalog_token")))
                except Exception as error:
                    if isinstance(error, RevisionConflictError):
                        self._move_files(job, "progress")
                        try:
                            job["preview"] = json_value(self.store.preview_reconciliation(reconciliation)
                                if reconciliation is not None else self.store.preview_bundle(job["bundle"]))
                            if reconciliation is not None:
                                job["bundle"] = deepcopy(job["preview"].get("bundle", {}))
                            job["issues"] = job["preview"].get("issues", []) + job.get("proposal_issues", [])
                        except (ValueError, RevisionConflictError):
                            pass
                        job.update(status="ready", revision=job["revision"] + 1,
                                   error="Catalog changed during confirmation; review and confirm again")
                        self._save()
                    elif not recovering:
                        try:
                            self._move_files(job, "progress")
                        except OSError:
                            job.update(status="confirming", error="Confirmation requires file recovery; retry after restart")
                            self._save()
                            raise
                        job.update(status="ready", error="Confirmation failed; no unconfirmed proposal was accepted")
                        self._save()
                    raise
                # Persist the committed receipt before post-commit memory/file work.
                job.update(commit_result=result, status="committed", error=None)
                self._save()
                self._finish_confirmation(job)
            return self.get(identifier)

    def _compact_terminal_source_state(self, job: dict[str, Any]) -> None:
        for key in list(job):
            if key.startswith("_source_") or key in {"profile_updates", "baseline_update"}:
                job.pop(key, None)
        if job["status"] == "confirmed":
            if job.get("source_comparison"):
                job["source_comparison"] = public_source_comparison(job["source_comparison"])
            sid, scope = job.get("supplier_id"), job.get("source_scope", "default")
            if sid:
                for previous in self.state["jobs"].values():
                    if (previous["id"] != job["id"] and previous.get("status") == "confirmed"
                            and previous.get("supplier_id") == sid and previous.get("source_scope", "default") == scope
                            and previous.pop("source_comparison", None) is not None):
                        previous["source_comparison_archived"] = True
        else:
            job.pop("source_comparison", None)

    def _finish_confirmation(self, job: dict[str, Any]) -> None:
        result = job["commit_result"]
        filenames = [self.state["uploads"][key]["filename"] for key in job["upload_ids"]]
        remember_supplier_confirmation(self.store, job, filenames)
        self.manager.adopt_confirmed_supplier(job)
        record_source_baseline(self, job)
        changes = dict(result.get("changes", {}))
        changes.setdefault("total", sum(int(changes.get(key, 0)) for key in ("added", "updated", "deleted")))
        job.update(status="confirmed", changes=changes, preview=result, error=None,
                   revision=job["revision"] + 1, updated_at=now())
        self._compact_terminal_source_state(job)
        self._save()
        if job.get("supplier_id") and job.get("conversation_id") in getattr(self.manager, "_sessions", {}):
            self.manager._sessions[job["conversation_id"]]["summary"] = self.manager._supplier_summary(job["supplier_id"])
            atomic_json(self.manager._sessions_path, self.manager._sessions)
        self._remove_draft(job["id"])

    async def cancel(self, identifier: str) -> dict[str, Any]:
        async with self._job_locks.setdefault(identifier, asyncio.Lock()):
            job = self._job(identifier)
            if job["status"] == "cancelled":
                return self.get(identifier)
            if job["status"] in {"confirmed", "confirming", "committed"}:
                raise ValueError("a confirmed or committing review cannot be cancelled")
            await self.manager.abort(identifier)
            with self._lock:
                # A provider may finish despite cancellation. Proposal acceptance
                # swaps a validated job copy, so refetch the current durable state.
                job = self._job(identifier)
                self._return_files(job)
                self._remove_draft(identifier)
                for key in ("pending_bundle", "pending_issues", "pending_reconciliation", "waiting_reason"):
                    job.pop(key, None)
                job.update(status="cancelled", bundle=None, original_bundle=None, preview=None,
                           issues=[], questions=[], question_summary=None, error=None, changes=dict(_ZERO), revision=job["revision"] + 1, updated_at=now())
                self._compact_terminal_source_state(job)
                self._save()
            return self.get(identifier)

    def _remove_draft(self, identifier: str) -> None:
        directory = self.agent_dir / "runs" / identifier
        if directory.parent == self.agent_dir / "runs" and directory.is_dir():
            shutil.rmtree(directory)
        catalogs = self.agent_dir / "catalogs"
        for pattern in (identifier + ".json", identifier + ".official-sources.md",
                        identifier + ".attempt-*.json", identifier + ".attempt-*.official-sources.md"):
            for path in catalogs.glob(pattern):
                path.unlink(missing_ok=True)

    async def recover(self) -> None:
        for identifier in list(self.state["jobs"]):
            job = self._job(identifier)
            try:
                if job["status"] == "confirming":
                    await self.confirm(identifier, job["revision"], recovering=True)
                elif job["status"] == "committed":
                    self._finish_confirmation(job)
                elif job["status"] in {"preparing", "queued", "running", "cancelling"}:
                    await self.manager.abort(identifier)
                    self._return_files(job)
                    self._remove_draft(identifier)
                    job.update(status="failed", error="Processing was interrupted; originals were returned to input",
                               bundle=None, preview=None, changes=dict(_ZERO))
                    self._save()
            except Exception as error:
                job["error"] = "Recovery pending: " + type(error).__name__
                self._save()
