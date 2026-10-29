"""Same-origin CRUD, normalized catalog, deterministic previews and exports."""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
import json
from typing import Literal
from urllib.parse import quote as urlquote

from fastapi import APIRouter, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import ValidationError
from .schemas import (
    BulkQuoteRequest, BulkQuoteDeleteRequest, BulkQuoteDeleteResponse, BulkWriteResponse, Contract, FilterOptions, PricingBundle,
    QuoteListResponse, QuotePatch, QuoteResponse, PricingContextUpdate, ViewExportRequest,
    SupplierWrite, SupplierReconciliationRequest, SupplierProfileWrite,
)
from .store import IdempotencyConflictError, QuoteNotFoundError, RevisionConflictError, WorkspaceStore
from .units import pricing_units_payload

EXPORT_COLUMNS = (
    "id", "model_id", "model_name", "vendor", "billing_item", "supplier_id", "supplier",
    "source_scope", "source_entry_id", "source_sheet",
    "supplier_model_id", "tags", "provider_type", "amount", "pricing_unit_id",
    "discount_rate", "pricing_basis", "raw_discount_rate", "derived_amount",
    "converted_supplier_amount", "converted_supplier_pricing_unit_id", "conversion_status",
    "conversion_evidence", "discount_status", "discount_evidence", "normalized", "display_issues",
    "currency", "unit", "unit_quantity", "legacy_unit", "conditions", "billing_metric", "pricing_conditions",
    "discount_scope", "discount_items", "official_prices", "official_match", "effective_discount_rules", "dimension", "direction", "effective_date",
    "source", "source_ref", "conversion", "note", "pricing_issue", "status", "revision",
    "created_at", "updated_at",
)


class CatalogWriteRequest(Contract):
    bundle: PricingBundle
    actor: str = "workspace-ui"
    import_id: str | None = None
    expected_catalog_token: str | None = None


class ReconciliationApplyRequest(Contract):
    request: SupplierReconciliationRequest
    expected_catalog_token: str
    expected_review_digest: str
    actor: str = "workspace-ui"
    import_id: str | None = None


def _detail(error: Exception) -> str:
    return " ".join(str(error).split())[:1000] or "request could not be processed"


def _error(error: Exception) -> HTTPException:
    status = 404 if isinstance(error, QuoteNotFoundError) else (
        409 if isinstance(error, (RevisionConflictError, IdempotencyConflictError)) else 422)
    return HTTPException(status_code=status, detail=_detail(error))


def _ids(value: str | None) -> list[int] | None:
    if value is None:
        return None
    try:
        result = [int(part.strip()) for part in value.split(",") if part.strip()]
    except ValueError:
        raise HTTPException(status_code=422, detail="ids must be comma-separated integers") from None
    if not result or any(v <= 0 for v in result):
        raise HTTPException(status_code=422, detail="ids must contain positive integers")
    return list(dict.fromkeys(result))


def _export_values(item: QuoteResponse) -> list[object]:
    value = item.model_dump(mode="json")
    result = []
    for column in EXPORT_COLUMNS:
        cell = value.get(column)
        if isinstance(cell, (dict, list)):
            cell = json.dumps(cell, ensure_ascii=False, sort_keys=True)
        if isinstance(cell, str) and cell.startswith(("=", "+", "-", "@", "\t", "\r")):
            cell = "'" + cell
        result.append("" if cell is None else cell)
    return result


def _view_cell(value: object) -> object:
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _view_filename(filename: str, format: str) -> str:
    suffix = "." + format
    return filename if filename.casefold().endswith(suffix) else filename + suffix


def _csv_response(items: list[QuoteResponse]) -> Response:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(EXPORT_COLUMNS)
    writer.writerows(_export_values(item) for item in items)
    return Response(content=output.getvalue().encode("utf-8-sig"), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="tokensql-quotes.csv"'})


def _xlsx_response(items: list[QuoteResponse]) -> StreamingResponse:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Quotes")
    sheet.append(EXPORT_COLUMNS)
    for item in items:
        sheet.append(_export_values(item))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="tokensql-quotes.xlsx"'})


def create_data_router(store: WorkspaceStore) -> APIRouter:
    router = APIRouter(prefix="/api/workspace", tags=["workspace-data"])

    @router.get("/pricing-units")
    def pricing_units():
        return pricing_units_payload()

    @router.get("/pricing-context")
    def pricing_context():
        return store.pricing_context()

    @router.get("/pricing-context/history")
    def pricing_context_history():
        return {"items": store.pricing_context_history()}

    @router.put("/pricing-context")
    def update_pricing_context(request: PricingContextUpdate):
        try:
            return store.set_pricing_context(**request.model_dump(mode="python"))
        except (ValueError, RevisionConflictError) as error:
            raise _error(error) from None

    @router.get("/catalog")
    def catalog():
        return store.catalog_snapshot()

    @router.get("/suppliers")
    def suppliers():
        return {"items": store.list_suppliers()}

    @router.post("/suppliers")
    def confirm_supplier(request: SupplierWrite):
        try:
            return store.confirm_supplier(request)
        except (ValueError, RevisionConflictError) as error:
            raise _error(error) from None


    @router.get("/suppliers/{supplier_id}/profile")
    def supplier_profile(supplier_id: str, source_scope: str | None = None):
        try:
            return store.get_supplier_profile(supplier_id, source_scope)
        except ValueError as error:
            raise _error(error) from None

    @router.get("/suppliers/{supplier_id}/profile/history")
    def supplier_profile_history(supplier_id: str):
        try:
            return {"items": store.supplier_profile_history(supplier_id)}
        except ValueError as error:
            raise _error(error) from None

    @router.put("/suppliers/{supplier_id}/profile")
    def update_supplier_profile(supplier_id: str, request: SupplierProfileWrite):
        try:
            return store.update_supplier_profile(supplier_id, request)
        except (ValueError, RevisionConflictError) as error:
            raise _error(error) from None

    @router.post("/supplier-updates/preview")
    def preview_supplier_update(request: SupplierReconciliationRequest):
        try:
            return store.preview_reconciliation(request)
        except (ValueError, RevisionConflictError, QuoteNotFoundError) as error:
            raise _error(error) from None

    @router.post("/supplier-updates/apply")
    def apply_supplier_update(request: ReconciliationApplyRequest):
        try:
            return store.apply_reconciliation(request.request, request.actor, request.import_id,
                expected_catalog_token=request.expected_catalog_token,
                expected_review_digest=request.expected_review_digest)
        except (ValueError, RevisionConflictError, IdempotencyConflictError, QuoteNotFoundError) as error:
            raise _error(error) from None

    @router.post("/catalog/preview")
    def preview_catalog(bundle: PricingBundle):
        try:
            return store.preview_bundle(bundle)
        except (ValueError, RevisionConflictError, QuoteNotFoundError) as error:
            raise _error(error) from None

    @router.post("/catalog")
    def write_catalog(request: CatalogWriteRequest):
        try:
            return store.apply_bundle(request.bundle, request.actor, request.import_id,
                                      expected_catalog_token=request.expected_catalog_token)
        except (ValueError, RevisionConflictError, IdempotencyConflictError, QuoteNotFoundError) as error:
            raise _error(error) from None

    @router.post("/quotes/preview")
    def preview_quotes(request: BulkQuoteRequest):
        try:
            return store.preview_bundle({"create": request.items})
        except (ValueError, RevisionConflictError, QuoteNotFoundError) as error:
            raise _error(error) from None

    @router.get("/quotes", response_model=QuoteListResponse)
    def list_quotes(
        q: str | None = None, model_id: str | None = None, model_name: str | None = None,
        manufacturer: str | None = None, vendor: str | None = None, provider: str | None = None,
        supplier: str | None = None, currency: str | None = None,
        pricing_unit_id: str | None = None, metric: str | None = None,
        billing_item: str | None = None, dimension: str | None = None, tag: str | None = None,
        supplier_model_id: str | None = None,
        limit: int = Query(default=100, ge=1, le=500), offset: int = Query(default=0, ge=0),
    ) -> QuoteListResponse:
        items, total = store.list_quotes(
            q=q, model_id=model_id, model_name=model_name, manufacturer=manufacturer, vendor=vendor,
            provider=provider, supplier=supplier, currency=currency, pricing_unit_id=pricing_unit_id,
            metric=metric,
            billing_item=billing_item, dimension=dimension, tag=tag,
            supplier_model_id=supplier_model_id, limit=limit, offset=offset,
        )
        return QuoteListResponse(items=items, total=total,
                                 filter_options=FilterOptions.model_validate(store.filter_options()))

    @router.post("/quotes/bulk-delete", response_model=BulkQuoteDeleteResponse)
    def delete_quotes(request: BulkQuoteDeleteRequest) -> BulkQuoteDeleteResponse:
        try:
            result = store.delete_many(request.items, actor=request.actor)
            return BulkQuoteDeleteResponse.model_validate(result)
        except (ValueError, QuoteNotFoundError, RevisionConflictError) as error:
            raise _error(error) from None

    @router.get("/quotes/{quote_id}", response_model=QuoteResponse)
    def get_quote(quote_id: int) -> QuoteResponse:
        try:
            return store.get_quote(quote_id)
        except QuoteNotFoundError as error:
            raise _error(error) from None

    @router.post("/quotes", response_model=BulkWriteResponse, status_code=201)
    def create_quotes(request: BulkQuoteRequest) -> BulkWriteResponse:
        try:
            items, reused = store.create_many(request.items, actor=request.actor,
                                               idempotency_key=request.idempotency_key)
        except (ValueError, RevisionConflictError, IdempotencyConflictError) as error:
            raise _error(error) from None
        return BulkWriteResponse(items=items, created=0 if reused else len(items), reused=reused)

    @router.patch("/quotes/{quote_id}", response_model=QuoteResponse)
    def update_quote(quote_id: int, request: QuotePatch) -> QuoteResponse:
        changes = request.model_dump(mode="python", exclude_unset=True, exclude={"expected_revision", "actor"})
        try:
            return store.update_quote(quote_id, changes, expected_revision=request.expected_revision,
                                       actor=request.actor)
        except (ValueError, ValidationError, QuoteNotFoundError, RevisionConflictError) as error:
            raise _error(error) from None

    @router.delete("/quotes/{quote_id}", response_model=QuoteResponse)
    def delete_quote(quote_id: int, expected_revision: int = Query(ge=1),
                     actor: str = "workspace-ui") -> QuoteResponse:
        try:
            return store.delete_quote(quote_id, expected_revision=expected_revision, actor=actor)
        except (QuoteNotFoundError, RevisionConflictError) as error:
            raise _error(error) from None

    @router.post("/exports/view")
    def export_view(request: ViewExportRequest) -> Response:
        """Serialize a client-computed display view without persisting or recalculating it."""
        headings = [_view_cell(column.label) for column in request.columns]
        rows = [[_view_cell(row.get(column.key)) for column in request.columns]
                for row in request.rows]
        filename = _view_filename(request.filename, request.format)
        fallback = "".join(character if character.isascii() and (
            character.isalnum() or character in "._-") else "_" for character in filename)
        disposition = {"Content-Disposition":
            f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{urlquote(filename)}"}
        if request.format == "csv":
            output = StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow(headings)
            writer.writerows(rows)
            return Response(content=output.getvalue().encode("utf-8-sig"),
                            media_type="text/csv; charset=utf-8", headers=disposition)
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("View")
        sheet.append(headings)
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return StreamingResponse(output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=disposition)

    @router.get("/exports")
    def export_quotes(
        format: Literal["csv", "xlsx"] = "csv", ids: str | None = None, q: str | None = None,
        model_id: str | None = None, model_name: str | None = None, manufacturer: str | None = None,
        vendor: str | None = None, provider: str | None = None, supplier: str | None = None,
        currency: str | None = None, pricing_unit_id: str | None = None,
        metric: str | None = None, billing_item: str | None = None,
        dimension: str | None = None, tag: str | None = None, supplier_model_id: str | None = None,
    ) -> Response:
        items, _ = store.list_quotes(
            q=q, model_id=model_id, model_name=model_name, manufacturer=manufacturer, vendor=vendor,
            provider=provider, supplier=supplier, currency=currency, pricing_unit_id=pricing_unit_id,
            metric=metric, billing_item=billing_item,
            dimension=dimension, tag=tag, supplier_model_id=supplier_model_id,
            ids=_ids(ids), limit=100000, offset=0,
        )
        return _csv_response(items) if format == "csv" else _xlsx_response(items)

    return router
