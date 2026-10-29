"""Deterministic source-table snapshots, content alignment and conservative reuse.

Cell coordinates locate evidence only. They never map a position to a quote or
supply business meaning. Ambiguous context is re-read by the model.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
import csv
from datetime import date, datetime, time
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
from zipfile import BadZipFile
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
import xlrd

PARSER_VERSION = "v24-content-1"
MAX_CELLS = 500000


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def fingerprint(value: Any) -> str:
    return hashlib.sha256(_json(value).encode("utf-8")).hexdigest()


def prompt_fingerprint(prompt: str) -> str:
    return fingerprint(str(prompt).strip())


def _value(value: Any, cell_type: str | None = None) -> tuple[str, Any]:
    if value is None:
        return "empty", None
    if cell_type == "f":
        return "formula", str(value)
    if isinstance(value, bool):
        return "boolean", value
    if isinstance(value, (datetime, date, time)):
        return "date", value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        return "number", str(Decimal(str(value)))
    if cell_type == "e":
        return "error", str(value)
    return "text", str(value)



def _format_kind(value: str) -> str:
    value = value or "General"
    if is_date_format(value):
        return "date:" + value
    active, literal, index = [], [], 0
    while index < len(value):
        char = value[index]
        if char == "\\" and index + 1 < len(value):
            literal.append(value[index + 1])
            index += 2
            continue
        if char == '"':
            index += 1
            while index < len(value) and value[index] != '"':
                literal.append(value[index])
                index += 1
        else:
            active.append(char)
        index += 1
    pattern = "".join(active)
    percent = "%" in pattern
    symbols = sorted(set(re.findall(r"USD|HKD|CNY|RMB|[$\u00a5\uffe5\u20ac\u00a3]", value, re.I)))
    kind = "percent" if percent else "currency:" + ",".join(symbols) if symbols else "number"
    if not percent and "%" in literal:
        kind += ":literal_percent"
    # Trailing commas scale displayed numbers; grouping commas and decimal
    # precision do not change the stored value's interpretation.
    scale = re.search(r"[0#?](,+)[^0#?]*$", pattern.split(";")[0])
    return kind + (":scale:" + str(len(scale.group(1))) if scale else "")

def _cell_meaning(cell: dict) -> dict:
    return {key: cell.get(key) for key in
            ("column", "type", "value", "format_kind", "formula", "cached_value", "comment")}


def _row_key(row: dict) -> str:
    return fingerprint([_cell_meaning(cell) for cell in row["cells"]])



def _numeric_like(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    return bool(re.fullmatch(
        r"(?:(?:USD|HKD|CNY|RMB|[$\u00a5\uffe5\u20ac\u00a3])\s*)?[+-]?"
        r"(?:\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d*)?|\.\d+)"
        r"(?:[eE][+-]?\d+)?\s*(?:%|\u6298|\u5143|USD|HKD|CNY|RMB)?", value.strip(), re.I))


def _numeric_cell(cell: dict) -> bool:
    return cell["type"] in {"number", "formula", "date"} or _numeric_like(cell.get("value"))

def _text_key(row: dict) -> str | None:
    values = [(cell["column"], cell["value"]) for cell in row["cells"]
              if cell["type"] in {"text", "boolean"} and not _numeric_like(cell.get("value")) and cell["value"] not in (None, "")]
    return fingerprint(values) if values else None


def _color(value: Any) -> Any:
    return None if value is None else {"type": value.type, "value": value.value, "tint": value.tint}


def _style(cell: Any) -> dict:
    return {"font": {"name": cell.font.name, "size": str(cell.font.sz), "bold": cell.font.b,
                     "italic": cell.font.i, "color": _color(cell.font.color)},
            "fill": {"type": cell.fill.fill_type, "color": _color(cell.fill.fgColor)},
            "alignment": {"horizontal": cell.alignment.horizontal, "vertical": cell.alignment.vertical,
                          "wrap": cell.alignment.wrap_text},
            "number_format": cell.number_format}


def _finish_sheet(sheet: dict) -> dict:
    for row in sheet["rows"]:
        row["semantic_key"] = _row_key(row)
        row["text_key"] = _text_key(row)
    sheet["semantic_fingerprint"] = fingerprint({
        "rows": [row["semantic_key"] for row in sheet["rows"]],
        "merged_ranges": sheet["merged_ranges"], "tables": sheet.get("tables", []),
        "defined_names": sheet.get("defined_names", [])})
    sheet["presentation_fingerprint"] = fingerprint({
        "rows": [{"row": row["row"], "height": row.get("height"), "hidden": row.get("hidden"),
                  "cells": [cell.get("style") for cell in row["cells"]]} for row in sheet["rows"]],
        "columns": sheet.get("columns", [])})
    return sheet


def _xlsx(path: Path, file_key: str) -> tuple[list[dict], dict]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    result, count = [], 0

    metadata = {"defined_names": [{"name": item.name, "reference": item.attr_text,
                                   "local_sheet_id": item.localSheetId}
                                  for item in workbook.defined_names.values()],
                "calculation": {key: getattr(workbook.calculation, key, None) for key in
                    ("calcMode", "fullCalcOnLoad", "forceFullCalc", "iterate", "iterateCount", "iterateDelta")}}
    metadata["defined_names"].sort(key=lambda item: (item["name"], str(item["local_sheet_id"])))
    try:
        for index, source in enumerate(workbook.worksheets):
            if source.max_row * source.max_column > MAX_CELLS * 4:
                raise ValueError("workbook dimensions exceed the safe snapshot budget")
            sheet = {"sheet_key": file_key + "/" + str(index), "name": source.title,
                     "rows": [], "merged_ranges": sorted(str(value) for value in source.merged_cells.ranges),
                     "columns": [{"column": key, "width": dim.width, "hidden": dim.hidden}
                                 for key, dim in source.column_dimensions.items()],
                     "formula_metadata_available": True,
                     "defined_names": [{"name": item.name, "reference": item.attr_text}
                                       for item in source.defined_names.values()],
                     "tables": [{"name": table.name, "ref": table.ref,
                                 "columns": [column.name for column in table.tableColumns]}
                                for table in source.tables.values()]}
            for cells in source.iter_rows():
                row = {"row": cells[0].row, "cells": [],
                       "height": source.row_dimensions[cells[0].row].height,
                       "hidden": source.row_dimensions[cells[0].row].hidden}
                for cell in cells:
                    if cell.value is None and cell.comment is None:
                        continue
                    count += 1
                    if count > MAX_CELLS:
                        raise ValueError("workbook exceeds the safe snapshot cell budget")
                    kind, value = _value(cell.value, cell.data_type)
                    row["cells"].append({"column": cell.column, "address": cell.coordinate,
                        "type": kind, "value": value, "number_format": cell.number_format,
                        "format_kind": _format_kind(cell.number_format),
                        "formula": value if kind == "formula" else None,
                        "cached_value": _value(cached[source.title][cell.coordinate].value)[1]
                            if kind == "formula" else None,
                        "comment": cell.comment.text if cell.comment else None,
                        "style": _style(cell)})
                if row["cells"]:
                    sheet["rows"].append(row)
            result.append(_finish_sheet(sheet))
    finally:
        workbook.close()
        cached.close()
    return result, metadata



def _xls_has_formulas(path: Path) -> bool:
    """Conservatively detect formula-bearing BIFF records, never evaluate them."""
    from struct import unpack_from
    from xlrd import biffh
    from xlrd.compdoc import CompDoc, SIGNATURE
    data = path.read_bytes()
    try:
        if data[:8] == SIGNATURE:
            document = CompDoc(data)
            data = document.get_named_stream("Workbook") or document.get_named_stream("Book")
            if data is None:
                return True
        formula_records = {*biffh.XL_FORMULA_OPCODES, biffh.XL_ARRAY, biffh.XL_ARRAY2,
                           biffh.XL_SHRFMLA, biffh.XL_TABLEOP, biffh.XL_TABLEOP2, biffh.XL_TABLEOP_B2}
        offset = 0
        while offset + 4 <= len(data):
            code, size = unpack_from("<HH", data, offset)
            if code in formula_records:
                return True
            offset += 4 + size
            if offset > len(data):
                return True
        return any(data[offset:])
    except (OSError, ValueError, TypeError, IndexError):
        return True

def _xls(path: Path, file_key: str) -> list[dict]:
    workbook = xlrd.open_workbook(str(path), formatting_info=True)
    formulas_present = _xls_has_formulas(path)
    result, count = [], 0
    try:
        for index, source in enumerate(workbook.sheets()):
            rows = []
            for row_index in range(source.nrows):
                row = {"row": row_index + 1, "cells": [],
                       "height": getattr(source.rowinfo_map.get(row_index), "height", None),
                       "hidden": getattr(source.rowinfo_map.get(row_index), "hidden", None)}
                for col_index in range(source.ncols):
                    cell = source.cell(row_index, col_index)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    count += 1
                    if count > MAX_CELLS:
                        raise ValueError("workbook exceeds the safe snapshot cell budget")
                    raw = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        raw = xlrd.xldate_as_datetime(raw, workbook.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        raw = bool(raw)
                    kind, value = _value(raw, "e" if cell.ctype == xlrd.XL_CELL_ERROR else None)
                    xf = workbook.xf_list[cell.xf_index]
                    format_value = workbook.format_map[xf.format_key].format_str
                    row["cells"].append({"column": col_index + 1,
                        "address": get_column_letter(col_index + 1) + str(row_index + 1),
                        "type": kind, "value": value, "number_format": format_value,
                        "format_kind": _format_kind(format_value), "formula": None, "cached_value": None,
                        "comment": getattr(source.cell_note_map.get((row_index, col_index)), "text", None),
                        "style": {"font": {key: getattr(workbook.font_list[xf.font_index], key, None)
                                          for key in ("name", "height", "bold", "italic", "colour_index")},
                                  "format": format_value}})
                if row["cells"]:
                    rows.append(row)
            merges = [get_column_letter(c1 + 1) + str(r1 + 1) + ":" + get_column_letter(c2) + str(r2)
                      for r1, r2, c1, c2 in source.merged_cells]
            result.append(_finish_sheet({"sheet_key": file_key + "/" + str(index), "name": source.name,
                "rows": rows, "merged_ranges": sorted(merges), "columns": [],
                "formula_metadata_available": not formulas_present}))
    finally:
        workbook.release_resources()
    return result


def _csv(path: Path, file_key: str) -> list[dict]:
    rows = []
    with path.open(encoding="utf-8-sig", newline="") as source:
        for index, values in enumerate(csv.reader(source, delimiter="\t" if path.suffix.casefold() == ".tsv" else ","), 1):
            cells = [{"column": col, "address": get_column_letter(col) + str(index),
                      "type": "text", "value": value, "number_format": "General", "format_kind": "number",
                      "formula": None, "cached_value": None, "comment": None, "style": {}, "numeric_literal": bool(re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?%?", value.strip()))}
                     for col, value in enumerate(values, 1) if value != ""]
            if cells:
                rows.append({"row": index, "cells": cells, "height": None, "hidden": False})
    if sum(len(row["cells"]) for row in rows) > MAX_CELLS:
        raise ValueError("table exceeds the safe snapshot cell budget")
    return [_finish_sheet({"sheet_key": file_key + "/0", "name": "CSV",
        "rows": rows, "merged_ranges": [], "columns": [], "formula_metadata_available": True})]


def read_workbooks(files: list[dict | str | Path]) -> dict:
    """Read supplied originals; never mutate files, workbooks or application data."""
    result = {"version": 1, "parser_version": PARSER_VERSION, "files": [], "supported": True, "issues": []}
    seen = Counter()
    for item in files:
        item = {"path": item} if isinstance(item, (str, Path)) else item
        path = Path(item["path"])
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if item.get("sha256") and item["sha256"] != digest:
            raise ValueError("source workbook hash changed before snapshot")
        seen[digest] += 1
        file_key = digest[:24] + "-" + str(seen[digest])
        entry = {"file_key": file_key, "filename": item.get("filename") or path.name,
                 "sha256": digest, "sheets": [], "supported": True, "issues": []}
        try:
            suffix = path.suffix.casefold()
            if suffix in {".xlsx", ".xlsm"}:
                entry["sheets"], entry["workbook_metadata"] = _xlsx(path, file_key)
            elif suffix == ".xls":
                entry["sheets"] = _xls(path, file_key)
            elif suffix in {".csv", ".tsv"}:
                entry["sheets"] = _csv(path, file_key)
            else:
                raise ValueError("source type requires full model reading")
        except (OSError, ValueError, KeyError, TypeError, EOFError, BadZipFile, InvalidFileException, xlrd.XLRDError) as error:
            entry["supported"] = False
            entry["issues"] = [" ".join(str(error).split())[:500]]
            result["supported"] = False
            result["issues"].extend(entry["issues"])
        entry["semantic_fingerprint"] = fingerprint({
            "sheets": [{"name": sheet["name"], "meaning": sheet["semantic_fingerprint"]} for sheet in entry["sheets"]],
            "workbook_metadata": entry.get("workbook_metadata", {})})
        entry["structure_fingerprint"] = fingerprint(sorted(sheet["name"] for sheet in entry["sheets"]))
        result["files"].append(entry)
    result["semantic_fingerprint"] = fingerprint(sorted(
        entry["semantic_fingerprint"] for entry in result["files"]))
    return result


def _pair(items_old: list[dict], items_new: list[dict], keys: tuple[str, ...],
          single_fallback: bool = True) -> list[tuple[dict | None, dict | None]]:
    old, new = list(items_old), list(items_new)
    pairs = []
    for key in keys:
        left, right = defaultdict(list), defaultdict(list)
        for item in old:
            left[item[key]].append(item)
        for item in new:
            right[item[key]].append(item)
        for value in left.keys() & right.keys():
            if len(left[value]) == len(right[value]) == 1:
                one, two = left[value][0], right[value][0]
                pairs.append((one, two))
                old.remove(one)
                new.remove(two)
    if single_fallback and len(old) == len(new) == 1:
        pairs.append((old.pop(), new.pop()))
    return pairs + [(item, None) for item in old] + [(None, item) for item in new]


def _location(file: dict, sheet: dict, row: dict) -> dict:
    return {"file_key": file["file_key"], "filename": file["filename"],
            "sheet_key": sheet["sheet_key"], "sheet": sheet["name"], "row": row["row"]}


def _context_rows(sheet: dict) -> set[int]:
    result = set()
    if sheet["rows"]:
        result.add(sheet["rows"][0]["row"])
    for row in sheet["rows"]:
        cells = row["cells"]
        numeric = any(_numeric_cell(cell) for cell in cells)
        if not numeric or any(cell.get("comment") or
                              (cell["type"] == "text" and len(str(cell["value"])) > 100) for cell in cells):
            result.add(row["row"])
    for merge in sheet["merged_ranges"]:
        _, first, _, last = range_boundaries(merge)
        if result & set(range(first, last + 1)):
            result.update(range(first, last + 1))
    return result


def _group_rows(sheet: dict, affected: set[int]) -> set[int]:
    """Expand vertical merged groups and adjacent continuation rows."""
    result = set(affected)
    for merge in sheet["merged_ranges"]:
        _, first, _, last = range_boundaries(merge)
        if result & set(range(first, last + 1)):
            result.update(range(first, last + 1))
    rows = sheet["rows"]
    leading = min((cell["column"] for row in rows for cell in row["cells"]), default=1)
    anchor = None
    groups = defaultdict(set)
    for row in rows:
        first = next((cell for cell in row["cells"] if cell["column"] == leading), None)
        if first and first["value"] not in (None, ""):
            anchor = row["row"]
        if anchor is not None:
            groups[anchor].add(row["row"])
    for members in groups.values():
        if result & members:
            result.update(members)
    return result


def _formula_dependencies(sheet: dict) -> list[tuple[int, list[tuple[str | None, int, int]], bool]]:
    result = []
    for row in sheet["rows"]:
        for cell in row["cells"]:
            formula = cell.get("formula")
            if not formula:
                continue
            deps, unknown = [], bool(re.search(r"\b(?:INDIRECT|OFFSET)\s*\(", formula, re.I))
            try:
                for token in Tokenizer(formula).items:
                    if token.type != "OPERAND" or token.subtype != "RANGE":
                        continue
                    match = re.fullmatch(r"(?:(?:'([^']+)'|([^!]+))!)?(\$?[A-Z]+\$?\d+)(?::(\$?[A-Z]+\$?\d+))?", token.value)
                    if not match:
                        unknown = True
                        continue
                    target = match.group(1) or match.group(2)
                    first = int(re.search(r"\d+", match.group(3))[0])
                    last = int(re.search(r"\d+", match.group(4) or match.group(3))[0])
                    deps.append((target, min(first, last), max(first, last)))
            except (ValueError, TypeError):
                unknown = True
            result.append((row["row"], deps, unknown))
    return result


def compare_snapshots(old: dict, new: dict) -> dict:
    groups = {key: [] for key in ("added", "modified", "deleted", "moved", "presentation")}
    result = {"version": 1, "groups": groups, "row_matches": [], "affected_old_rows": {},
              "affected_new_rows": {}, "regions": [], "requires_full": False,
              "issues": [], "stats": {}, "old_snapshot": old, "new_snapshot": new}
    records = []
    for old_file, new_file in _pair(old["files"], new["files"],
                                    ("semantic_fingerprint", "structure_fingerprint")):
        if (old_file and not old_file["supported"]) or (new_file and not new_file["supported"]):
            result["requires_full"] = True
            result["issues"].append("Unsupported workbook metadata requires full reading")
        old_sheets = old_file["sheets"] if old_file else []
        new_sheets = new_file["sheets"] if new_file else []
        for old_sheet, new_sheet in _pair(old_sheets, new_sheets, ("name", "semantic_fingerprint")):
            before = old_sheet["rows"] if old_sheet else []
            after = new_sheet["rows"] if new_sheet else []
            old_affected, new_affected = set(), set()
            old_tainted, new_tainted = set(), set()
            row_pairs = _pair(before, after, ("semantic_key", "text_key"), single_fallback=False)
            # Identical positional duplicates are safe only if the entire sheet is identical.
            if old_sheet and new_sheet and old_sheet["semantic_fingerprint"] == new_sheet["semantic_fingerprint"]:
                row_pairs = list(zip(before, after))
            for left, right in row_pairs:
                old_loc = _location(old_file, old_sheet, left) if left else None
                new_loc = _location(new_file, new_sheet, right) if right else None
                same = bool(left and right and left["semantic_key"] == right["semantic_key"])
                item = {"old": old_loc, "new": new_loc, "before": left, "after": right}
                if left and right:
                    result["row_matches"].append({"old": old_loc, "new": new_loc, "unchanged": same})
                if same:
                    if old_loc["row"] != new_loc["row"] or old_loc["sheet"] != new_loc["sheet"]:
                        groups["moved"].append({**item, "reason": "content_position"})
                        old_tainted.add(left["row"])
                        new_tainted.add(right["row"])
                    if fingerprint({"height": left.get("height"), "hidden": left.get("hidden"),
                                    "cells": [cell["style"] for cell in left["cells"]]}) != fingerprint({
                                        "height": right.get("height"), "hidden": right.get("hidden"),
                                        "cells": [cell["style"] for cell in right["cells"]]}):
                        groups["presentation"].append({**item, "reason": "display_format"})
                else:
                    kind = "modified" if left and right else "deleted" if left else "added"
                    groups[kind].append({**item, "reason": "cell_meaning"})
                    if left:
                        old_affected.add(left["row"])
                    if right:
                        new_affected.add(right["row"])
            reason = "changed_content"
            if old_sheet and new_sheet:
                old_context, new_context = _context_rows(old_sheet), _context_rows(new_sheet)
                context_change = (bool(old_affected & old_context or new_affected & new_context)
                    or old_sheet["merged_ranges"] != new_sheet["merged_ranges"]
                    or old_sheet["name"] != new_sheet["name"])
                if any(item["old"] and item["new"] and item["old"]["sheet_key"] == old_sheet["sheet_key"]
                       and item["old"]["row"] in old_context for item in groups["moved"]):
                    context_change = True
                # BIFF does not expose formula text. Only byte-identical sources prove unchanged formulas.
                uncertain = (not old_sheet.get("formula_metadata_available", True)
                             or not new_sheet.get("formula_metadata_available", True)) and old_file["sha256"] != new_file["sha256"]
                context_change = context_change or old_sheet.get("tables", []) != new_sheet.get("tables", [])
                if context_change or uncertain:
                    old_affected = {row["row"] for row in before}
                    new_affected = {row["row"] for row in after}
                    reason = "legacy_formula_metadata" if uncertain else "header_unit_note_or_structure"
            if old_file and new_file and (old_file.get("workbook_metadata", {}) != new_file.get("workbook_metadata", {})
                    or (old_sheet or {}).get("defined_names", []) != (new_sheet or {}).get("defined_names", [])):
                for sheet, target in ((old_sheet, old_affected), (new_sheet, new_affected)):
                    if sheet:
                        target.update(row["row"] for row in sheet["rows"]
                                      if any(cell.get("formula") for cell in row["cells"]))
                reason = "workbook_formula_metadata"
            if old_sheet:
                old_affected = _group_rows(old_sheet, old_affected)
            if new_sheet:
                new_affected = _group_rows(new_sheet, new_affected)
            records.append({"old_file": old_file, "new_file": new_file, "old_sheet": old_sheet,
                            "new_sheet": new_sheet, "old_affected": old_affected,
                            "new_affected": new_affected, "old_tainted": old_tainted,
                            "new_tainted": new_tainted, "reason": reason})
    # Workbook notes and shared unit/discount explanations can govern other
    # sheets without an Excel formula. Unknown business dependencies are never
    # treated as evidence of independence.
    for side in ("old", "new"):
        context_files = {record[side + "_file"]["file_key"] for record in records
                         if record[side + "_file"] and record["reason"] in {"header_unit_note_or_structure", "workbook_header_unit_note_or_structure"}
                         and record[side + "_affected"]}
        for record in records:
            if record[side + "_file"] and record[side + "_file"]["file_key"] in context_files:
                record[side + "_affected"].update(row["row"] for row in record[side + "_sheet"]["rows"])
                record["reason"] = "workbook_header_unit_note_or_structure"
    # Formula dependencies may cross sheets. Resolve to a fixed point; unknown references
    # conservatively include the formula sheet whenever its workbook meaning changed.
    for side in ("old", "new"):
        changed = True
        while changed:
            changed = False
            by_file = defaultdict(list)
            for record in records:
                if record[side + "_sheet"]:
                    by_file[record[side + "_file"]["file_key"]].append(record)
            for related in by_file.values():
                any_change = any(record[side + "_affected"] or record[side + "_tainted"] for record in related)
                for record in related:
                    sheet = record[side + "_sheet"]
                    target = record[side + "_affected"]
                    for row, dependencies, unknown in _formula_dependencies(sheet):
                        affected = unknown and any_change
                        for name, first, last in dependencies:
                            affected = affected or any(
                                (name is None and candidate is record or
                                 name == candidate[side + "_sheet"]["name"])
                                and any(first <= value <= last for value in candidate[side + "_affected"] | candidate[side + "_tainted"])
                                for candidate in related)
                        if affected and row not in target:
                            target.add(row)
                            record[side + "_tainted"].add(row)
                            record["reason"] = "formula_dependency"
                            changed = True
                    expanded = _group_rows(sheet, target)
                    if expanded != target:
                        target.update(expanded)
                        changed = True
    for record in records:
        region = {"reason": record["reason"]}
        useful = False
        for side in ("old", "new"):
            sheet, file = record[side + "_sheet"], record[side + "_file"]
            affected = record[side + "_affected"]
            region[side + "_sheet_key"] = sheet["sheet_key"] if sheet else None
            region[side + "_file_key"] = file["file_key"] if file else None
            region[side + "_sheet"] = sheet["name"] if sheet else None
            region[side + "_filename"] = file["filename"] if file else None
            region[side + "_workbook_metadata"] = deepcopy(file.get("workbook_metadata", {})) if file else {}
            region[side + "_tables"] = deepcopy(sheet.get("tables", [])) if sheet else []
            region[side + "_defined_names"] = deepcopy(sheet.get("defined_names", [])) if sheet else []
            region[side + "_rows"] = []
            if sheet and affected:
                useful = True
                result["affected_" + side + "_rows"][sheet["sheet_key"]] = sorted(affected)
                context = _context_rows(sheet)
                context.update(value + offset for value in affected for offset in (-1, 1))
                region[side + "_rows"] = [{**deepcopy(row), "parse": row["row"] in affected}
                    for row in sheet["rows"] if row["row"] in affected | context]
        if useful:
            result["regions"].append(region)
    result["semantic_changed"] = bool(result["regions"] or result["requires_full"])
    result["stats"] = {key + "_rows": len(items) for key, items in groups.items()}
    result["stats"].update(old_rows=sum(len(s["rows"]) for f in old["files"] for s in f["sheets"]),
                           new_rows=sum(len(s["rows"]) for f in new["files"] for s in f["sheets"]),
                           changed_context_characters=len(_json(_compact_regions(result["regions"]))))
    return result



def _reference_parts(reference: str) -> tuple[str, str] | None:
    """Separate a canonical location from a bounded quoted-business annotation."""
    if "(" not in reference:
        return reference, ""
    primary, tail = reference.split("(", 1)
    tail = "(" + tail
    if not primary.strip() or len(tail) > 4000 or not tail.endswith(")"):
        return None
    depth = 0
    for char in tail:
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
            if depth < 0:
                return None
    return (primary.strip(), tail[1:-1].strip()) if depth == 0 else None

def _entry_refs(entry: dict, snapshot: dict) -> list[tuple[dict, dict, set[int]]] | None:
    reference = str(entry.get("source_ref") or "").strip()
    if not reference:
        return None
    parts = _reference_parts(reference)
    if parts is None:
        return None
    reference, annotation = parts
    sheets = [(file, sheet) for file in snapshot["files"] for sheet in file["sheets"]]
    source = str(entry.get("source") or "")
    named = [(file, sheet) for file, sheet in sheets
             if file["filename"] == source or source.endswith("/" + file["filename"])]
    if named:
        sheets = named
    if entry.get("source_sheet"):
        sheets = [(file, sheet) for file, sheet in sheets if sheet["name"] == entry["source_sheet"]]
    # Full cell/range references, with optional quoted sheet prefix.
    match = re.fullmatch(r"(?:(?:'([^']+)'|([^!]+))!)?(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?(?:\s*[,;]\s*\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)*)", reference)
    coordinate_match = match
    rows = set()
    if match:
        sheet_name = match.group(1) or match.group(2)
        if sheet_name:
            sheets = [(file, sheet) for file, sheet in sheets if sheet["name"] == sheet_name]
        for cell_range in re.split(r"\s*[,;]\s*", match.group(3)):
            _, first, _, last = range_boundaries(cell_range.replace("$", ""))
            if last - first > 10000:
                return None
            rows.update(range(first, last + 1))
    else:
        match = re.fullmatch(r"rows?\s*(\d+)(?:\s*[-:]\s*(\d+))?", reference, re.I)
        if not match:
            return None
        first, last = int(match.group(1)), int(match.group(2) or match.group(1))
        if last < first or last - first > 10000:
            return None
        rows.update(range(first, last + 1))

    if annotation:
        # Secondary locations must be covered by the primary coordinate range.
        # Never infer extra dependencies from arbitrary price/model prose.
        if "!" in annotation or not coordinate_match:
            return None
        primary_ranges = [range_boundaries(part.replace("$", ""))
                          for part in re.split(r"\s*[,;]\s*", coordinate_match.group(3))]
        tokens = re.findall(r"(?<![A-Za-z0-9_])\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?(?![A-Za-z0-9_])", annotation)
        for token in tokens:
            left, first, right, last = range_boundaries(token.replace("$", ""))
            if not any(left >= col1 and right <= col2 and first >= row1 and last <= row2
                       for col1, row1, col2, row2 in primary_ranges):
                return None
    return [(sheets[0][0], sheets[0][1], rows)] if len(sheets) == 1 else None


def _relocate_entry(entry: dict, snapshot: dict, comparison: dict) -> tuple[dict, bool]:
    refs = _entry_refs(entry, snapshot)
    if not refs:
        return deepcopy(entry), False
    matches = {(item["old"]["sheet_key"], item["old"]["row"]): item["new"]
               for item in comparison["row_matches"] if item["unchanged"]}
    targets = [matches.get((sheet["sheet_key"], row)) for _, sheet, rows in refs for row in rows]
    if not targets or any(target is None for target in targets):
        return deepcopy(entry), False
    locations = {(item["filename"], item["sheet"]) for item in targets}
    if len(locations) != 1:
        return deepcopy(entry), False
    result = deepcopy(entry)
    primary_reference, annotation = _reference_parts(str(entry["source_ref"]))
    file_name, sheet_name = next(iter(locations))
    old_rows = sorted(row for _, _, rows in refs for row in rows)
    mapping = {row: matches[(refs[0][1]["sheet_key"], row)]["row"] for row in old_rows}

    # A source range denotes every intermediate row, not just its endpoints.
    for first, last in re.findall(r"\$?[A-Z]{1,3}\$?(\d+):\$?[A-Z]{1,3}\$?(\d+)", primary_reference):
        original = list(range(int(first), int(last) + 1))
        values = [mapping[row] for row in original]
        if values != list(range(values[0], values[0] + len(values))):
            return deepcopy(entry), False
    if re.fullmatch(r"rows?\s*\d+(?:\s*[-:]\s*\d+)?", primary_reference, re.I):
        values = sorted(set(mapping.values()))
        # Noncontiguous source rows cannot be represented by one simple row range.
        if values != list(range(values[0], values[-1] + 1)):
            return deepcopy(entry), False
        result["source_ref"] = "rows " + str(values[0]) + (("-" + str(values[-1])) if len(values) > 1 else "")
    else:
        reference = primary_reference
        cells = reference.split("!", 1)[-1]
        cells = re.sub(r"(\$?[A-Z]{1,3}\$?)(\d+)",
                       lambda match: match.group(1) + str(mapping[int(match.group(2))]), cells)
        result["source_ref"] = "'" + sheet_name.replace("'", "''") + "'!" + cells

    if annotation:
        def relocate_annotation(match: Any) -> str:
            # Only explicit cell references are rewritten. Quoted business
            # literals and model names that resemble coordinates stay literal.
            prefix, suffix = annotation[:match.start()], annotation[match.end():]
            if prefix.count("'") % 2 or prefix.count('"') % 2:
                return match.group(0)
            explicit = bool(re.search(r"(?:\bcells?|\u5355\u5143\u683c)\s*$", prefix, re.I)
                            or suffix.lstrip().startswith("="))
            if not explicit:
                return match.group(0)
            return re.sub(r"(\$?[A-Z]{1,3}\$?)(\d+)",
                          lambda value: value.group(1) + str(mapping[int(value.group(2))]), match.group(0))
        annotation = re.sub(r"(?<![A-Za-z0-9_])\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?(?![A-Za-z0-9_])",
                            relocate_annotation, annotation)
        result["source_ref"] += " (" + annotation + ")"
    if "source_sheet" in entry:
        result["source_sheet"] = sheet_name
    if entry.get("source"):
        result["source"] = file_name
    return result, True



def _compact_regions(regions: list[dict]) -> list[dict]:
    result = []
    fields = ("column", "address", "type", "value", "number_format", "format_kind",
              "formula", "cached_value", "comment")
    for region in regions:
        item = {key: deepcopy(value) for key, value in region.items()
                if key not in {"old_rows", "new_rows"}}
        for side in ("old", "new"):
            item[side + "_rows"] = [{"row": row["row"], "parse": row["parse"],
                "cells": [{key: cell.get(key) for key in fields if cell.get(key) is not None}
                          for cell in row["cells"]]} for row in region[side + "_rows"]]
        result.append(item)
    return result

def build_reuse_plan(baseline: dict | None, new_snapshot: dict, comparison: dict | None = None, *,
                     profile_fingerprint: str, prompt_fingerprint: str,
                     force_full: bool = False) -> dict:
    empty = {"models": [], "official_prices": [], "create": [], "update": [], "delete": []}
    result = {"reused_bundle": deepcopy(empty), "needs_model": True, "reuse_mode": "none",
              "changed_context": None, "coverage": {"complete_read": False, "covered_sheets": []},
              "stats": {"reused_entries": 0, "reinterpreted_entries": 0, "model_calls": 0},
              "reason": "first_confirmed_source" if baseline is None else "interpretation_changed"}
    if not baseline:
        return result
    comparison = comparison or compare_snapshots(baseline["snapshot"], new_snapshot)
    result["comparison"] = comparison
    compatible = (not force_full and baseline.get("parser_version") == PARSER_VERSION
                  and baseline.get("profile_fingerprint") == profile_fingerprint
                  and baseline.get("prompt_fingerprint") == prompt_fingerprint)
    bundle = baseline.get("bundle") or empty
    if not compatible or comparison["requires_full"]:
        result["reason"] = "forced_full" if force_full else "interpretation_or_reader_changed"
        result["stats"]["reinterpreted_entries"] = sum(len(bundle.get(key, [])) for key in ("official_prices", "create"))
        return result
    covered = sorted({sheet["name"] for file in new_snapshot["files"] for sheet in file["sheets"]})
    if not comparison["semantic_changed"]:
        reused = deepcopy(bundle)
        reused["models"] = []
        for key in ("official_prices", "create"):
            reused[key] = []
            for entry in bundle.get(key, []):
                relocated, safe = _relocate_entry(entry, baseline["snapshot"], comparison)
                if not safe and _entry_refs(entry, baseline["snapshot"]):
                    result["reason"] = "source_region_requires_full_read"
                    return result
                reused[key].append(relocated)
        result.update(reused_bundle=reused, needs_model=False, reuse_mode="all",
                      reason="confirmed_content_unchanged",
                      coverage={"complete_read": True, "covered_sheets": covered})
    else:
        missing_refs = False
        for key in ("official_prices", "create"):
            for entry in bundle.get(key, []):
                refs = _entry_refs(entry, baseline["snapshot"])
                if not refs:
                    missing_refs = True
                    break
                changed = any(rows & set(comparison["affected_old_rows"].get(sheet["sheet_key"], []))
                              for _, sheet, rows in refs)
                if not changed:
                    relocated, safe = _relocate_entry(entry, baseline["snapshot"], comparison)
                    if not safe:
                        missing_refs = True
                        break
                    result["reused_bundle"][key].append(relocated)
            if missing_refs:
                break
        if missing_refs:
            result["reused_bundle"] = deepcopy(empty)
            result["reason"] = "source_provenance_requires_full_read"
        else:
            result.update(reuse_mode="partial", reason="changed_regions",
                changed_context={"version": 1, "parser_version": PARSER_VERSION,
                                 "regions": _compact_regions(comparison["regions"]),
                                 "instruction": "Interpret new rows marked parse=true. Other rows and old rows are context only. Do not repeat unchanged entries outside those regions."},
                coverage={"complete_read": True, "covered_sheets": covered})
            # Deletions without new affected source rows need no model interpretation.
            if not any(row.get("parse") for region in comparison["regions"] for row in region["new_rows"]):
                result["needs_model"] = False
    result["stats"]["reused_entries"] = sum(len(result["reused_bundle"].get(key, []))
                                              for key in ("official_prices", "create"))
    result["stats"]["reinterpreted_entries"] = max(0, sum(len(bundle.get(key, []))
        for key in ("official_prices", "create")) - result["stats"]["reused_entries"])
    result["stats"]["changed_context_characters"] = len(_json(result["changed_context"])) if result["changed_context"] else 0
    return result
