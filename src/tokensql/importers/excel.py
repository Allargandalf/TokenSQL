"""Pure Excel profiling and strict quote-string classification.

Formulae are never executed. The workbook is opened twice through openpyxl:
once for formula text and once for the cached result stored in the file.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import hashlib
import posixpath
from pathlib import Path
import re
from typing import Any, Dict, Iterable, Mapping, Optional, Tuple
import unicodedata
from zipfile import ZipFile

from .types import (
    CellRecord,
    ColumnHeader,
    ExcelParseResult,
    FormulaKind,
    ParsedPrice,
    ParserProfile,
    PriceKind,
    RowRecord,
    SheetProfile,
    StagedQuoteRecord,
    WorkbookProfile,
)


_NUMBER = r"[+-]?(?:\d+(?:\.\d+)?|\.\d+)"
_CNY = r"(?:元|人民币|cny|rmb|¥)"
_WS = re.compile(r"\s+")
_CONDITIONAL_MARKERS = ("免费", "超出", "超过", "以内", "以上", "以下", "起步", "封顶", "按需")


@dataclass(frozen=True)
class _FormulaMetadata:
    kind: FormulaKind
    text: Optional[str]
    formula_range: Optional[str]
    shared_index: Optional[str]


def normalize_header(value: Any) -> str:
    """Return NFKC + trimmed + whitespace-collapsed + casefolded text."""

    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return _WS.sub(" ", text.strip()).casefold()


def _decimal(text: str) -> Optional[Decimal]:
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _unknown(raw_value: Any, raw_text: str, reason: str) -> ParsedPrice:
    return ParsedPrice(
        raw_value=raw_value,
        raw_text=raw_text,
        kind=PriceKind.UNKNOWN,
        needs_review=True,
        review_reasons=(reason,),
    )


def _condition_dimensions(text: str) -> Tuple[Optional[str], Optional[str]]:
    currency = "CNY" if re.search(_CNY, text, flags=re.IGNORECASE) else None
    if re.search(r"百万tokens?", text, flags=re.IGNORECASE):
        return currency, "million_tokens"
    if re.search(r"tokens?", text, flags=re.IGNORECASE):
        return currency, "token"
    if "万字符" in text:
        return currency, "ten_thousand_characters"
    if "个音色" in text:
        return currency, "voice"
    if "秒" in text:
        return currency, "second"
    if "次" in text:
        return currency, "request"
    if "张" in text:
        return currency, "image"
    return currency, None


def parse_price(value: Any, profile: Optional[ParserProfile] = None) -> ParsedPrice:
    """Classify a quote value without guessing missing context.

    Decimal is used for every parsed number. Conditional prose intentionally
    keeps amount unset even when it contains a numeric fragment.
    """

    profile = profile or ParserProfile()
    raw_text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", raw_text).strip()
    compact = _WS.sub("", text)
    folded = compact.casefold()

    if not compact:
        return _unknown(value, raw_text, "empty_price")

    if any(marker in folded for marker in _CONDITIONAL_MARKERS):
        currency, unit = _condition_dimensions(folded)
        return ParsedPrice(
            raw_value=value,
            raw_text=raw_text,
            kind=PriceKind.CONDITIONAL,
            currency=currency,
            unit=unit,
            condition_text=text,
            needs_review=True,
            review_reasons=("conditional_price_requires_review",),
        )

    discount = re.fullmatch(rf"({_NUMBER})折", folded)
    if discount:
        amount = _decimal(discount.group(1))
        if amount is None:
            return _unknown(value, raw_text, "invalid_decimal")
        valid = Decimal("0") < amount <= Decimal("10")
        return ParsedPrice(
            raw_value=value,
            raw_text=raw_text,
            kind=PriceKind.DISCOUNT,
            amount=amount,
            unit="discount_multiplier",
            discount_multiplier=amount / Decimal("10"),
            needs_review=not valid,
            review_reasons=() if valid else ("discount_out_of_range",),
        )

    rules = (
        (PriceKind.CNY_PER_MILLION_TOKENS, "million_tokens", rf"({_NUMBER}){_CNY}(?:/|每)百万tokens?"),
        (PriceKind.CNY_PER_TOKEN, "token", rf"({_NUMBER}){_CNY}(?:/|每)tokens?"),
        (PriceKind.CNY_PER_TEN_THOUSAND_CHARACTERS, "ten_thousand_characters", rf"({_NUMBER}){_CNY}(?:/|每)万字符"),
        (PriceKind.CNY_PER_VOICE, "voice", rf"({_NUMBER}){_CNY}(?:/|每)个音色"),
        (PriceKind.CNY_PER_SECOND, "second", rf"({_NUMBER}){_CNY}(?:/|每)秒"),
        (PriceKind.CNY_PER_REQUEST, "request", rf"({_NUMBER}){_CNY}(?:/|每)次"),
        (PriceKind.CNY_PER_IMAGE, "image", rf"({_NUMBER}){_CNY}(?:/|每)张"),
    )
    for kind, unit, pattern in rules:
        match = re.fullmatch(pattern, folded, flags=re.IGNORECASE)
        if match:
            amount = _decimal(match.group(1))
            if amount is None:
                return _unknown(value, raw_text, "invalid_decimal")
            valid = amount >= Decimal("0")
            return ParsedPrice(
                raw_value=value,
                raw_text=raw_text,
                kind=kind,
                amount=amount,
                currency="CNY",
                unit=unit,
                needs_review=not valid,
                review_reasons=() if valid else ("negative_price",),
            )

    missing_currency = re.fullmatch(rf"({_NUMBER})(?:/|每)秒", folded)
    if missing_currency:
        amount = _decimal(missing_currency.group(1))
        if amount is None:
            return _unknown(value, raw_text, "invalid_decimal")
        reasons = ["currency_missing"]
        if amount < Decimal("0"):
            reasons.append("negative_price")
        return ParsedPrice(
            raw_value=value,
            raw_text=raw_text,
            kind=PriceKind.UNKNOWN_CURRENCY_PER_SECOND,
            amount=amount,
            unit="second",
            needs_review=True,
            review_reasons=tuple(reasons),
        )

    bare_number = re.fullmatch(_NUMBER, folded)
    if bare_number:
        amount = _decimal(bare_number.group(0))
        if amount is None:
            return _unknown(value, raw_text, "invalid_decimal")
        if profile.discount_scale == Decimal("10"):
            valid = Decimal("0") < amount <= Decimal("10")
            return ParsedPrice(
                raw_value=value,
                raw_text=raw_text,
                kind=PriceKind.DISCOUNT,
                amount=amount,
                unit="discount_multiplier",
                discount_multiplier=amount / Decimal("10"),
                needs_review=not valid,
                review_reasons=() if valid else ("discount_out_of_range",),
            )
        return _unknown(value, raw_text, "bare_number_requires_discount_scale_10")

    return _unknown(value, raw_text, "unrecognized_price_format")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _domain_value(value: Any) -> Any:
    if isinstance(value, float):
        return Decimal(str(value))
    return value


def _formula_text(value: Any) -> Optional[str]:
    candidate = value
    if not isinstance(candidate, str):
        candidate = getattr(value, "text", None)
    if not isinstance(candidate, str) or not candidate:
        return None
    if not candidate.startswith("="):
        candidate = "=" + candidate
    return candidate


def _relationship_target(target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join("xl", target))


def _formula_metadata(path: Path) -> Mapping[str, Mapping[str, _FormulaMetadata]]:
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    document_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"

    import xml.etree.ElementTree as element_tree

    with ZipFile(path) as archive:
        workbook_root = element_tree.fromstring(archive.read("xl/workbook.xml"))
        rels_root = element_tree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            rel.get("Id"): _relationship_target(rel.get("Target", ""))
            for rel in rels_root.findall("{" + package_rel_ns + "}Relationship")
        }
        result: Dict[str, Dict[str, _FormulaMetadata]] = {}
        sheets = workbook_root.find("{" + spreadsheet_ns + "}sheets")
        if sheets is None:
            return result
        for sheet in sheets:
            title = sheet.get("name", "")
            relationship_id = sheet.get("{" + document_rel_ns + "}id")
            target = targets.get(relationship_id)
            metadata: Dict[str, _FormulaMetadata] = {}
            if target and target in archive.namelist():
                sheet_root = element_tree.fromstring(archive.read(target))
                for cell in sheet_root.findall(".//{" + spreadsheet_ns + "}c"):
                    formula = cell.find("{" + spreadsheet_ns + "}f")
                    coordinate = cell.get("r")
                    if formula is None or not coordinate:
                        continue
                    formula_type = formula.get("t")
                    if formula_type == "shared":
                        kind = FormulaKind.SHARED
                    elif formula_type == "array":
                        kind = FormulaKind.ARRAY
                    else:
                        kind = FormulaKind.NORMAL
                    formula_text = formula.text
                    if formula_text and not formula_text.startswith("="):
                        formula_text = "=" + formula_text
                    metadata[coordinate] = _FormulaMetadata(
                        kind=kind,
                        text=formula_text,
                        formula_range=formula.get("ref"),
                        shared_index=formula.get("si"),
                    )
            result[title] = metadata
        return result


def _merge_map(worksheet: Any) -> Tuple[Tuple[str, ...], Mapping[str, Tuple[str, str]]]:
    ranges = tuple(str(item) for item in worksheet.merged_cells.ranges)
    mapping: Dict[str, Tuple[str, str]] = {}
    for item in worksheet.merged_cells.ranges:
        source = worksheet.cell(row=item.min_row, column=item.min_col).coordinate
        merge_range = str(item)
        for row_index in range(item.min_row, item.max_row + 1):
            for column_index in range(item.min_col, item.max_col + 1):
                coordinate = worksheet.cell(row=row_index, column=column_index).coordinate
                mapping[coordinate] = (merge_range, source)
    return ranges, mapping


def _build_cell(
    formula_sheet: Any,
    cached_sheet: Any,
    coordinate: str,
    merge_mapping: Mapping[str, Tuple[str, str]],
    formula_metadata: Mapping[str, _FormulaMetadata],
    missing_header: bool,
) -> CellRecord:
    merge_range, source_coordinate = merge_mapping.get(coordinate, (None, coordinate))
    source_cell = formula_sheet[source_coordinate]
    cached_cell = cached_sheet[source_coordinate]
    metadata = formula_metadata.get(source_coordinate)
    formula_text = _formula_text(source_cell.value)
    if formula_text is None and metadata is not None:
        formula_text = metadata.text
    is_formula = metadata is not None or formula_text is not None or source_cell.data_type == "f"

    reasons = []
    if is_formula:
        reasons.append("formula_not_executed")
    if missing_header and (
        source_cell.value is not None or cached_cell.value is not None or is_formula
    ):
        reasons.append("missing_header")

    return CellRecord(
        coordinate=coordinate,
        column_index=formula_sheet[coordinate].column,
        row_index=formula_sheet[coordinate].row,
        raw_value=formula_text if is_formula else _domain_value(source_cell.value),
        cached_value=_domain_value(cached_cell.value) if is_formula else None,
        formula_text=formula_text,
        formula_kind=(metadata.kind if metadata else FormulaKind.NORMAL) if is_formula else None,
        formula_range=metadata.formula_range if metadata else None,
        shared_index=metadata.shared_index if metadata else None,
        merge_range=merge_range,
        source_coordinate=source_coordinate,
        needs_review=bool(reasons),
        review_reasons=tuple(reasons),
    )


def _non_empty_data_rows(rows: Iterable[RowRecord], header_row: int) -> int:
    count = 0
    for row in rows:
        if row.row_index <= header_row:
            continue
        if any(
            cell.raw_value is not None
            or cell.cached_value is not None
            or cell.formula_text is not None
            for cell in row.cells
        ):
            count += 1
    return count


def parse_excel(
    source_path: Any,
    profile: Optional[ParserProfile] = None,
) -> ExcelParseResult:
    """Profile an xlsx file and emit traceable staged quote candidates."""

    profile = profile or ParserProfile()
    path = Path(source_path).expanduser().resolve()
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError(
            "openpyxl is required to parse workbooks; dependency installation was not attempted"
        ) from error

    digest = _sha256(path)
    byte_size = path.stat().st_size
    metadata_warning = None
    try:
        all_formula_metadata = _formula_metadata(path)
    except Exception as error:
        all_formula_metadata = {}
        metadata_warning = "formula_metadata_unavailable:" + error.__class__.__name__

    formula_workbook = load_workbook(path, data_only=False, read_only=False)
    cached_workbook = load_workbook(path, data_only=True, read_only=False)
    sheet_profiles = []
    staged_quotes = []

    try:
        normalized_price_headers = {
            normalize_header(header) for header in profile.price_headers
        }
        for sheet_index, title in enumerate(formula_workbook.sheetnames):
            formula_sheet = formula_workbook[title]
            cached_sheet = cached_workbook[title]
            merge_ranges, merge_mapping = _merge_map(formula_sheet)
            sheet_formula_metadata = all_formula_metadata.get(title, {})

            headers = []
            for column_index in range(1, formula_sheet.max_column + 1):
                coordinate = formula_sheet.cell(
                    row=profile.header_row, column=column_index
                ).coordinate
                source_coordinate = merge_mapping.get(coordinate, (None, coordinate))[1]
                raw_header = _domain_value(formula_sheet[source_coordinate].value)
                normalized = normalize_header(raw_header)
                missing = not normalized
                headers.append(
                    ColumnHeader(
                        column_index=column_index,
                        column_letter=get_column_letter(column_index),
                        coordinate=coordinate,
                        raw_value=raw_header,
                        normalized_value=normalized,
                        needs_review=missing,
                        review_reasons=("empty_header",) if missing else (),
                    )
                )

            missing_header_columns = {
                header.column_index for header in headers if header.needs_review
            }
            rows = []
            for row_index in range(1, formula_sheet.max_row + 1):
                cells = tuple(
                    _build_cell(
                        formula_sheet=formula_sheet,
                        cached_sheet=cached_sheet,
                        coordinate=formula_sheet.cell(
                            row=row_index, column=column_index
                        ).coordinate,
                        merge_mapping=merge_mapping,
                        formula_metadata=sheet_formula_metadata,
                        missing_header=(
                            row_index > profile.header_row
                            and column_index in missing_header_columns
                        ),
                    )
                    for column_index in range(1, formula_sheet.max_column + 1)
                )
                row_reasons = tuple(
                    sorted({reason for cell in cells for reason in cell.review_reasons})
                )
                rows.append(
                    RowRecord(
                        row_index=row_index,
                        cells=cells,
                        needs_review=bool(row_reasons),
                        review_reasons=row_reasons,
                    )
                )

            sheet_reasons = []
            if missing_header_columns:
                sheet_reasons.extend(
                    "empty_header:" + get_column_letter(index)
                    for index in sorted(missing_header_columns)
                )
            if any(
                cell.formula_text is not None
                for row in rows
                for cell in row.cells
            ):
                sheet_reasons.append("contains_unexecuted_formulas")

            sheet_profile = SheetProfile(
                title=title,
                sheet_index=sheet_index,
                max_row=formula_sheet.max_row,
                max_column=formula_sheet.max_column,
                header_row=profile.header_row,
                data_row_count=_non_empty_data_rows(rows, profile.header_row),
                merge_ranges=merge_ranges,
                headers=tuple(headers),
                rows=tuple(rows),
                needs_review=bool(sheet_reasons),
                review_reasons=tuple(sheet_reasons),
            )
            sheet_profiles.append(sheet_profile)

            price_columns = {
                header.column_index: header
                for header in headers
                if header.normalized_value in normalized_price_headers
            }
            for row in rows:
                if row.row_index <= profile.header_row:
                    continue
                for column_index, header in price_columns.items():
                    cell = row.cell_at_column(column_index)
                    parse_value = (
                        cell.cached_value
                        if cell.formula_text is not None and cell.cached_value is not None
                        else cell.raw_value
                    )
                    if parse_value is None:
                        continue
                    parsed = parse_price(parse_value, profile=profile)
                    review_reasons = tuple(
                        dict.fromkeys(cell.review_reasons + parsed.review_reasons)
                    )
                    staged_quotes.append(
                        StagedQuoteRecord(
                            workbook_sha256=digest,
                            sheet_name=title,
                            source_row=row.row_index,
                            source_coordinate=cell.coordinate,
                            header=header,
                            cell=cell,
                            parsed_price=parsed,
                            needs_review=bool(review_reasons),
                            review_reasons=review_reasons,
                        )
                    )
    finally:
        formula_workbook.close()
        cached_workbook.close()

    workbook_reasons = []
    if metadata_warning:
        workbook_reasons.append(metadata_warning)
    workbook_reasons.extend(
        "sheet_requires_review:" + sheet.title
        for sheet in sheet_profiles
        if sheet.needs_review
    )
    workbook = WorkbookProfile(
        source_path=path,
        source_name=path.name,
        sha256=digest,
        byte_size=byte_size,
        sheets=tuple(sheet_profiles),
        needs_review=bool(workbook_reasons),
        review_reasons=tuple(workbook_reasons),
    )
    return ExcelParseResult(workbook=workbook, staged_quotes=tuple(staged_quotes))
