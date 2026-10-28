#!/usr/bin/env python3
"""Synthetic real-API/browser proof for simplified CRUD. Never touches production or Qwen."""
from __future__ import annotations
import argparse
import csv
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import httpx
from playwright.sync_api import expect, sync_playwright
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8893)
    parser.add_argument("--chromium", default="/home/user/.cache/ms-playwright/chromium-1243/chrome-linux64/chrome")
    parser.add_argument("--evidence-dir", type=Path, default=Path("/tmp/tokensql-simplify-real-evidence"))
    args = parser.parse_args()
    with socket.socket() as probe:
        if probe.connect_ex(("127.0.0.1", args.port)) == 0:
            raise RuntimeError("QA port is occupied; will not stop another process")
    root = Path(tempfile.mkdtemp(prefix="tokensql-simplify-real-"))
    args.evidence_dir.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    for name in ("workspace", "agent", "input", "progress", "data"):
        env["TOKENSQL_" + name.upper() + "_DIR"] = str(root / name)
    env.update(TOKENSQL_TEST_MODE="true", TOKENSQL_FX_AUTO_REFRESH="false",
               TOKENSQL_OPENCODE_URL="http://127.0.0.1:1", PYTHONDONTWRITEBYTECODE="1")
    log = (root / "server.log").open("w")
    process = subprocess.Popen([sys.executable, "-m", "uvicorn", "tokensql.workspace.app:app",
                                "--host", "127.0.0.1", "--port", str(args.port)],
                               cwd=PROJECT, env=env, stdout=log, stderr=subprocess.STDOUT)
    base = "http://127.0.0.1:" + str(args.port)
    client = httpx.Client(base_url=base, timeout=15)
    def request(path, body=None, method=None):
        result = client.request(method or ("POST" if body is not None else "GET"),
                                "/api/workspace/" + path, json=body)
        result.raise_for_status()
        return result.json()
    def quotes():
        return request("quotes?limit=500")["items"]
    try:
        for _ in range(100):
            if process.poll() is not None:
                raise RuntimeError((root / "server.log").read_text())
            try:
                if client.get("/health").json().get("isolated_test"): break
            except Exception: pass
            time.sleep(0.1)
        else: raise RuntimeError("Isolated service did not start")
        errors = []
        bulk_requests = []
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True, executable_path=args.chromium)
            page = browser.new_page(viewport={"width":1500,"height":960}, accept_downloads=True)
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.on("request", lambda req: bulk_requests.append(req.post_data_json)
                    if req.url.endswith("/quotes/bulk-delete") else None)
            page.goto(base + "/workspace", wait_until="networkidle")
            assert request("catalog")["models"] == []
            page.locator("#newQuoteButton").click()
            expect(page.locator('[name=currency]')).to_have_value("CNY")
            expect(page.locator('[name=pricing_unit_id]')).to_have_value("")
            page.locator('[name=model_id]').fill("qa-new-without-entity")
            page.locator('[name=discount_rate]').fill("1.5")
            page.locator("#saveQuoteButton").click()
            expect(page.locator("#quoteRows tr")).to_have_count(1)
            first = quotes()[0]
            assert first["discount_rate"] == "0.15" and first["currency"] == "CNY"
            assert first["pricing_unit_id"] is None and first["unit_quantity"] is None
            assert len(request("catalog")["models"]) == 1
            expect(page.locator("#quoteRows")).to_contain_text("源折扣 1.5折")
            assert "物理标识" not in page.content()
            seeded = request("quotes", {"items":[
                {"model_id":"qa-high-precision","discount_rate":"0.1555","currency":"USD",
                 "amount":"0.012345678901234567890"},
                {"model_id":"qa-half-005","discount_rate":"0.005"},
                {"model_id":"qa-half-155","discount_rate":"0.155"},
                {"model_id":"qa-below-half","discount_rate":"0.1549"},
                {"model_id":"qa-zero","discount_rate":"0"},
                {"model_id":"qa-one","discount_rate":"1"},
                {"model_id":"qa-premium","discount_rate":"1.255"},
                {"model_id":"qa-explicit-unknown","currency":None,"discount_rate":"0.85"},
            ]})["items"]
            expected_ratios = {"qa-high-precision":"0.16","qa-half-005":"0.01","qa-half-155":"0.16",
                               "qa-below-half":"0.15","qa-zero":"0.00","qa-one":"1.00","qa-premium":"1.26",
                               "qa-explicit-unknown":"0.85"}
            assert {item["model_id"]:item["discount_rate"] for item in seeded} == expected_ratios
            for item in seeded:
                page.locator("#refreshButton").click()
                row = page.locator('#quoteRows tr[data-id="' + str(item["id"]) + '"]')
                row.get_by_role("button", name="编辑报价 " + item["model_id"], exact=True).click()
                expect(page.locator('[name=currency]')).to_have_value(item["currency"] or "")
                if item["model_id"] == "qa-high-precision":
                    expect(page.locator('[name=discount_rate]')).to_have_value("1.6")
                page.locator('[name=note]').fill("unrelated edit preserves fixed-two ratio/currency/amount")
                page.locator("#saveQuoteButton").click()
                expect(page.locator("#editorDialog")).not_to_be_visible()
                current = request("quotes/" + str(item["id"]))
                assert current["discount_rate"] == item["discount_rate"]
                assert current["currency"] == item["currency"]
                assert current["amount"] == item["amount"]
            job = request("imports/drafts", {"bundle":{"official_prices":[{
                "model_id":"qa-new-without-entity","billing_metric":"input","billing_item":"input","amount":"20","pricing_unit_id":"cny_mtoken",
                "explicit_official":True,"source":"Synthetic fixture","source_ref":"Test!A1"}]}})
            request("imports/" + job["id"] + "/confirm", {
                "expected_revision":job["revision"],"idempotency_key":"synthetic-" + job["id"]})
            request("quotes", {"items":[{"model_id":"qa-new-without-entity","billing_metric":"input","amount":"0.0031",
                                        "pricing_unit_id":"cny_ktoken"}]})
            page.locator("#refreshButton").click()
            page.locator("#pricingConversionButton").click()
            expect(page.locator("#quoteRows")).to_contain_text("按官网反算折扣 1.6折")
            # Display/export rounds only the multiplier; exact ratio and raw amount remain available.
            before_exports = quotes()
            for format in ("csv", "xlsx"):
                page.locator("#exportFormat").select_option(format)
                with page.expect_download() as download:
                    page.locator("#exportFilteredButton").click()
                output = Path(download.value.path()).read_bytes()
                (args.evidence_dir / ("discount-two." + format)).write_bytes(output)
                if format == "xlsx":
                    book = load_workbook(BytesIO(output), read_only=True)
                    exported = list(book.active.iter_rows(values_only=True))
                else:
                    exported = list(csv.reader(StringIO(output.decode("utf-8-sig"))))
                columns = exported[0]
                rows = [dict(zip(columns, row)) for row in exported[1:]]
                calculated = next(row for row in rows if row["原价"] == "0.0031")
                assert calculated["按官网反算折扣值（两位小数）"] == "0.16"
                assert calculated["按官网反算折扣（折，一位小数）"] == "1.6"
                assert calculated["反算倍率精确分子"] == "31" and calculated["反算倍率精确分母"] == "200"
                assert str(calculated["反算倍率显示为近似"]).lower() == "true"
                assert "HALF_UP" in calculated["反算倍率舍入规则"]
                original = next(row for row in rows if row["模型 ID"] == "qa-high-precision")
                assert original["原始存储折扣值（两位小数）"] == "0.16"
                assert original["原价"] == "0.012345678901234567890"
            assert quotes() == before_exports
            request("quotes", {"items":[{"model_id":"qa-bulk-" + str(i)} for i in range(105)]})
            page.locator("#refreshButton").click()
            expect(page.locator("#pageLabel")).to_have_text("1 / 2")
            first_page = page.locator("#quoteRows tr").first
            first_id = int(first_page.get_attribute("data-id"))
            first_page.locator('input[type=checkbox]').check()
            page.locator("#nextPageButton").click()
            expect(page.locator("#pageLabel")).to_have_text("2 / 2")
            second_page = page.locator("#quoteRows tr").first
            second_id = int(second_page.get_attribute("data-id"))
            second_page.locator('input[type=checkbox]').check()
            expect(page.locator("#selectedCount")).to_have_text("2")
            before = quotes()
            def cancel(dialog):
                assert "2 条" in dialog.message and "永久删除" in dialog.message
                dialog.dismiss()
            page.once("dialog", cancel)
            page.locator("#deleteSelectedButton").click()
            assert bulk_requests == [] and quotes() == before
            # A concurrent edit makes the entire selected batch stale; no partial delete.
            chosen = next(item for item in before if item["id"] == first_id)
            request("quotes/" + str(first_id), {"expected_revision":chosen["revision"],"note":"concurrent"},
                    method="PATCH")
            page.once("dialog", lambda dialog: dialog.accept())
            page.locator("#deleteSelectedButton").click()
            expect(page.locator("#toastRegion")).to_contain_text("本批次全部未删除")
            expect(page.locator("#selectionBar")).not_to_be_visible()
            assert len(quotes()) == len(before) and len(bulk_requests) == 1
            # Re-select from refreshed pages, then delete once across both pages.
            page.locator("#previousPageButton").click()
            expect(page.locator("#pageLabel")).to_have_text("1 / 2")
            page.locator('#quoteRows tr[data-id="' + str(first_id) + '"] input').check()
            page.locator("#nextPageButton").click()
            expect(page.locator("#pageLabel")).to_have_text("2 / 2")
            page.locator('#quoteRows tr[data-id="' + str(second_id) + '"] input').check()
            catalog_before = request("catalog")
            page.once("dialog", lambda dialog: dialog.accept())
            page.locator("#deleteSelectedButton").click()
            expect(page.locator("#toastRegion")).to_contain_text("已永久删除 2 条报价")
            after = quotes()
            assert len(after) == len(before) - 2 and len(bulk_requests) == 2
            assert {item["id"] for item in bulk_requests[-1]["items"]} == {first_id, second_id}
            assert not {first_id,second_id} & {item["id"] for item in after}
            catalog_after = request("catalog")
            assert all(catalog_after[key] == catalog_before[key] for key in ("models","official_prices","suppliers"))
            with sqlite3.connect(root / "workspace" / "quotes.sqlite3") as db:
                assert db.execute("SELECT count(*) FROM pricing_quote WHERE id IN (?,?)",
                                  (first_id,second_id)).fetchone()[0] == 0
            page.screenshot(path=str(args.evidence_dir / "simplified-desktop.png"), full_page=True)
            page.set_viewport_size({"width":390,"height":844})
            assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
            page.screenshot(path=str(args.evidence_dir / "simplified-mobile.png"), full_page=True)
            browser.close()
        assert not errors, errors
        result = {"passed":True,"kind":"synthetic_real_api_browser","root":str(root),
                  "production_accessed":False,"real_model_called":False,
                  "auto_created_model":True,"default_cny":True,"foreign_and_null_preserved":True,
                  "discount_half_up_fixed_two_persisted":True,"amount_precision_preserved":True,
                  "derived_ratio_fixed_two_only_display":True,"fold_half_up_one_decimal":True,
                  "csv_xlsx_fixed_two_and_exact_evidence":True,"exports_no_quote_writes":True,
                  "cross_page_deleted_ids":[first_id,second_id],"cancel_no_request":True,
                  "conflict_entire_batch_unchanged":True,"physical_rows_deleted":True,
                  "models_and_official_unchanged":True}
        (args.evidence_dir / "result.json").write_text(json.dumps(result,ensure_ascii=False,indent=2) + "\n")
        print(json.dumps(result,ensure_ascii=False))
    finally:
        client.close()
        process.terminate()
        try: process.wait(timeout=10)
        except subprocess.TimeoutExpired: process.kill(); process.wait()
        log.close()


if __name__ == "__main__":
    main()
