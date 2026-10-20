#!/usr/bin/env python3
"""Remote browser regression; all API traffic is intercepted, never production writes."""
from __future__ import annotations

import argparse
import csv
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO, StringIO
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import json
import os
from pathlib import Path
from threading import Thread
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook, load_workbook
from playwright.sync_api import Route, expect, sync_playwright

STATIC = Path(__file__).resolve().parents[1] / "src/tokensql/workspace/static"

PRICING_UNITS = [
    {"id": f"{currency.lower()}_{scale}", "currency": currency, "unit": unit, "unit_quantity": quantity,
     "dimension": dimension, "scale": scale, "label": f"{currency} / {label}", "symbol": symbol}
    for currency, symbol in [("USD", "$"), ("HKD", "HK$"), ("CNY", "¥")]
    for scale, unit, quantity, dimension, label in [
        ("ktoken", "token", "1000", "token", "千 token"),
        ("mtoken", "token", "1000000", "token", "百万 token"),
        ("second", "second", "1", "time", "秒"),
        ("image", "image", "1", "image", "张"),
    ]
]
UNIT_BY_ID = {item["id"]: item for item in PRICING_UNITS}


def model(identifier: str = "smoke-a") -> dict:
    return {"model_id": identifier, "vendor": "Smoke Labs", "model_name": "Smoke Model " + identifier,
            "aliases": ["alias-" + identifier], "revision": 1}


def official(identifier: str = "smoke-a") -> dict:
    return {"model_id": identifier, "billing_item": "input", "amount": "20", "pricing_unit_id": "cny_mtoken",
            "currency": "CNY", "unit": "token", "unit_quantity": "1000000", "explicit_official": True,
            "source": "Explicit official spreadsheet label", "source_ref": "Price!B2", "revision": 1}


def quote(identifier: int = 1) -> dict:
    return {"id": identifier, "model_id": "smoke-a", "model_name": "Smoke Model smoke-a", "vendor": "Smoke Labs",
            "supplier": "Smoke Supplier", "provider": "Smoke Supplier", "billing_item": "input", "tags": ["premium", "cn"],
            "supplier_model_id": "raw-smoke-channel", "amount": "10", "discount_rate": "0.5",
            "pricing_unit_id": "cny_mtoken", "currency": "CNY", "unit": "token", "unit_quantity": "1000000",
            "conditions": {}, "source": "Smoke sheet",
            "source_ref": "Sheet1!D8", "revision": 1, "updated_at": "2026-10-16T00:00:00Z"}


def bundle() -> dict:
    item = {key: value for key, value in quote().items() if key not in {
        "id", "model_name", "vendor", "provider", "revision", "updated_at",
        "currency", "unit", "unit_quantity"}}
    item["supplier_model_id"] = "review-channel"
    return {"models": [], "official_prices": [], "create": [item], "update": [], "delete": []}


class Fixture:
    def __init__(self):
        self.quotes = [quote()]
        self.files = []
        self.jobs = {}
        self.count = 0
        self.drifted = False
        self.official_amount = "20"
        self.context = {"revision": 1, "updated_at": "2026-10-16T00:00:00Z", "fx_rates": [
            {"currency": "USD", "cny_per_unit": "7", "source": "Synthetic FX", "observed_at": "2026-10-16T00:00:00Z"},
            {"currency": "HKD", "cny_per_unit": "0.9", "source": "Synthetic FX", "observed_at": "2026-10-16T00:00:00Z"}], "token_equivalences": []}
        self.confirmed = 0
        self.cancelled = 0
        self.exports = []
        self.view_exports = []
        self.starts = []
        self.preview_requests = []
        self.lookup_requests = []
        self.next_failure = None
        self.fail_poll_count = 0
        self.history_reply = None
        self.next_history_reply = None
        self.mutations = []
        self.sessions = [{"id": "saved-session", "title": "Persistent session"}]
        self.memory = {"revision": 1, "rules": [{"id": "rule-1", "scope": "supplier:Smoke Supplier", "text": "保留原渠道 ID",
                        "enabled": True, "source": "confirmed", "import_id": "old-import"}]}

    def catalog(self):
        price = official(); price["amount"] = self.official_amount
        price["normalized"] = self.normalized(price)
        return {"models": [model()], "official_prices": [price],
                "suppliers": [{"name": "Smoke Supplier", "supplier_id": "smoke-supplier", "revision": 1}]}

    def normalized(self, item):
        amount = item.get("amount")
        if amount is None: amount = item.get("derived_amount")
        rates = {"CNY": Decimal("1"), **{row["currency"]: Decimal(row["cny_per_unit"]) for row in self.context["fx_rates"]}}
        currency = item.get("currency", "CNY")
        result = {"status": "exact", "amount": None, "currency": "CNY", "unit": "token", "unit_quantity": "1000000",
                  "issues": [], "evidence": {"context_revision": self.context["revision"], "fx": self.context["fx_rates"]}}
        if currency not in rates:
            result.update(status="unavailable", issues=[{"code": "missing_fx", "message": "请补充汇率来源"}]); return result
        if amount is None:
            result.update(status="unavailable", issues=[{"code": "missing_price", "message": "缺少官网金额依据"}]); return result
        per_unit = Decimal(amount) / Decimal(item["unit_quantity"]) * rates[currency]
        if item["unit"] != "token":
            factors = [factor for factor in self.context["token_equivalences"] if factor["model_id"] == item["model_id"] and
                       factor["billing_item"] == item["billing_item"] and factor["unit"] == item["unit"]]
            if not factors:
                result.update(status="unavailable", issues=[{"code": "missing_token_equivalence", "message": "请补充模型专属 token 参数"}]); return result
            per_unit /= Decimal(factors[0]["tokens_per_unit"]); result["status"] = "estimated"
            result["evidence"]["token_equivalence"] = factors[0]
        result["amount"] = str(per_unit * Decimal("1000000"))
        return result

    def compute(self, item):
        value = deepcopy(item)
        pricing_unit_id = value.get("pricing_unit_id")
        if pricing_unit_id:
            unit = UNIT_BY_ID[pricing_unit_id]
            value.update(currency=unit["currency"], unit=unit["unit"], unit_quantity=unit["unit_quantity"])
        if value.get("discount_rate") is not None:
            value["discount_rate"] = format(Decimal(value["discount_rate"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), ".2f")
        value["raw_discount_rate"] = value.get("discount_rate")
        value["derived_amount"] = None
        value["normalized"] = {}
        value["discount_status"] = "raw"
        value["discount_evidence"] = {}
        return value

    def preview(self, proposal):
        return {"items": [self.compute(item) for item in proposal.get("create", [])],
                "changes": {"added": len(proposal.get("create", [])), "updated": 0, "deleted": 0},
                "catalog": self.catalog(), "issues": []}

    def route(self, route: Route):
        request = route.request
        url = urlparse(request.url)
        path = url.path.removeprefix("/api/workspace")
        method = request.method
        data = request.post_data_json if method in {"POST", "PATCH", "PUT"} and "application/json" in request.headers.get("content-type", "") else {}
        if method in {"POST", "PATCH", "PUT", "DELETE"} and path != "/exports/view":
            self.mutations.append((method, path, deepcopy(data)))
        def respond(payload, status=200):
            route.fulfill(status=status, json=payload)
        if path == "/catalog":
            respond(self.catalog())
        elif path == "/pricing-units":
            respond({"version": "1", "items": PRICING_UNITS})
        elif path == "/pricing-context":
            if method == "PUT":
                assert data["expected_revision"] == self.context["revision"]
                self.context.update(revision=self.context["revision"]+1, fx_rates=data["fx_rates"], token_equivalences=data["token_equivalences"])
            respond(self.context)
        elif path == "/pricing-context/refresh-fx":
            assert data["expected_revision"] == self.context["revision"]
            self.context["revision"] += 1
            self.context["fx_rates"][0]["source"] = "ECB synthetic fixture"
            respond(self.context)
        elif path == "/quotes/preview":
            self.preview_requests.append(data)
            respond({"items": [self.compute(item) for item in data["items"]], "issues": []})
        elif path == "/quotes" and method == "GET":
            query = parse_qs(url.query)
            values = self.quotes
            if query.get("model_id"):
                values = [item for item in values if item["model_id"] == query["model_id"][0]]
            offset = int(query.get("offset", ["0"])[0])
            limit = int(query.get("limit", ["100"])[0])
            respond({"items": [self.compute(item) for item in values[offset:offset + limit]], "total": len(values), "filter_options": {
                "vendors": ["Smoke Labs"], "suppliers": ["Smoke Supplier"], "billing_items": ["input"], "tags": ["premium", "cn"], "currencies": ["CNY", "USD", "HKD"]}})
        elif path == "/quotes" and method == "POST":
            item = self.compute(data["items"][0]); item.update(id=len(self.quotes)+10, revision=1)
            self.quotes.append(item); respond({"items": [item]})
        elif path.startswith("/quotes/") and method in {"PATCH", "DELETE"}:
            identifier = int(path.split("/")[-1])
            old = next(item for item in self.quotes if item["id"] == identifier)
            if method == "DELETE":
                assert parse_qs(url.query)["expected_revision"] == [str(old["revision"])]
                self.quotes.remove(old); route.fulfill(status=204)
            else:
                assert data["expected_revision"] == old["revision"]
                old.update(self.compute({**old, **{k: v for k, v in data.items() if k != "expected_revision"}}))
                old["revision"] += 1; respond(old)
        elif path == "/exports":
            self.exports.append(parse_qs(url.query))
            route.fulfill(body="model_id,amount\nsmoke-a,10", headers={"Content-Disposition": "attachment; filename=smoke.csv",
                          "Content-Type": "text/csv"})
        elif path == "/exports/view":
            self.view_exports.append(data)
            columns, rows = data["columns"], data["rows"]
            safe = lambda value: "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
            if data["format"] == "xlsx":
                workbook = Workbook(); sheet = workbook.active
                sheet.append([column["label"] for column in columns])
                for row in rows: sheet.append([safe(row.get(column["key"], "")) for column in columns])
                output = BytesIO(); workbook.save(output)
                route.fulfill(body=output.getvalue(), headers={"Content-Disposition": "attachment; filename=converted.xlsx",
                              "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
            else:
                output = StringIO(); writer = csv.writer(output)
                writer.writerow([column["label"] for column in columns])
                writer.writerows([[safe(row.get(column["key"], "")) for column in columns] for row in rows])
                route.fulfill(body=output.getvalue(), headers={"Content-Disposition": "attachment; filename=converted.csv",
                              "Content-Type": "text/csv; charset=utf-8"})
        elif path == "/agent/status":
            respond({"available": True, "status": "ready", "model": "intercepted-smoke"})
        elif path == "/agent/sessions":
            respond({"items": self.sessions})
        elif path.startswith("/agent/sessions/") and path.endswith("/messages"):
            messages = [{"role":"user","text":"已保存的用户消息"},
                        {"role":"assistant","text":"持久会话 <img src=x onerror=alert(1)>"}]
            if self.history_reply:
                messages.extend([{"role":"user","text":"本轮新的用户请求"},
                                 {"role":"assistant","text":self.history_reply}])
            respond({"conversation_id":path.split("/")[-2],"messages":messages})
        elif path == "/agent/memory":
            if method == "PATCH":
                assert data["expected_revision"] == self.memory["revision"]
                self.memory = {"revision": self.memory["revision"]+1, "rules": data["rules"]}
            respond(self.memory)
        elif path == "/imports/files":
            respond({"items": self.files})
        elif path == "/imports/upload":
            assert method == "POST"
            upload = {"id": "file-" + str(len(self.files)+1), "filename": "browser.xlsx", "status": "input", "size": 12}
            self.files.append(upload); respond({"items": [upload]}, 201)
        elif path == "/imports":
            respond({"items": list(self.jobs.values())})
        elif path == "/official-prices/lookup":
            self.lookup_requests.append(deepcopy(data))
            identifier = "official-lookup-" + str(len(self.lookup_requests))
            proposal = {"models": [], "official_prices": [official(data["model_id"])], "create": [], "update": [], "delete": []}
            job = {"id": identifier, "status": "queued", "revision": 1,
                   "conversation_id": "fresh-official-session-" + str(len(self.lookup_requests)),
                   "intent": "official_price_lookup", "lookup_model_id": data["model_id"],
                   "lookup_billing_item": data.get("billing_item"), "bundle": proposal,
                   "preview": self.preview(proposal), "issues": [], "messages": [], "files": [], "polls": 0}
            self.sessions.append({"id": job["conversation_id"], "title": "Independent official lookup"})
            self.jobs[identifier] = job
            respond(job, 202)
        elif path in {"/imports/start", "/imports/drafts"}:
            self.count += 1; identifier = "review-" + str(self.count)
            proposal = data.get("bundle", bundle())
            job = {"id": identifier, "status": "ready" if path.endswith("/drafts") else "queued", "revision": 1,
                   "conversation_id": data.get("conversation_id") or "saved-session", "bundle": proposal,
                   "preview": self.preview(proposal), "issues": [], "assistant_message": None, "messages": [], "files": deepcopy(self.files), "polls": 0}
            if path.endswith("/start"):
                self.starts.append(data)
                for file in self.files:
                    if file["id"] in data["upload_ids"]: file["status"] = "progress"
            if self.next_history_reply:
                self.history_reply = self.next_history_reply; self.next_history_reply = None
            if self.next_failure is not None:
                old_reply = self.next_failure.get("old_reply") if isinstance(self.next_failure, dict) else None
                job.update(status="failed", error="proposal validation failed",
                           assistant_message=self.next_failure if isinstance(self.next_failure, str) else None,
                           messages=[{"role":"user","text":"previous turn"},{"role":"assistant","text":old_reply}] if old_reply else [])
                self.next_failure = None
            self.jobs[identifier] = job; respond(job, 202)
        elif path.startswith("/imports/"):
            identifier = path.split("/")[2]; job = self.jobs[identifier]
            if method == "GET" and self.fail_poll_count:
                self.fail_poll_count -= 1; route.abort("failed"); return
            if path.endswith("/proposal"):
                assert data["expected_revision"] == job["revision"]
                job.update(bundle=data["bundle"], revision=job["revision"]+1)
                job["preview"] = self.preview(job["bundle"])
            elif path.endswith("/confirm"):
                assert data["expected_revision"] == job["revision"]
                assert data["idempotency_key"] == "workspace-" + identifier
                if not self.drifted:
                    self.drifted = True; self.official_amount = "25"
                    job["revision"] += 1; job["preview"] = self.preview(job["bundle"])
                    respond({"detail": "Catalog changed; review again"}, 409)
                    return
                self.confirmed += 1
                for item in job["preview"]["items"]:
                    item = deepcopy(item); item.update(id=len(self.quotes)+20, revision=1); self.quotes.append(item)
                job.update(status="confirmed", changes=job["preview"]["changes"])
                for file in self.files: file["status"] = "data"
            elif path.endswith("/cancel"):
                self.cancelled += 1; job["status"] = "cancelled"
                for file in self.files: file["status"] = "input"
            elif job["status"] in {"queued", "running"}:
                job["polls"] += 1
                job["status"] = "running" if job["polls"] == 1 or self.count == 3 else "ready"
            respond(job)
        else:
            raise AssertionError("Unexpected API request; never forwarded: " + method + " " + path)


class StaticHandler(SimpleHTTPRequestHandler):
    def do_GET(self):
        if urlparse(self.path).path == "/workspace":
            self.path = "/index.html"
        elif self.path.startswith("/static/"):
            self.path = self.path.removeprefix("/static")
        super().do_GET()

    def log_message(self, *_):
        pass


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", help="Existing service; omitted starts an isolated static-only server.")
    parser.add_argument("--evidence-dir", type=Path)
    parser.add_argument("--chromium", help="Optional remote cached executable path.")
    args = parser.parse_args()
    server = None
    if args.url:
        url = args.url.rstrip("/")
    else:
        server = ThreadingHTTPServer(("127.0.0.1", 0), partial(StaticHandler, directory=str(STATIC)))
        Thread(target=server.serve_forever, daemon=True).start()
        url = "http://127.0.0.1:" + str(server.server_port)
    fixture = Fixture()
    page_errors, console_errors = [], []
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, executable_path=args.chromium)
        page = browser.new_page(viewport={"width": 1500, "height": 960}, accept_downloads=True)
        page.on("pageerror", lambda error: page_errors.append(str(error)))
        page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
        page.route("**/api/workspace/**", fixture.route)
        page.goto(url + "/workspace?session=saved-session", wait_until="networkidle")
        expect(page.locator("#quoteRows tr")).to_have_count(1)
        expect(page.locator("#quoteRows")).to_contain_text("raw-smoke-channel")
        expect(page.locator("#quoteRows")).to_contain_text("premium")
        expect(page.locator("#quoteRows")).to_contain_text("20")
        expect(page.locator("#conversation")).to_contain_text("已保存的用户消息")
        expect(page.locator("#conversation img")).to_have_count(0)
        assert len(page.locator(".content-grid").evaluate("el => getComputedStyle(el).gridTemplateColumns.split(' ')")) == 2

        # Browser conversion math preserves large/tiny Decimal strings and marks repeating values approximate.
        exact_values = page.evaluate("""() => {
          const exact = window.TokenSQLConversion._exact;
          return {
            tiny: exact.decimalString(exact.parseDecimal("1E-90")),
            huge: exact.decimalString(exact.parseDecimal("123456789012345678901234567890123456789012345678901234567890")),
            repeating: exact.decimalString(exact.divide(exact.parseDecimal("1"), exact.parseDecimal("3"))),
            fold15: window.TokenSQLConversion.discountFold("0.15"),
            fold155: window.TokenSQLConversion.discountFold("0.155"),
            foldTiny: window.TokenSQLConversion.discountFold("0.000001"),
            ratio15: window.TokenSQLConversion.ratioFromFold("1.5"),
            ratioRound: window.TokenSQLConversion.ratioFromFold("1.55"),
            ratios: ["0.005", "0.155", "0.1549", "0", "1", "0.8", "1.255"].map(window.TokenSQLConversion.discountRatio),
            fold1549: window.TokenSQLConversion.discountFold("0.1549"),
            ratioZero: window.TokenSQLConversion.ratioFromFold("0"),
            ratioOne: window.TokenSQLConversion.ratioFromFold("10"),
            ratioPremium: window.TokenSQLConversion.ratioFromFold("12.6"),
            negativeChecks: [
              () => window.TokenSQLConversion.discountRatio("-0.0001"),
              () => window.TokenSQLConversion.ratioFromFold("-0.04"),
              () => window.TokenSQLConversion.ratioFromFold("-0.05"),
            ].map((check) => { try { check(); return false; } catch (_) { return true; } }),
            negativeZero: window.TokenSQLConversion.discountRatio("-0.000"),
          };
        }""")
        assert exact_values["tiny"] == "0." + ("0" * 89) + "1"
        assert exact_values["huge"] == "123456789012345678901234567890123456789012345678901234567890"
        assert exact_values["repeating"].endswith("…（近似）")
        assert exact_values["fold15"] == "1.5" and exact_values["fold155"] == "1.6"
        assert exact_values["foldTiny"] == "0.0" and exact_values["ratio15"] == "0.15"
        assert exact_values["ratioRound"] == "0.16"
        assert exact_values["ratios"] == ["0.01", "0.16", "0.15", "0.00", "1.00", "0.80", "1.26"]
        assert exact_values["fold1549"] == "1.5"
        assert exact_values["negativeChecks"] == [True, True, True]
        assert exact_values["negativeZero"] == "0.00"
        assert [exact_values[name] for name in ("ratioZero", "ratioOne", "ratioPremium")] == ["0.00", "1.00", "1.26"]

        # Manual raw inputs stay independent; arithmetic preview is frontend-only.
        page.locator("#newQuoteButton").click()
        page.locator('[name=model_id]').fill("smoke-a")
        page.locator('[name=billing_item]').fill("input")
        page.locator('[name=supplier]').fill("Manual Supplier")
        page.locator('[name=discount_rate]').fill("8.0")
        page.locator('[name=pricing_unit_id]').select_option("cny_ktoken")
        expect(page.locator("#quotePreview")).to_contain_text("16 CNY")
        page.locator('[name=supplier]').press("Enter")
        assert len(fixture.quotes) == 1
        page.locator("#saveQuoteButton").click()
        expect(page.locator("#quoteRows tr")).to_have_count(2)
        page.locator("#quoteRows tr").last.get_by_role("button", name="编辑报价 smoke-a").click()
        page.locator('[name=amount]').fill("0.012")
        expect(page.locator("#quotePreview")).to_contain_text("前端反算付款倍率：0.60 × / 6.0 折")
        page.locator("#saveQuoteButton").click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("原始付款倍率 0.80 × / 8.0 折")
        assert fixture.preview_requests == []
        assert fixture.quotes[-1]["amount"] == "0.012" and fixture.quotes[-1]["discount_rate"] == "0.80"
        assert fixture.quotes[-1]["pricing_unit_id"] == "cny_ktoken"

        # Both CSV and XLSX export use normalized filters, selections and model IDs.
        page.locator("#exportFormat").select_option("csv")
        with page.expect_download():
            page.locator("#exportFilteredButton").click()
        page.locator("#selectAll").check()
        with page.expect_download():
            page.locator("#exportSelectedButton").click()
        page.locator("#exportFormat").select_option("xlsx")
        with page.expect_download():
            page.locator("#quoteRows tr").first.get_by_role("button", name="导出此模型 smoke-a").click()
        assert fixture.exports[0]["format"] == ["csv"]
        assert "ids" in fixture.exports[1]
        assert fixture.exports[2]["model_id"] == ["smoke-a"]
        assert fixture.exports[2]["format"] == ["xlsx"]
        page.on("dialog", lambda dialog: dialog.accept())
        page.locator("#quoteRows tr").last.get_by_role("button", name="删除报价 smoke-a").click()
        expect(page.locator("#quoteRows tr")).to_have_count(1)

        page.locator("#pricingConversionButton").click()
        # Picker and clipboard attachment UI; parsing only produces a review.
        page.locator("#filePicker").set_input_files({"name": "pricing.xlsx", "mimeType":
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": b"mock-workbook"})
        page.locator("#promptInput").evaluate("""el => {
            const dt = new DataTransfer(); dt.items.add(new File(["paste"], "clipboard.xlsx"));
            el.dispatchEvent(new ClipboardEvent("paste", {clipboardData: dt, bubbles: true, cancelable: true}));
        }""")
        expect(page.locator("#fileChips .file-chip")).to_have_count(2)
        page.locator("#promptInput").fill("解析明示官网价格，等待审核")
        page.locator("#sendButton").click()
        expect(page.locator("#reviewDialog")).to_be_visible(timeout=15000)
        assert fixture.starts[-1]["conversation_id"] == "saved-session"
        assert len(fixture.quotes) == 1
        page.locator(".bundle-details summary").click()
        page.locator("#reviewBundle").fill('"invalid bundle"')
        page.locator("#applyReviewJsonButton").click()
        expect(page.locator("#reviewIssues")).to_contain_text("必须是 JSON 对象")
        page.locator("#reviewBundle").fill(json.dumps(fixture.jobs["review-1"]["bundle"]))
        page.locator("#applyReviewJsonButton").click()
        page.locator("#cancelReviewButton").click()
        expect(page.locator("#reviewDialog")).not_to_be_visible()
        assert fixture.cancelled == 1 and fixture.confirmed == 0 and len(fixture.quotes) == 1
        expect(page.locator("#inputFileCount")).to_have_text("1")

        # Reuse returned input file; correction recalculates multiplier before commit.
        page.locator(".input-files summary").click()
        page.locator("#inputFiles input").check()
        page.locator("#promptInput").fill("重新生成报价方案")
        page.locator("#sendButton").click()
        expect(page.locator("#reviewDialog")).to_be_visible(timeout=15000)
        page.get_by_role("textbox", name="新增供应商报价 1 原币金额", exact=True).fill("12")
        page.locator("#previewReviewButton").click()
        expect(page.locator("#reviewPreview")).to_contain_text("前端反算倍率 0.60 × / 6.0 折")
        assert fixture.jobs["review-2"]["revision"] == 2
        assert len(fixture.quotes) == 1
        if args.evidence_dir:
            args.evidence_dir.mkdir(parents=True, exist_ok=True)
            page.screenshot(path=str(args.evidence_dir / "workspace-review-desktop.png"), full_page=True)
        page.locator("#confirmReviewButton").click()
        expect(page.locator("#reviewIssues")).to_contain_text("已重新加载最新计算预览")
        expect(page.locator("#reviewPreview")).to_contain_text("前端反算倍率 0.48 × / 4.8 折")
        assert fixture.jobs["review-2"]["revision"] == 3 and fixture.confirmed == 0
        page.locator("#confirmReviewButton").click()
        expect(page.locator("#reviewDialog")).not_to_be_visible()
        expect(page.locator("#quoteRows tr")).to_have_count(2)
        assert fixture.confirmed == 1 and fixture.files[0]["status"] == "data"

        # A running job can be aborted before it becomes reviewable.
        page.locator("#composer").evaluate("""el => {
            const dt = new DataTransfer(); dt.items.add(new File(["drag"], "dragged.xlsx"));
            el.dispatchEvent(new DragEvent("drop", {dataTransfer: dt, bubbles: true, cancelable: true}));
        }""")
        expect(page.locator("#fileChips .file-chip")).to_have_count(1)
        page.locator("#promptInput").fill("启动后中止")
        page.locator("#sendButton").click()
        expect(page.locator("#abortButton")).to_be_visible()
        page.locator("#abortButton").click()
        expect(page.locator("#jobStrip")).not_to_be_visible()
        expect(page.locator("#conversation .run-progress")).to_have_count(0)
        assert fixture.cancelled == 2 and fixture.confirmed == 1
        expect(page.locator("#sendButton")).to_be_enabled()

        # Memory is displayed with scope and persisted by its revision.
        page.locator("#memoryButton").click()
        expect(page.locator("#memoryRules")).to_contain_text("confirmed")
        page.get_by_role("textbox", name="规则 1 内容").fill("确认后持久保留原渠道 ID")
        page.locator("#saveMemoryButton").click()
        expect(page.locator("#memoryDialog")).not_to_be_visible()
        assert fixture.memory["revision"] == 2
        assert fixture.memory["rules"][0]["scope"] == "supplier:Smoke Supplier"
        page.reload(wait_until="networkidle")
        expect(page.locator("#conversation")).to_contain_text("已保存的用户消息")
        expect(page.locator("#sessionSelect")).to_have_value("saved-session")
        page.locator("#memoryButton").click()
        expect(page.get_by_role("textbox", name="规则 1 内容")).to_have_value("确认后持久保留原渠道 ID")
        page.locator("#closeMemoryButton").click()

        # Manual catalog edits enter the same independent review view.
        page.locator("#catalogButton").click()
        expect(page.locator("#catalogContents")).to_contain_text("alias-smoke-a")
        page.locator("#newModelButton").click()
        expect(page.locator("#reviewDialog")).to_be_visible()
        expect(page.get_by_role("textbox", name="模型实体 1 唯一模型 ID", exact=True)).to_be_visible()
        page.locator("#cancelReviewButton").click()
        if page.locator("#pricingConversionButton").get_attribute("aria-pressed") == "true":
            page.locator("#pricingConversionButton").click()
        # Original USD survives; preview arithmetic is frontend-only, never a server projection.
        page.locator("#newQuoteButton").click()
        page.locator('[name=model_id]').fill("smoke-a")
        page.locator('[name=billing_item]').fill("input")
        page.locator('[name=supplier]').fill("USD Supplier")
        page.locator('[name=pricing_unit_id]').select_option("usd_ktoken")
        page.locator('[name=amount]').fill("2")
        expect(page.locator("#quotePreview")).to_contain_text("USD 2")
        expect(page.locator("#quotePreview")).to_contain_text("前端反算付款倍率：560.00 × / 5600.0 折")
        expect(page.locator("#quotePreview").locator(".display-conversion")).to_have_count(0)
        page.locator("#saveQuoteButton").click()
        expect(page.locator("#quoteRows tr")).to_have_count(3)
        expect(page.locator("#quoteRows tr").last).to_contain_text("USD 2")
        expect(page.locator("#quoteRows tr").last.locator(".display-conversion")).to_have_count(0)
        expect(page.locator("#quoteRows tr").last).not_to_contain_text("换算 CNY")

        # Currency and unit display conversions are independent, composable and exact.
        page.locator("#currencyConversionButton").click()
        expect(page.locator("#currencyConversionDialog")).to_be_visible()
        page.locator("#currencyConversionOptions").get_by_role("button", name="CNY", exact=True).click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("显示换算 CNY 14 / 千 token")
        page.locator("#unitConversionButton").click()
        page.locator("#unitConversionOptions").get_by_role("button", name="百万 token", exact=True).click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("显示换算 CNY 14000 / 百万 token")
        page.locator("#closeUnitConversionButton").click()
        expect(page.locator("#conversionSummary")).to_contain_text("CSV / XLSX")

        # An expired FX version is rejected rather than reused for display or export.
        fixture.context["fx_rates"][0]["expires_at"] = "2000-01-01T00:00:00Z"
        page.evaluate("(context) => window.dispatchEvent(new CustomEvent('tokensql:pricing-context', {detail: context}))",
                      fixture.context)
        expect(page.locator("#quoteRows tr").last).to_contain_text("USD 汇率已过期")
        del fixture.context["fx_rates"][0]["expires_at"]
        page.evaluate("(context) => window.dispatchEvent(new CustomEvent('tokensql:pricing-context', {detail: context}))",
                      fixture.context)
        expect(page.locator("#quoteRows tr").last).to_contain_text("显示换算 CNY 14000 / 百万 token")

        # Active conversions export through the pure serializer as a real XLSX with evidence.
        page.locator("#exportFormat").select_option("xlsx")
        with page.expect_download() as converted_download:
            page.locator("#exportFilteredButton").click()
        converted_book = load_workbook(BytesIO(Path(converted_download.value.path()).read_bytes()), read_only=True)
        converted_rows = list(converted_book.active.iter_rows(values_only=True))
        assert "原价" in converted_rows[0] and "换算价" in converted_rows[0] and "FX 来源" in converted_rows[0]
        assert any(row[converted_rows[0].index("换算价")] == "14000" for row in converted_rows[1:])
        assert fixture.view_exports[-1]["format"] == "xlsx"
        assert all(not isinstance(value, float) for row in fixture.view_exports[-1]["rows"] for value in row.values())
        page.locator("#exportFormat").select_option("csv")
        with page.expect_download():
            page.locator("#exportFilteredButton").click()
        assert fixture.view_exports[-1]["format"] == "csv"
        assert fixture.view_exports[-1]["rows"][0]["converted_amount_denominator"]
        assert fixture.view_exports[-1]["rows"][0]["multiplier_denominator"]

        # Per-model converted export fetches every matching quote across multiple API pages.
        original_quote_count = len(fixture.quotes)
        fixture.quotes.extend(quote(1000 + index) for index in range(501))
        with page.expect_download():
            page.locator("#quoteRows tr").first.get_by_role("button", name="导出此模型 smoke-a").click()
        assert len(fixture.view_exports[-1]["rows"]) == original_quote_count + 501
        assert {row["model_id"] for row in fixture.view_exports[-1]["rows"]} == {"smoke-a"}
        fixture.quotes = fixture.quotes[:original_quote_count]

        # Restoring one axis leaves the other active; restoring both returns the untouched view.
        page.locator("#unitConversionButton").click()
        page.locator("#unitConversionOptions").get_by_role("button", name="恢复原单位", exact=True).click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("显示换算 CNY 14 / 千 token")
        page.locator("#closeUnitConversionButton").click()
        page.locator("#currencyConversionButton").click()
        page.locator("#currencyConversionOptions").get_by_role("button", name="恢复原币", exact=True).click()
        expect(page.locator("#quoteRows .display-conversion")).to_have_count(0)

        # Without either display conversion, only the original price is rendered.
        expect(page.locator("#priceView")).to_have_count(0)
        expect(page.locator("#quoteRows .normalized-price")).to_have_count(0)
        expect(page.locator("#quoteRows tr").last).to_contain_text("USD 2")
        expect(page.locator("#quoteRows tr").last).to_contain_text("USD / 千 token")
        expect(page.locator("#quoteRows tr").last).not_to_contain_text("换算 CNY")

        # Selecting HKD adds only the frontend HKD conversion, never an extra CNY/1M view.
        page.locator("#currencyConversionButton").click()
        page.locator("#currencyConversionOptions").get_by_role("button", name="HKD", exact=True).click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("显示换算 HKD")
        expect(page.locator("#quoteRows tr").last).not_to_contain_text("换算 CNY")
        expect(page.locator("#quoteRows .normalized-price")).to_have_count(0)
        page.locator("#currencyConversionButton").click()
        page.locator("#currencyConversionOptions").get_by_role("button", name="恢复原币", exact=True).click()

        video = quote(101)
        video.update(pricing_unit_id="cny_second", unit="second", unit_quantity="1", amount="1",
                     supplier_model_id="video-channel", discount_rate=None)
        fixture.quotes.append(video)
        page.locator("#refreshButton").click()
        expect(page.locator("#quoteRows tr")).to_have_count(4)
        expect(page.locator("#quoteRows tr").last).to_contain_text("CNY 1")
        expect(page.locator("#quoteRows tr").last).to_contain_text("CNY / 秒")
        expect(page.locator("#quoteRows tr").last.locator(".display-conversion")).to_have_count(0)

        # Cross-dimension display conversion blocks missing factors, then accepts a scoped temporary factor.
        page.locator("#unitConversionButton").click()
        page.locator("#unitConversionOptions").get_by_role("button", name="百万 token", exact=True).click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("未换算：缺少该模型的 tokens_per_second 合法倍率")
        page.locator("#temporaryFactorModel").fill("smoke-a")
        page.locator("#temporaryFactorBillingItem").fill("input")
        page.locator("#temporaryTokensPerSecond").fill("100000")
        page.locator("#temporaryFactorSource").fill("Synthetic browser-only measurement")
        page.locator("#applyTemporaryFactorsButton").click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("工程估算 CNY 10 / 百万 token")
        assert page.evaluate("window.TokenSQLConversion.state.customFactors.has('smoke-a\\u0000input')")
        assert fixture.context["token_equivalences"] == []
        page.locator("#closeUnitConversionButton").click()
        page.locator("#resetConversionButton").click()
        expect(page.locator("#quoteRows .display-conversion")).to_have_count(0)

        page.locator("#contextButton").click()
        expect(page.locator("#contextMeta")).to_contain_text("参数版本 1")
        factor = {"model_id": "smoke-a", "billing_item": "input", "unit": "second", "tokens_per_unit": "100000",
                  "source": "Synthetic measurement", "source_ref": "test only", "observed_at": "2026-10-16T00:00:00Z",
                  "assumptions": "Explicit test factor, never a production default"}
        page.locator("#contextFactors").fill(json.dumps([factor]))
        page.locator("#saveContextButton").click()
        expect(page.locator("#contextMeta")).to_contain_text("参数版本 2")
        page.locator("#refreshFxButton").click()
        expect(page.locator("#contextMeta")).to_contain_text("参数版本 3")
        expect(page.locator("#contextFx")).to_have_value(json.dumps(fixture.context["fx_rates"], indent=2))
        page.locator("#closeContextButton").click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("CNY 1")
        expect(page.locator("#quoteRows tr").last.locator(".display-conversion")).to_have_count(0)
        page.locator("#pricingConversionButton").click()
        expect(page.locator("#quoteRows tr").last).to_contain_text("工程估算倍率")
        assert fixture.context["token_equivalences"][0]["source"] == "Synthetic measurement"

        # Every calculation is a frontend projection, including both raw fields and null/zero inputs.
        mutation_count = len(fixture.mutations)
        originals = deepcopy(fixture.quotes)
        computed = page.evaluate("""() => {
          const c = window.TokenSQLConversion;
          const official = {model_id:"smoke-a",billing_item:"input",amount:"20",pricing_unit_id:"cny_mtoken",revision:7,source_ref:"official-test"};
          const base = {model_id:"smoke-a",billing_item:"input",official_price:official};
          const read = (quote) => { const r=c.computePricing(quote); return {
            status:r.status,reason:r.reason,discount:r.calculatedDiscount,derived:r.derivedAmount,
            issue:r.discountIssue,estimated:r.estimated,unit:r.officialUnit?.id,
            numerator:r.discountNumerator,denominator:r.discountDenominator,
            fold:r.calculatedDiscountFold,rounded:r.discountRounded}; };
          return {
            same:read({...base,amount:"10",pricing_unit_id:"cny_mtoken"}),
            scale:read({...base,amount:"0.012",pricing_unit_id:"cny_ktoken"}),
            fx:read({...base,amount:"2",pricing_unit_id:"usd_ktoken"}),
            discount:read({...base,amount:null,discount_rate:"0.8",pricing_unit_id:null}),
            zeroDiscount:read({...base,amount:null,discount_rate:"0",pricing_unit_id:null}),
            zeroAmount:read({...base,amount:"0",pricing_unit_id:"cny_mtoken"}),
            both:read({...base,amount:"10",discount_rate:"0.8",pricing_unit_id:"cny_mtoken"}),
            partial:read({...base,amount:"10",currency:"CNY",unit:"token",unit_quantity:null}),
            missing:read({model_id:"missing-official",amount:null,discount_rate:"0.8"}),
            bothUnknownCurrency:read({...base,official_price:{...official,pricing_unit_id:null,currency:null,unit:"token",unit_quantity:"1000000"},amount:"100",pricing_unit_id:null,currency:null,unit:"token",unit_quantity:"1000000"}),
            supplierUnknownCurrency:read({...base,amount:"100",pricing_unit_id:null,currency:null,unit:"token",unit_quantity:"1000000"}),
            officialUnknownCurrency:read({...base,official_price:{...official,pricing_unit_id:null,currency:null,unit:"token",unit_quantity:"1000000"},amount:"100",pricing_unit_id:"cny_mtoken"}),
            zeroOfficial:read({...base,official_price:{...official,amount:"0"},amount:"1",pricing_unit_id:"cny_mtoken"}),
            estimated:read({...base,amount:"1",pricing_unit_id:"cny_second"}),
            half005:read({...base,amount:"0.1",pricing_unit_id:"cny_mtoken"}),
            half155:read({...base,amount:"3.1",pricing_unit_id:"cny_mtoken"}),
            below1549:read({...base,amount:"3.098",pricing_unit_id:"cny_mtoken"}),
            repeating:read({...base,official_price:{...official,amount:"3"},amount:"1",pricing_unit_id:"cny_mtoken"}),
          };
        }""")
        assert computed["same"]["discount"] == "0.50"
        assert computed["scale"]["discount"] == "0.60"
        assert computed["fx"]["discount"] == "700.00"
        assert computed["discount"]["derived"] == "16" and computed["discount"]["unit"] == "cny_mtoken"
        assert computed["zeroDiscount"]["derived"] == "0" and computed["zeroAmount"]["discount"] == "0.00"
        assert computed["both"]["discount"] == "0.50" and computed["both"]["derived"] == "16"
        assert computed["partial"]["issue"] and computed["missing"]["reason"] == "缺失官网价"
        assert "币种未提供" in computed["bothUnknownCurrency"]["reason"]
        assert "币种未提供" in computed["supplierUnknownCurrency"]["issue"]
        assert "币种未提供" in computed["officialUnknownCurrency"]["reason"]
        assert all(computed[name].get("discount") is None for name in ("bothUnknownCurrency", "supplierUnknownCurrency", "officialUnknownCurrency"))
        assert "为 0" in computed["zeroOfficial"]["issue"]
        assert computed["estimated"]["estimated"] is True and computed["estimated"]["discount"] == "0.50"
        assert [(computed[key]["discount"], computed[key]["fold"]) for key in ("half005", "half155", "below1549", "repeating")] == [
            ("0.01", "0.1"), ("0.16", "1.6"), ("0.15", "1.5"), ("0.33", "3.3")]
        assert all(computed[key]["rounded"] for key in ("half005", "half155", "below1549", "repeating"))
        assert (computed["half155"]["numerator"], computed["half155"]["denominator"]) == ("31", "200")
        assert fixture.quotes == originals and len(fixture.mutations) == mutation_count

        # No supplier, billing item or unit is required, and zero is a real raw discount value.
        page.locator("#newQuoteButton").click()
        page.locator('[name=model_id]').fill("smoke-a")
        expect(page.locator('[name=billing_item]')).to_have_value("")
        expect(page.locator('[name=pricing_unit_id]')).to_have_value("")
        page.locator('[name=discount_rate]').fill("0")
        page.locator("#saveQuoteButton").click()
        expect(page.locator("#quoteRows tr")).to_have_count(5)
        raw = fixture.quotes[-1]
        assert raw["amount"] is None and raw["discount_rate"] == "0.00"
        assert raw["supplier"] is None and raw["billing_item"] is None and raw["pricing_unit_id"] is None
        expect(page.locator("#quoteRows tr").last).to_contain_text("原始付款倍率 0.00 × / 0.0 折")
        expect(page.locator("#quoteRows tr").last).to_contain_text("缺失官网价")

        # Display-only discount and derived-price columns export, retaining null originals.
        page.locator("#exportFormat").select_option("xlsx")
        with page.expect_download() as pricing_download:
            page.locator("#exportFilteredButton").click()
        pricing_book = load_workbook(BytesIO(Path(pricing_download.value.path()).read_bytes()), read_only=True)
        pricing_rows = list(pricing_book.active.iter_rows(values_only=True))
        assert "原始付款倍率（两位小数）" in pricing_rows[0] and "前端反算付款倍率（两位小数）" in pricing_rows[0]
        assert fixture.view_exports[-1]["rows"][-1]["original_amount"] is None
        assert fixture.view_exports[-1]["rows"][-1]["original_discount_rate"] == "0.00"
        assert fixture.view_exports[-1]["rows"][-1]["frontend_pricing_reason"] == "缺失官网价"

        # Quick official import is a brand-new zero-file session, and still awaits human review.
        raw_before_lookup = deepcopy(fixture.quotes)
        page.locator("#quoteRows tr").last.locator(".official-lookup").first.click()
        expect(page.locator("#reviewDialog")).to_be_visible(timeout=15000)
        assert fixture.lookup_requests[-1] == {"model_id":"smoke-a","billing_item":None,"prompt":""}
        assert page.locator("#sessionSelect").input_value().startswith("fresh-official-session-")
        assert fixture.quotes == raw_before_lookup
        page.locator("#cancelReviewButton").click()
        expect(page.locator("#reviewDialog")).not_to_be_visible()

        # Failed drafts immediately render the assistant explanation before a separate task error.
        fixture.next_failure = "Tokens 未标明 K 或 M；vendor 未提供，保留为空，不猜测。"
        page.locator("#promptInput").fill("模拟校验失败但模型有解释")
        page.locator("#sendButton").click()
        expect(page.locator("#conversation")).to_contain_text("Tokens 未标明 K 或 M")
        expect(page.locator("#conversation")).to_contain_text("任务未完成：proposal validation failed")
        texts = page.locator("#conversation .message-body p").all_text_contents()
        assert next(i for i,t in enumerate(texts) if "Tokens 未标明" in t) < next(i for i,t in enumerate(texts) if "任务未完成：proposal" in t)
        expect(page.locator("#abortButton")).to_have_text("取消并退回文件")
        page.locator("#abortButton").click()
        expect(page.locator("#jobStrip")).not_to_be_visible()

        # Canonical null means no answer for this turn; never re-label an older answer.
        fixture.next_failure = {"old_reply":"OBSOLETE_OLD_TURN_NOT_THIS_REQUEST"}
        page.locator("#promptInput").fill("模拟HTTP错误且只有旧会话历史")
        page.locator("#sendButton").click()
        expect(page.locator("#conversation .message-body p").last).to_contain_text("任务未完成：proposal validation failed")
        expect(page.locator("#conversation")).not_to_contain_text("OBSOLETE_OLD_TURN_NOT_THIS_REQUEST")
        page.locator("#abortButton").click()
        expect(page.locator("#jobStrip")).not_to_be_visible()

        # Interrupted polling can still recover an answer anchored after the pre-request history.
        fixture.next_history_reply = "网络中断前的部分真实说明：该行只有付款倍率。"
        fixture.fail_poll_count = 2
        page.locator("#promptInput").fill("模拟轮询网络中断")
        page.locator("#sendButton").click()
        expect(page.locator("#conversation")).to_contain_text("网络中断前的部分真实说明")
        expect(page.locator("#conversation")).to_contain_text("任务未完成：Failed to fetch")
        page.locator("#abortButton").click()
        expect(page.locator("#jobStrip")).not_to_be_visible()
        fixture.history_reply = None

        # A slow history response must not overwrite a newer selected session.
        fixture.sessions.extend([{"id":"stale-a","title":"Slow session A"},{"id":"fresh-b","title":"Fresh session B"}])
        stale_page = browser.new_page(viewport={"width":1200,"height":800})
        held = []
        def history_route(route):
            path = urlparse(route.request.url).path
            if path.endswith("/sessions/stale-a/messages"):
                held.append(route)
            elif path.endswith("/sessions/fresh-b/messages"):
                route.fulfill(json={"messages":[{"role":"assistant","text":"CURRENT_SESSION_B"}]})
            else: fixture.route(route)
        stale_page.route("**/api/workspace/**", history_route)
        stale_page.goto(url + "/workspace", wait_until="networkidle")
        with stale_page.expect_request("**/sessions/stale-a/messages"):
            stale_page.locator("#sessionSelect").select_option("stale-a")
        stale_page.locator("#sessionSelect").select_option("fresh-b")
        expect(stale_page.locator("#conversation")).to_contain_text("CURRENT_SESSION_B")
        assert len(held) == 1
        held[0].fulfill(json={"messages":[{"role":"assistant","text":"OBSOLETE_SESSION_A"}]})
        stale_page.wait_for_timeout(100)
        expect(stale_page.locator("#conversation")).not_to_contain_text("OBSOLETE_SESSION_A")
        stale_page.close()
        if args.evidence_dir:
            page.screenshot(path=str(args.evidence_dir / "workspace-smoke-desktop.png"), full_page=True)

        page.set_viewport_size({"width": 390, "height": 844})
        assert len(page.locator(".content-grid").evaluate("el => getComputedStyle(el).gridTemplateColumns.split(' ')")) == 1
        assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
        page.locator("#catalogButton").click()
        page.locator("#newOfficialButton").click()
        expect(page.locator("#reviewDialog")).to_be_visible()
        assert page.locator("#reviewDialog").evaluate("el => el.getBoundingClientRect().width") <= 390
        manual_official = json.loads(page.locator("#reviewBundle").input_value())["official_prices"][0]
        assert manual_official["explicit_official"] is None and manual_official["pricing_unit_id"] is None
        checkbox = page.get_by_role("checkbox", name="官网计费项 1 文件 / 来源明确标注官网价", exact=True)
        checkbox.check(); checkbox.uncheck()
        assert json.loads(page.locator("#reviewBundle").input_value())["official_prices"][0]["explicit_official"] is None
        unit_select = page.get_by_role("combobox", name="官网计费项 1 计价单位", exact=True)
        unit_select.select_option("usd_ktoken")
        edited_official = json.loads(page.locator("#reviewBundle").input_value())["official_prices"][0]
        assert edited_official["currency"] == "USD"
        assert all(edited_official[key] is None for key in ("unit", "unit_quantity"))
        if args.evidence_dir:
            page.screenshot(path=str(args.evidence_dir / "workspace-review-mobile.png"), full_page=True)
        page.locator("#cancelReviewButton").click()
        if args.evidence_dir:
            page.screenshot(path=str(args.evidence_dir / "workspace-smoke-mobile.png"), full_page=True)

        browser.close()
    if server:
        server.shutdown()
    assert not page_errors, page_errors
    expected_conflicts = [message for message in console_errors if "409 (Conflict)" in message]
    assert len(expected_conflicts) == 1, console_errors
    assert not [message for message in console_errors if "409 (Conflict)" not in message and "net::ERR_FAILED" not in message], console_errors
    print("workspace browser smoke passed: raw nullable fields/zero, frontend-only official alignment/discount/derived amount, immediate failure+history fallback, independent official-lookup review, canonical units, exact dual-axis conversion, CSV/real XLSX export, multi-page model export, CRUD/review/session/memory, desktop/mobile")


if __name__ == "__main__":
    main()
