#!/usr/bin/env python3
"""Synthetic real-API/browser acceptance for scopes, official scenarios and display-only prices."""
from __future__ import annotations

import argparse
import csv
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

PROJECT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8896)
    parser.add_argument("--chromium", default="/home/user/.cache/ms-playwright/chromium-1243/chrome-linux64/chrome")
    parser.add_argument("--evidence-dir", type=Path, default=Path("/tmp/tokensql-conditions-real-evidence"))
    args = parser.parse_args()
    with socket.socket() as preflight:
        if preflight.connect_ex(("127.0.0.1", args.port)) == 0:
            raise RuntimeError("QA port is in use; will not touch another process")
    root = Path(tempfile.mkdtemp(prefix="tokensql-conditions-real-"))
    args.evidence_dir.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    for name in ("workspace", "agent", "input", "progress", "data"):
        env["TOKENSQL_" + name.upper() + "_DIR"] = str(root / name)
    env.update(TOKENSQL_TEST_MODE="true", TOKENSQL_FX_AUTO_REFRESH="false",
               TOKENSQL_OPENCODE_URL="http://127.0.0.1:1", PYTHONDONTWRITEBYTECODE="1",
               PYTHONPATH=str(PROJECT / "src"))
    log = (root / "server.log").open("w")
    process = subprocess.Popen([sys.executable, "-m", "uvicorn", "tokensql.workspace.app:app",
                                "--host", "127.0.0.1", "--port", str(args.port)],
                               cwd=PROJECT, env=env, stdout=log, stderr=subprocess.STDOUT)
    base = "http://127.0.0.1:" + str(args.port)

    def request(path, body=None, method=None):
        payload = None if body is None else json.dumps(body).encode()
        req = Request(base + path, data=payload, method=method or ("POST" if body is not None else "GET"),
                      headers={"Content-Type": "application/json"})
        try:
            with urlopen(req, timeout=20) as response:
                return json.load(response) if response.status != 204 else None
        except HTTPError as error:
            raise AssertionError((error.code, error.read().decode())) from error

    def snapshot():
        result = {key: request("/api/workspace/" + key) for key in
                  ("quotes?limit=500", "catalog", "pricing-context", "agent/memory")}
        with sqlite3.connect(root / "workspace" / "quotes.sqlite3") as connection:
            result["audit_count"] = connection.execute("SELECT COUNT(*) FROM pricing_audit").fetchone()[0]
        return result

    def official(model, metric, amount, conditions=None, item=None):
        return dict(model_id=model, billing_metric=metric, billing_item=item or metric, amount=amount,
                    pricing_unit_id="usd_mtoken", pricing_conditions=conditions or {}, explicit_official=True,
                    source="Synthetic official fixture", source_ref="https://example.invalid/synthetic/" + model)

    try:
        for _ in range(100):
            if process.poll() is not None:
                raise RuntimeError((root / "server.log").read_text())
            try:
                if request("/health").get("isolated_test"):
                    break
            except Exception:
                time.sleep(0.1)
        else:
            raise RuntimeError("Isolated QA service failed to start")
        prices = [official("qa-opus", "input", "5"), official("qa-opus", "output", "25"),
                  official("qa-opus", "cache_write", "6.25", {"cache_ttl_seconds": 300}, "cache_write_5m"),
                  official("qa-opus", "cache_write", "10", {"cache_ttl_seconds": 3600}, "cache_write_1h"),
                  official("qa-opus", "cache_read", "0.5")]
        prices += [official("qa-region", "input", "1", {"market": "cn"}),
                   official("qa-region", "input", "2", {"market": "us"}),
                   official("qa-rules", "input", "5"), official("qa-rules", "output", "2"),
                   official("qa-range", "input", "2", {"input_tokens": {"min": 0, "max": 100}})]
        quotes = [
            dict(model_id="qa-opus", supplier="Whole", discount_scope="model", discount_rate="0.70"),
            dict(model_id="qa-opus", supplier="Unknown", discount_rate="0.70"),
            dict(model_id="qa-opus", supplier="Priced", billing_metric="input", amount="3.5",
                 discount_rate="0.65", pricing_unit_id="usd_mtoken"),
            dict(model_id="qa-missing", supplier="Missing", discount_scope="model", discount_rate="0.70"),
            dict(model_id="qa-region", supplier="Ambiguous", billing_metric="input", discount_rate="0.70"),
            dict(model_id="qa-rules", supplier="Rules", supplier_model_id="same-channel",
                 discount_scope="model", discount_rate="0.70"),
            dict(model_id="qa-rules", supplier="Rules", supplier_model_id="same-channel",
                 billing_metric="input", discount_rate="0.20"),
            dict(model_id="qa-rules", supplier="Rules", supplier_model_id="same-channel",
                 billing_metric="input", discount_rate="0.30"),
            dict(model_id="qa-opus", supplier="CompleteUnits", discount_scope="model", discount_rate="0.70", pricing_unit_id="usd_mtoken"),
            dict(model_id="qa-opus", supplier="NullUnits", discount_scope="model", discount_rate="0.70", currency=None),
            dict(model_id="qa-opus", supplier="PartialUnits", discount_scope="model", discount_rate="0.70", currency="CNY", unit="token"),
            dict(model_id="qa-range", supplier="Range", billing_metric="input", discount_rate="0.70",
                 pricing_conditions={"input_tokens": {"min": 20, "max": 50}}),
        ]
        request("/api/workspace/catalog", {"bundle": {"models": [{"model_id": model} for model in
            ("qa-opus", "qa-region", "qa-rules", "qa-missing", "qa-range")],
            "official_prices": prices, "create": quotes}})
        before = snapshot()
        items = before["quotes?limit=500"]["items"]
        by_supplier = {item["supplier"]: item for item in items}
        whole = by_supplier["Whole"]
        assert len(whole["official_prices"]) == 5 and whole["official_price"] is None
        assert whole["official_match"]["status"] == "multiple"
        assert by_supplier["Unknown"]["official_match"]["status"] == "needs_scope"
        assert by_supplier["Ambiguous"]["official_match"]["status"] == "ambiguous"
        assert by_supplier["Missing"]["official_match"]["status"] == "missing"
        assert by_supplier["Range"]["official_match"]["status"] == "matched"
        hour = next(price for price in whole["official_prices"] if price["pricing_conditions"].get("cache_ttl_seconds") == 3600)
        errors = []
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True, executable_path=args.chromium)
            page = browser.new_page(viewport={"width": 1600, "height": 1000}, accept_downloads=True)
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(base + "/workspace", wait_until="networkidle")
            expect(page.locator("#quoteRows > tr")).to_have_count(12)

            def row(item):
                return page.locator('#quoteRows > tr[data-id="' + str(item["id"]) + '"]')

            for supplier in ("CompleteUnits", "NullUnits", "PartialUnits"):
                expect(row(by_supplier[supplier])).to_contain_text("源折扣 7.0折")
                assert "源折扣 0.70折" not in row(by_supplier[supplier]).inner_text()
            main = row(whole)
            expect(main).to_contain_text("源折扣 7.0折")
            expect(main).to_contain_text("已有 5 项官网价")
            assert "0.70 ×" not in page.locator("#quoteRows").inner_text()
            assert not main.locator(".official-lookup").count()
            assert not row(by_supplier["Unknown"]).locator(".official-lookup").count()
            assert not row(by_supplier["Ambiguous"]).locator(".official-lookup").count()
            assert row(by_supplier["Missing"]).locator(".official-lookup").count()
            expect(row(by_supplier["Unknown"])).to_contain_text("适用于整模型")
            expect(row(by_supplier["Ambiguous"])).to_contain_text("适用条件")
            page.locator("#pricingConversionButton").click()
            pricing = main.locator("td").nth(7)
            pricing.locator(".official-details > summary").click()
            expect(pricing.locator(".official-detail-item")).to_have_count(5)
            for amount in ("3.5", "17.5", "4.375", "7", "0.35"):
                expect(pricing).to_contain_text("折后价 USD " + amount + " /")
            expect(pricing).to_contain_text("缓存有效期（秒）：3600")
            expect(row(by_supplier["Priced"])).to_contain_text("源折扣 6.5折")
            expect(row(by_supplier["Priced"])).to_contain_text("按官网反算折扣 7.0折")
            rules_model = next(item for item in items if item["supplier"] == "Rules" and item["discount_scope"] == "model")
            rules = row(rules_model)
            rules.locator("td").nth(7).locator(".official-details > summary").click()
            expect(rules).to_contain_text("同等优先级的源折扣冲突")
            expect(rules).to_contain_text("折后价 USD 1.4 /")
            pricing.locator("select").select_option(hour["official_price_id"])
            expect(main.locator("td").nth(7).locator(".official-detail-item")).to_have_count(1)
            expect(main.locator("td").nth(7)).to_contain_text("折后价 USD 7 /")

            def export(format):
                page.locator("#exportFormat").select_option(format)
                with page.expect_download() as download:
                    page.locator("#exportFilteredButton").click()
                data = Path(download.value.path()).read_bytes()
                (args.evidence_dir / ("conditions-export." + format)).write_bytes(data)
                if format == "csv":
                    return list(csv.DictReader(StringIO(data.decode("utf-8-sig"))))
                book = load_workbook(BytesIO(data), read_only=True)
                values = list(book.active.iter_rows(values_only=True))
                return [dict(zip(values[0], values_row)) for values_row in values[1:]]

            selected_export = export("csv")
            chosen = [item for item in selected_export if str(item["源报价 ID"]) == str(whole["id"])]
            assert len(chosen) == 1 and chosen[0]["前端折后价"] == "7"
            assert "3600" in chosen[0]["官网适用条件"]
            main.locator("td").nth(7).locator("select").select_option("")
            xlsx = export("xlsx")
            all_prices = [item for item in xlsx if str(item["源报价 ID"]) == str(whole["id"])]
            assert {item["前端折后价"] for item in all_prices} == {"3.5", "17.5", "4.375", "7", "0.35"}
            assert {item["原始存储折扣值（两位小数）"] for item in all_prices} == {"0.70"}
            assert {item["源折扣（折，一位小数）"] for item in all_prices} == {"7.0"}
            priced = next(item for item in xlsx if str(item["源报价 ID"]) == str(by_supplier["Priced"]["id"]))
            assert priced["源折扣（折，一位小数）"] == "6.5" and priced["按官网反算折扣（折，一位小数）"] == "7.0"
            page.screenshot(path=str(args.evidence_dir / "desktop.png"), full_page=True)
            page.set_viewport_size({"width": 390, "height": 844})
            assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
            page.screenshot(path=str(args.evidence_dir / "mobile.png"), full_page=True)
            page.set_viewport_size({"width": 1600, "height": 1000})
            page.reload(wait_until="networkidle")
            expect(row(whole)).to_contain_text("已有 5 项官网价")
            assert not row(whole).locator(".official-lookup").count()
            # A failed display module must never relabel the raw 0.70 ratio as 0.70-fold.
            fallback_page = browser.new_page()
            fallback_page.route("**/static/conversion.js", lambda route: route.abort())
            fallback_page.goto(base + "/workspace", wait_until="networkidle")
            fallback_row = fallback_page.locator('#quoteRows > tr[data-id="' + str(by_supplier["NullUnits"]["id"]) + '"]')
            expect(fallback_row).to_contain_text("源折扣暂不可显示")
            assert "0.70折" not in fallback_row.inner_text()
            fallback_page.close()
            unsupported = page.evaluate("""() => {
              const units = window.TokenSQL.state.pricingUnits.items;
              const target = units.find(unit => unit.id === "usd_mtoken");
              return ["usd_mtoken_hour", "usd_request", "usd_account_month"].map(id =>
                window.TokenSQLConversion.convertTo({model_id:"qa-opus", amount:"5", pricing_unit_id:id}, target).status);
            }""")
            assert unsupported == ["unavailable", "unavailable", "unavailable"]

            # Existing quote preview must use the backend range projection and make no business writes.
            range_row = row(by_supplier["Range"])
            range_row.get_by_role("button", name="编辑报价 qa-range").click()
            expect(page.locator("#quotePreview")).to_contain_text("折后价 USD 1.4 /")
            assert "联网补" not in page.locator("#quotePreview").inner_text()
            page.locator("[data-close-editor]").first.click()
            assert before == snapshot(), "Display, selection, preview or exports changed business data"

            # Normal edit persists typed scope and independent original text; amount/rate stay intact.
            row(by_supplier["Unknown"]).get_by_role("button", name="编辑报价 qa-opus").click()
            editor = page.locator("#quoteConditionsEditor .conditions-editor")
            editor.locator("> summary").click()
            editor.locator('[data-condition-field="discount_scope"]').select_option("items")
            editor.locator(".discount-item-choices input").first.check()
            editor.locator(".condition-extra > summary").click()
            editor.locator('[data-condition-field="market"]').fill("cn")
            editor.locator('[data-condition-field="time_band"]').select_option("busy")
            page.locator('#quoteForm [name="conditions"]').fill('{"source_note":"synthetic retained text"}')
            page.locator("#saveQuoteButton").click()
            expect(page.locator("#editorDialog")).not_to_be_visible()
            changed = request("/api/workspace/quotes?limit=500")["items"]
            updated = next(item for item in changed if item["id"] == by_supplier["Unknown"]["id"])
            assert updated["discount_scope"] == "items" and len(updated["discount_items"]) == 1
            assert updated["amount"] is None and updated["discount_rate"] == "0.70"
            assert updated["conditions"]["source_note"] == "synthetic retained text"
            assert updated["pricing_conditions"]["market"] == "cn" and updated["pricing_conditions"]["time_band"] == "busy"

            # Resolve one conflicting synthetic rule and prove item discount overrides, never multiplies.
            conflicting = next(item for item in items if item["model_id"] == "qa-rules" and item["discount_rate"] == "0.30")
            request("/api/workspace/quotes/" + str(conflicting["id"]),
                    {"expected_revision": conflicting["revision"], "discount_rate": "0.20"}, "PATCH")
            page.locator("#refreshButton").click()
            page.locator("#pricingConversionButton").click()
            row(rules_model).locator("td").nth(7).locator(".official-details > summary").click()
            expect(row(rules_model)).to_contain_text("采用指定计费项源折扣 2.0折")
            expect(row(rules_model)).to_contain_text("折后价 USD 1 /")
            assert "折后价 USD 0.7 /" not in row(rules_model).inner_text()

            # Official catalog editing must retain the stable price ID and its revision.
            page.locator("#catalogButton").click()
            catalog = page.locator(".catalog-model").filter(has=page.get_by_role("heading", name="未提供模型商 · qa-range", exact=True))
            catalog.get_by_role("button", name="编辑官网计费项").click()
            official_editor = page.locator('#reviewSections .conditions-editor').last
            official_editor.locator("> summary").click()
            official_editor.locator(".condition-extra > summary").click()
            official_editor.locator('[data-condition-field="market"]').fill("cn")
            page.locator("#previewReviewButton").click()
            expect(page.locator("#reviewIssues")).to_have_text("")
            proposal = page.locator("#reviewBundle").input_value()
            parsed = json.loads(proposal)["official_prices"][0]
            original_range = next(item for item in before["catalog"]["official_prices"] if item["model_id"] == "qa-range")
            assert parsed["official_price_id"] == original_range["official_price_id"]
            assert parsed["expected_revision"] == original_range["revision"]
            assert parsed["pricing_conditions"]["market"] == "cn"
            page.locator("#cancelReviewButton").click()
            expect(page.locator("#reviewDialog")).not_to_be_visible()
            browser.close()
        assert not errors, errors
        result = dict(passed=True, kind="synthetic_real_api_browser", production_accessed=False,
                      real_model_called=False, original_quotes=12, five_official_prices=True,
                      range_preview_matches_backend=True, all_nullable_units_fold_correct=True, failed_module_no_mislabel=True, new_dimensions_no_token_conversion=True, display_export_preview_zero_business_writes=True,
                      source_fold="7.0", effective_rule_conflict_verified=True, effective_rule_override_verified=True,
                      explicit_scope_edit=True, official_stable_id_revision_edit=True,
                      evidence_dir=str(args.evidence_dir), database=str(root / "workspace" / "quotes.sqlite3"))
        (args.evidence_dir / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps(result, ensure_ascii=False))
    finally:
        process.terminate()
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait()
        log.close()


if __name__ == "__main__":
    main()
