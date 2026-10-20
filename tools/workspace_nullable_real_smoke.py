#!/usr/bin/env python3
"""Isolated real-API browser proof; synthetic data only, never calls OpenCode."""
from __future__ import annotations
import argparse
from io import BytesIO
import json
import os
from pathlib import Path
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

PROJECT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8789)
    parser.add_argument("--chromium", default="/home/user/.cache/ms-playwright/chromium-1243/chrome-linux64/chrome")
    parser.add_argument("--evidence-dir", type=Path, default=Path("/tmp/tokensql-nullable-real-evidence"))
    args = parser.parse_args()
    with socket.socket() as preflight:
        if preflight.connect_ex(("127.0.0.1", args.port)) == 0:
            raise RuntimeError("QA port already in use; will not touch another process")
    root = Path(tempfile.mkdtemp(prefix="tokensql-nullable-real-"))
    args.evidence_dir.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    for name in ("workspace", "agent", "input", "progress", "data"):
        env["TOKENSQL_" + (name.upper() + "_DIR")] = str(root / name)
    env.update(TOKENSQL_TEST_MODE="true", TOKENSQL_FX_AUTO_REFRESH="false",
               TOKENSQL_OPENCODE_URL="http://127.0.0.1:1", PYTHONDONTWRITEBYTECODE="1")
    log = (root / "server.log").open("w")
    process = subprocess.Popen([sys.executable, "-m", "uvicorn", "tokensql.workspace.app:app",
                               "--host", "127.0.0.1", "--port", str(args.port)],
                              cwd=PROJECT, env=env, stdout=log, stderr=subprocess.STDOUT)
    base = "http://127.0.0.1:" + str(args.port)

    def request(path, body=None, method=None):
        payload = None if body is None else json.dumps(body).encode()
        req = Request(base + path, data=payload, method=method or ("POST" if body is not None else "GET"),
                      headers={"Content-Type": "application/json"})
        try:
            with urlopen(req, timeout=15) as response:
                return json.load(response)
        except HTTPError as error:
            raise AssertionError((error.code, error.read().decode())) from error

    def snapshot():
        result = {key: request("/api/workspace/" + key) for key in
                  ("quotes?limit=500", "catalog", "pricing-context", "agent/memory")}
        with sqlite3.connect(root / "workspace" / "quotes.sqlite3") as connection:
            result["audit_count"] = connection.execute("SELECT COUNT(*) FROM pricing_audit").fetchone()[0]
        return result

    try:
        for _ in range(100):
            if process.poll() is not None: raise RuntimeError((root / "server.log").read_text())
            try:
                health = request("/health")
                if health.get("isolated_test"): break
            except Exception: time.sleep(0.1)
        else: raise RuntimeError("isolated test service failed to start")
        proposal = {
            "models": [{"model_id": name} for name in ("qa-null-a", "qa-null-b", "qa-null-c")],
            "official_prices": [{"model_id": "qa-null-a", "billing_item": "input", "amount": "20",
                                 "pricing_unit_id": "cny_mtoken", "explicit_official": True,
                                 "source": "Synthetic official fixture", "source_ref": "Test!B2"}],
            "create": [
                {"model_id": "qa-null-a", "billing_item": "input", "discount_rate": "0.8"},
                {"model_id": "qa-null-a", "billing_item": "input", "supplier": "Synthetic both",
                 "amount": "0.012", "discount_rate": "0.8", "pricing_unit_id": "cny_ktoken"},
                {"model_id": "qa-null-b", "discount_rate": "0", "currency": "CNY", "unit": "token"},
                {"model_id": "qa-null-c"},
            ],
        }
        job = request("/api/workspace/imports/drafts", {"bundle": proposal, "prompt": "Synthetic real-API fixture"})
        confirmed = request("/api/workspace/imports/" + job["id"] + "/confirm",
                            {"expected_revision": job["revision"], "idempotency_key": "isolated-proof-" + job["id"]})
        assert confirmed["status"] == "confirmed"
        before = snapshot()
        assert before["quotes?limit=500"]["total"] == 4
        items = before["quotes?limit=500"]["items"]
        assert all(item["vendor"] is None for item in items)
        assert all(item.get("derived_amount") is None for item in items)
        assert all(item.get("normalized") == {} for item in items)
        zero = next(item for item in items if item["model_id"] == "qa-null-b")
        assert zero["discount_rate"] == "0.00" and zero["pricing_unit_id"] is None
        assert zero["unit_quantity"] is None and zero["supplier"] is None
        both = next(item for item in items if item.get("supplier") == "Synthetic both")
        assert both["amount"] == "0.012" and both["discount_rate"] == "0.80"
        errors = []
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True, executable_path=args.chromium)
            page = browser.new_page(viewport={"width":1500, "height":960}, accept_downloads=True)
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(base + "/workspace", wait_until="networkidle")
            expect(page.locator("#quoteRows tr")).to_have_count(4)
            page.locator("#pricingConversionButton").click()
            expect(page.locator("#quoteRows")).to_contain_text("前端反算倍率 0.60 × / 6.0 折")
            expect(page.locator("#quoteRows")).to_contain_text("前端折后价 CNY 16")
            expect(page.locator("#quoteRows")).to_contain_text("缺失官网价")
            assert not {"null", "undefined"} & set(page.locator("#quoteRows td > span").all_text_contents())
            for format in ("csv", "xlsx"):
                page.locator("#exportFormat").select_option(format)
                with page.expect_download() as download:
                    page.locator("#exportFilteredButton").click()
                output = Path(download.value.path()).read_bytes()
                (args.evidence_dir / ("real-export." + format)).write_bytes(output)
                if format == "xlsx":
                    book = load_workbook(BytesIO(output), read_only=True)
                    rows = list(book.active.iter_rows(values_only=True))
                    assert len(rows) == 5
                    assert "原始付款倍率（两位小数）" in rows[0] and "前端折后价" in rows[0]
                    assert "原始折扣（折，一位小数）" in rows[0] and "前端反算折扣（折，一位小数）" in rows[0]
                    fold_index = rows[0].index("前端反算折扣（折，一位小数）")
                    assert any(row[fold_index] == "6.0" for row in rows[1:])
                    derived_index = rows[0].index("前端折后价")
                    rate_index = rows[0].index("前端反算付款倍率（两位小数）")
                    assert any(row[derived_index] == "16" for row in rows[1:])
                    assert any(row[rate_index] == "0.60" for row in rows[1:])
            page.screenshot(path=str(args.evidence_dir / "real-desktop.png"), full_page=True)
            page.set_viewport_size({"width":390, "height":844})
            assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
            page.screenshot(path=str(args.evidence_dir / "real-mobile.png"), full_page=True)
            browser.close()
        after = snapshot()
        assert before == after, "frontend conversion/export mutated data, catalog, context, memory or audit"
        assert not errors, errors
        evidence = {"passed":True, "kind":"synthetic_real_api_browser", "database":str(root / "workspace" / "quotes.sqlite3"),
                    "original_quotes":4, "audit_before":before["audit_count"], "audit_after":after["audit_count"],
                    "frontend_derived_amount":"16", "frontend_calculated_discount":"0.60",
                    "business_snapshot_unchanged":True, "real_model_called":False,
                    "production_accessed":False, "evidence_dir":str(args.evidence_dir)}
        (args.evidence_dir / "result.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps(evidence, ensure_ascii=False))
    finally:
        process.terminate()
        try: process.wait(timeout=10)
        except subprocess.TimeoutExpired: process.kill(); process.wait()
        log.close()


if __name__ == "__main__":
    main()
